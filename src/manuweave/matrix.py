"""Optional developer tests of source-pinned templates using installed local tools.

The matrix is an execution test, not an assertion of journal submission compliance.
It is not a prerequisite for normal Skill use. Every selected identifier receives
a persisted result, including failures.
"""
from __future__ import annotations

import ctypes
import copy
from datetime import datetime, timezone
import io
from pathlib import Path
import re
import shutil
import sys
import time

from .io import digest, write_json

MARKER = "MWBODY"
PROBES = {
    "MW_BODY_SIZE_PT": "body_size_pt", "MW_BODY_PT": "body_size_pt",
    "MW_BODY_FONT_FAMILY": "body_font_family", "MW_BODY_FONT_NAME": "body_font_name",
    "MW_FONT": "body_font_name", "MW_COLUMNS": "columns",
    "MW_PAPER_WIDTH_PT": "paper_width_pt", "MW_PAPER_HEIGHT_PT": "paper_height_pt",
    "MW_BASELINE_PT": "baseline_pt", "MW_TEXT_WIDTH_PT": "text_width_pt",
    "MW_TEXT_HEIGHT_PT": "text_height_pt",
}


def _status(statuses):
    statuses = list(statuses)
    if "FAIL" in statuses:
        return "FAIL"
    if not statuses or any(s in {"UNVERIFIED", "UNKNOWN"} for s in statuses):
        return "UNVERIFIED"
    return "PASS"


def parse_tex_metrics(log: str) -> dict:
    """Read final TeX engine log probes; retain conflicting observations."""
    observations = {}
    for match in re.finditer(r"(?m)^\s*(MW_[A-Z_]+)\s*=\s*([^\r\n]*)", log):
        key, value = match.groups()
        if key not in PROBES:
            continue
        canonical = PROBES[key]
        value = value.strip()
        if canonical not in {"body_font_family", "body_font_name"}:
            number = re.match(r"^(-?(?:\d+(?:\.\d*)?|\.\d+))(?:pt)?\s*$", value)
            value = float(number.group(1)) if number else value
        observations.setdefault(canonical, []).append(value)
    return {"actual": {key: values[-1] for key, values in observations.items()},
            "observations": observations,
            "conflicts": {key: values for key, values in observations.items() if len(set(values)) > 1}}


def compare_tex_metrics(parsed: dict, expected: dict) -> dict:
    actual, checks = parsed.get("actual", {}), []
    tolerance = expected.get("size_tolerance_pt", 0.03)
    for key in ("body_size_pt", "columns", "paper_width_pt", "paper_height_pt", "baseline_pt", "text_width_pt", "text_height_pt", "body_font_family"):
        if key not in expected:
            continue
        value, target = actual.get(key), expected[key]
        if value is None:
            status = "UNVERIFIED"
        elif key == "body_font_family":
            status = "PASS" if value == target else "FAIL"
        elif not isinstance(value, (int, float)):
            status = "FAIL"
        else:
            delta = 0 if key == "columns" else tolerance if key == "body_size_pt" else 0.1
            status = "PASS" if abs(value - target) <= delta else "FAIL"
        checks.append({"property": key, "status": status, "expected": target, "actual": value,
                       "unit": "tex_pt" if key.endswith("_pt") else None,
                       "tolerance": 0 if key == "columns" else tolerance if key == "body_size_pt" else 0.1})
    if "body_font_regex" in expected:
        value = actual.get("body_font_name")
        checks.append({"property": "body_font_name", "status": "UNVERIFIED" if value is None else
                       "PASS" if re.search(expected["body_font_regex"], str(value)) else "FAIL",
                       "expected_regex": expected["body_font_regex"], "actual": value})
    if parsed.get("conflicts"):
        checks.append({"property": "conflicting_probe_values", "status": "FAIL", "actual": parsed["conflicts"]})
    if not checks:
        checks.append({"property": "source_expected_metrics", "status": "UNVERIFIED",
                       "reason": "No source-established rendered metric expectations"})
    return {"status": _status(c["status"] for c in checks), "checks": checks, **parsed}


def pdf_body_metrics(pdf: str | Path, marker: str = MARKER) -> dict:
    """Measure actual glyphs at the body marker, not nominal XML properties."""
    import pypdfium2 as pdfium
    occurrences, text = [], []
    with pdfium.PdfDocument(pdf) as document:
        for number in range(len(document)):
            page = document[number]
            textpage = page.get_textpage()
            text.append(textpage.get_text_range())
            search = textpage.search(marker, match_case=True)
            while found := search.get_next():
                start, count = found
                glyphs = []
                for index in range(start, start + count):
                    flags = ctypes.c_int()
                    size = float(pdfium.raw.FPDFText_GetFontSize(textpage.raw, index))
                    length = pdfium.raw.FPDFText_GetFontInfo(textpage.raw, index, None, 0, ctypes.byref(flags))
                    buffer = ctypes.create_string_buffer(length) if 0 < length < 65536 else None
                    name = None
                    if buffer is not None:
                        pdfium.raw.FPDFText_GetFontInfo(textpage.raw, index, buffer, length, ctypes.byref(flags))
                        name = buffer.value.decode("utf-8", errors="replace")
                    glyphs.append({"index": index, "size_pt": size, "font_name": name, "font_flags": flags.value})
                occurrences.append({"page": number + 1, "start": start, "count": count, "glyphs": glyphs})
            search.close()
            textpage.close()
            page.close()
    return {"status": "PASS" if occurrences else "UNVERIFIED", "marker": marker,
            "unit": "pdf_bp", "occurrences": occurrences, "text": "\n".join(text)}


def pdf_page_boxes(pdf: str | Path, expected: dict | None = None) -> dict:
    """Record physical PDF boxes separately from nominal TeX length probes."""
    import pypdfium2 as pdfium
    expected = expected or {}
    pages = []
    with pdfium.PdfDocument(pdf) as document:
        for index in range(len(document)):
            page = document[index]
            boxes = {name: getattr(page, method)(fallback_ok=False) for name, method in
                     (("MediaBox", "get_mediabox"), ("CropBox", "get_cropbox"),
                      ("TrimBox", "get_trimbox"), ("BleedBox", "get_bleedbox"))}
            item = {"page": index + 1, **boxes, "effective_bbox": page.get_bbox()}
            nominal = {}
            media = boxes["MediaBox"]
            for dimension, key, indices in (("width", "paper_width_pt", (0, 2)), ("height", "paper_height_pt", (1, 3))):
                if key in expected and media is not None:
                    target = expected[key] * 72 / 72.27
                    actual = media[indices[1]] - media[indices[0]]
                    nominal[dimension] = {"nominal_tex_pt": expected[key], "nominal_pdf_bp": target,
                                          "actual_media_bp": actual, "delta_bp": actual - target}
            item["nominal_paper_comparison"] = nominal
            pages.append(item)
            page.close()
    return {"measurement_status": "PASS", "pdf_sha256": digest(pdf), "unit": "pdf_bp", "pages": pages,
            "physical_equivalence_status": "UNVERIFIED",
            "reason": "Nominal TeX paper probes do not establish identical output Media/Crop/Trim/Bleed boxes; source-specific driver and page-box review remains required.",
            "method": "PDFium explicit page boxes with fallbacks disabled; effective_bbox records the actual page viewport.",
            "inheritance_note": "PDFium box getters may not inherit parent page-tree entries; an absent direct box is recorded as null rather than guessed."}


def _font_key(name):
    from .fonts import pdf_font_family
    return pdf_font_family(name)


def _chinese_font(profile):
    name = profile.get("rules", {}).get("body", {}).get("east_asia_font")
    return name if name and (re.search(r"[\u3400-\u9fff]", name) or _font_key(name) in
                             {"simsun", "simhei", "kaiti", "fangsong", "microsoftyahei"}) else None


def compare_word_metrics(measured: dict, profile: dict) -> dict:
    target = profile.get("rules", {}).get("body", {})
    glyphs = [g for occurrence in measured.get("occurrences", []) for g in occurrence["glyphs"]]
    checks = []
    if "size_pt" in target:
        actual = sorted({round(g["size_pt"], 6) for g in glyphs})
        checks.append({"property": "body_size_pt", "expected": target["size_pt"], "actual": actual,
                       "unit": "pdf_bp", "tolerance": 0.06,
                       "status": "UNVERIFIED" if not actual else "PASS" if all(abs(v - target["size_pt"]) <= 0.06 + 1e-6 for v in actual) else "FAIL"})
    if "font" in target:
        actual = sorted({g["font_name"] for g in glyphs if g["font_name"]})
        expected = _font_key(target["font"])
        checks.append({"property": "body_font_name", "expected": target["font"], "actual": actual,
                       "status": "UNVERIFIED" if not actual else "PASS" if all(_font_key(v) == expected for v in actual) else "FAIL"})
    if _chinese_font(profile):
        east = [g for occurrence in measured.get("east_asia", {}).get("occurrences", []) for g in occurrence["glyphs"]]
        actual = sorted({g["font_name"] for g in east if g["font_name"]})
        expected = _font_key(target["east_asia_font"])
        checks.append({"property": "body_east_asia_font_name", "expected": target["east_asia_font"], "actual": actual,
                       "status": "UNVERIFIED" if not actual else "PASS" if all(_font_key(v) == expected for v in actual) else "FAIL"})
        if "size_pt" in target:
            actual_sizes = sorted({round(g["size_pt"], 6) for g in east})
            checks.append({"property": "body_east_asia_size_pt", "expected": target["size_pt"], "actual": actual_sizes,
                           "unit": "pdf_bp", "tolerance": 0.06,
                           "status": "UNVERIFIED" if not actual_sizes else "PASS" if all(abs(v - target["size_pt"]) <= 0.06 + 1e-6 for v in actual_sizes) else "FAIL"})
    if not checks:
        checks.append({"property": "body_metrics", "status": "NOT_APPLICABLE",
                       "reason": "The profile prescribes no body font family or size"})
    status = "NOT_APPLICABLE" if all(c["status"] == "NOT_APPLICABLE" for c in checks) else _status(c["status"] for c in checks)
    return {"status": status, "checks": checks,
            "actual": {"occurrences": measured.get("occurrences", []),
                       "east_asia_occurrences": measured.get("east_asia", {}).get("occurrences", []), "unit": "pdf_bp"}}


def reference_order(pdf_text: str, expected: list[str] | None) -> dict:
    if expected is None:
        return {"status": "NOT_APPLICABLE", "reason": "No source-established reference order expectation"}
    # Titles exist only in the synthetic bibliography; narrative citation order is different.
    names = [m.group(1).capitalize() for m in re.finditer(r"Synthetic\s+(Alpha|Zulu)\s+test\s+reference", pdf_text, re.I)]
    method = "Order of unique synthetic bibliography titles in extracted PDF text"
    if not names:
        # REVTeX styles legitimately omit article titles. Require the exact invented
        # journal on that same reference line, so narrative author mentions do not count.
        names = [m.group(1).capitalize() for m in re.finditer(
            r"(?<![A-Za-z])(Alpha|Zulu)\b[^\r\n]{0,160}?Invented\s+Fixture\s+Journal", pdf_text, re.I)]
        method = "Order of synthetic authors followed by the exact invented journal on the same bibliography line"
    observed = list(dict.fromkeys(names))
    return {"status": "UNVERIFIED" if not observed else "PASS" if observed == expected else "FAIL",
            "expected": expected, "actual": observed,
            "method": method}


def _word_fixture(destination: Path) -> tuple[Path, dict, dict]:
    """Create explicit synthetic roles, native math, merged cells and a real image."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.shared import Inches, Pt
    from PIL import Image, ImageDraw
    from .word import inspect_docx
    destination.mkdir(parents=True)
    source = destination / "source.docx"
    document = Document()
    labels = {}

    def paragraph(text, role):
        p = document.add_paragraph(text)
        labels[text] = role
        for run in p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(9)
        return p

    paragraph("Synthetic Template Matrix Fixture", "title")
    paragraph("Abstract", "abstract_heading")
    paragraph("This synthetic software fixture checks document mechanics and makes no research claim.", "abstract")
    paragraph("Keywords: synthetic fixture; local validation", "keywords")
    paragraph("1 Methods", "heading1")
    paragraph("MWBODY Synthetic body values are 12, 24 and 0.125. All are invented software test data.", "body")
    paragraph("MWCNBODY 中文正文仅用于字体和排版软件测试，不包含真实研究结果。", "body")
    paragraph("1.1 Native objects", "heading2")
    paragraph("The test retains editable native math, a merged table, an image and local citations [1, 2].", "body")
    p = document.add_paragraph()
    math = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    value = OxmlElement("m:t")
    value.text = "E=mc²"
    run.append(value)
    math.append(run)
    p._p.append(math)
    paragraph("Table 1. Synthetic fixture values.", "table_caption")
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Synthetic merged header"
    table.cell(0, 2).text = "Count"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "fixture-a"
    table.cell(1, 2).text = "12"
    table.cell(2, 0).text = "B"
    table.cell(2, 1).text = "fixture-b"
    table.cell(2, 2).text = "24"
    image = Image.new("RGB", (500, 200), "white")
    draw = ImageDraw.Draw(image)
    points = [(40, 165), (130, 140), (220, 110), (310, 70), (400, 40)]
    draw.line([(30, 20), (30, 180), (460, 180)], fill="black", width=3)
    draw.line(points, fill="navy", width=4)
    for x, y in points:
        draw.ellipse((x - 4, y - 4, x + 4, y + 4), fill="navy")
    data = io.BytesIO()
    image.save(data, "PNG")
    data.seek(0)
    document.add_picture(data, width=Inches(3.5))
    paragraph("Figure 1. Synthetic five-point software-test plot.", "figure_caption")
    paragraph("2 Conclusion", "heading1")
    paragraph("This fixture supports no scientific conclusion.", "body")
    paragraph("References", "references_heading")
    paragraph("[1] Synthetic Alpha. Synthetic Alpha test reference. Invented Fixture Journal, 2025.", "reference")
    paragraph("[2] Synthetic Zulu. Synthetic Zulu test reference. Invented Fixture Journal, 2026.", "reference")
    document.save(source)
    ir = inspect_docx(source)
    roles, evidence = {}, {}
    for block in ir["blocks"]:
        if block["text"] in labels:
            role = labels[block["text"]]
            roles[block["id"]] = role
            evidence[block["id"]] = {"role": role, "source": block["source"],
                                      "basis": "Exact synthetic fixture paragraph label", "quote": block["text"]}
        elif block["role"] in {"table_body", "equation"}:
            roles[block["id"]] = block["role"]
            evidence[block["id"]] = {"role": block["role"], "source": block["source"],
                                      "basis": "Native OOXML table/math structure", "quote": block["text"]}
    return source, roles, evidence


def _log_issues(log, root=None, main=None):
    patterns = {"overflow": r"Overfull \\[hv]box", "missing_glyph": r"Missing character:",
                "font_missing": r"Font.*not found", "undefined_reference": r"undefined references|Citation.*undefined|Reference.*undefined",
                "font_substitution": r"LaTeX Font Warning: Font shape .* undefined|Some font shapes were not available|defaults substituted",
                "compile_error": r"^!|^error:"}
    issues = [{"code": code, "message": line} for line in log.splitlines() for code, pattern in patterns.items() if re.search(pattern, line)]
    if root is not None and main is not None:
        from .templates import classify_native_diagnostic
        issues = [classify_native_diagnostic(root, main, issue) for issue in issues]
    return issues


def _run_profile(profile, directory):
    from . import pipeline
    from .fonts import check_fonts
    from .render import render
    from .templates import audit_requirements, materialize_fixture, native_template, validate_template
    directory.mkdir()
    result = {"id": profile["id"], "profile_version": profile["version"],
              "source_version": profile.get("template", {}).get("source_version"),
              "execution_status": "UNVERIFIED", "compliance_status": "UNVERIFIED",
              "visual_status": "UNVERIFIED", "layout_status": "UNVERIFIED", "issues": [], "directory": str(directory),
              "visual_reason": "Page images are retained for host review; the matrix does not inspect images."}
    write_json(directory / "profile.json", profile)
    try:
        if "template" in profile:
            validate_template(profile)
        native = native_template(profile)
        if native:
            from .latex import CLASS
            source = materialize_fixture(profile, directory / "input")
            original = source.read_text(encoding="utf-8")
            source.write_text(CLASS.sub(lambda _: r"\documentclass{article}", original, count=1), encoding="utf-8")
            result["engine"] = "latex"
            result["fixture_migration"] = {"source_class": "article", "target_class": profile["latex"]["documentclass"],
                                             "fixture": profile["template"]["fixture"]}
            roles, evidence = {}, {}
            output = directory / "native"
        else:
            source, roles, evidence = _word_fixture(directory / "input")
            # Distinct caption rules apply when present; generic profiles share caption.
            for block, role in list(roles.items()):
                if role in {"figure_caption", "table_caption"} and role not in profile["rules"]:
                    roles[block] = "caption"
                    evidence[block]["role"] = "caption"
            result["engine"] = "docx"
            output = directory / "editable.docx"
        result["source"] = str(source)
        result["input_hash"] = pipeline.source_identity(source)
        write_json(directory / "roles.json", {"roles": roles, "evidence": evidence})
        edit_plan = pipeline.plan(source, profile, roles)
        write_json(directory / "plan.json", edit_plan)
        applied = pipeline.apply(edit_plan, output)
        write_json(directory / "apply.json", applied)
        target = output / applied["main"] if "main" in applied else output
        result["output"] = str(target)
        result["output_hash"] = digest(output)
        audited = pipeline.audit(source, target, profile, roles)
        write_json(directory / "audit.json", audited)
        result["audit"] = audited
        result["requirements"] = audit_requirements(target, profile)
        result["compliance_status"] = result["requirements"]["status"]
        if not native:
            result["fonts"] = check_fonts(profile)
        else:
            result["fonts"] = {"status": "NOT_APPLICABLE", "reason": "Native fonts measured from the TeX log"}
        rendered = render(target, directory / "render", profile)
        write_json(directory / "render.json", rendered)
        result["render"] = rendered
        metrics, references = {"status": "UNVERIFIED", "checks": []}, {"status": "UNVERIFIED"}
        if rendered.get("pdf"):
            measured = pdf_body_metrics(rendered["pdf"], marker=MARKER if not native else "Synthetic")
            if not native and _chinese_font(profile):
                measured["east_asia"] = pdf_body_metrics(rendered["pdf"], marker="中文正文")
            write_json(directory / "pdf-measurements.json", measured)
            expected = profile.get("template", {}).get("expected_metrics", {})
            result["pdf_page_boxes"] = pdf_page_boxes(rendered["pdf"], expected)
            write_json(directory / "pdf-page-boxes.json", result["pdf_page_boxes"])
            references = reference_order(measured["text"], expected.get("reference_order"))
            if native:
                engine_log = directory / "render" / (target.stem + ".log")
                result["engine_log"] = str(engine_log)
                if engine_log.exists():
                    content = engine_log.read_text(encoding="utf-8", errors="replace")
                    result["issues"].extend(_log_issues(content, target.parent, target))
                    metrics = compare_tex_metrics(parse_tex_metrics(content), expected)
                else:
                    metrics = {"status": "UNVERIFIED", "checks": [], "reason": "Final TeX engine log is absent"}
            else:
                metrics = compare_word_metrics(measured, profile)
        result["rendered_metrics"], result["references"] = metrics, references
        result["issues"].extend(rendered.get("issues", []))
        statuses = [pipeline.check_status(audited, "content"), pipeline.check_status(audited, "style"),
                    result["fonts"]["status"], rendered.get("status", "UNVERIFIED"), metrics["status"], references["status"]]
        geometry = rendered.get("visual", {})
        if isinstance(geometry, dict) and geometry.get("geometry") == "FAIL":
            statuses.append("FAIL")
        if any(issue.get("severity") != "review" for issue in result["issues"]):
            statuses.append("FAIL")
        result["layout_status"] = ("UNVERIFIED" if any(issue.get("severity") == "review" for issue in result["issues"])
                                   else "FAIL" if any(issue.get("code") == "overflow" for issue in result["issues"])
                                   else rendered.get("layout_status", "PASS" if rendered.get("status") == "PASS" else "UNVERIFIED"))
        result["execution_status"] = _status(statuses)
        if result["requirements"]["status"] == "PASS":
            # Automatic class/XML checks cannot replace all visual and journal checks.
            result["compliance_status"] = "UNVERIFIED"
            result["compliance_reason"] = "Automated source-rule checks passed; complete page and submission-rule review remains pending."
        if result["requirements"]["status"] == "NOT_APPLICABLE":
            result["compliance_status"] = "UNVERIFIED"
            result["compliance_reason"] = "Generic engineering preset does not establish official submission compliance."
        if metrics["status"] not in {"PASS", "NOT_APPLICABLE"}:
            result["issues"].append({"code": "rendered_metrics", "status": metrics["status"], "checks": metrics.get("checks", [])})
        if references["status"] not in {"PASS", "NOT_APPLICABLE"}:
            result["issues"].append({"code": "reference_order", **references})
    except Exception as exc:
        result["execution_status"] = "FAIL"
        result["issues"].append({"code": "exception", "type": type(exc).__name__, "message": str(exc)})
    write_json(directory / "result.json", result)
    return result


def run_matrix(output: Path, names: list[str] | None = None) -> dict:
    """Developer-invoked tests using the installed native rendering adapters."""
    from .presets import list_presets, load_preset
    from .render import doctor
    output = Path(output).resolve()
    if output.exists():
        raise FileExistsError("Matrix output must be a new directory")
    catalog = list_presets()
    identifiers = [p["id"] for p in catalog] if names is None else [load_preset(name)["id"] for name in names]
    if len(identifiers) != len(set(identifiers)):
        raise ValueError("Duplicate matrix profile identifiers")
    output.mkdir(parents=True)
    installed = doctor().get("renderers", {})
    native_tools = {name: installed.get(name) for name in ("word_registered", "word_automation", "libreoffice")}
    native_tools.update({name: shutil.which(name) for name in ("latexmk", "xelatex", "lualatex", "pdflatex")})
    native_tools["python"] = sys.executable
    report = {"version": "1.0.0", "started_at": datetime.now(timezone.utc).isoformat(),
              "catalog_ids": [p["id"] for p in catalog], "selected_ids": identifiers,
              "all_catalog_selected": set(identifiers) == {p["id"] for p in catalog},
              "native_tools": native_tools, "results": [], "summary": {},
              "scope": "Actual native migrations, integrity audits, rendering and source-established metrics; visual review is separate."}
    for identifier in identifiers:
        start = time.monotonic()
        print(f"BEGIN {identifier}", file=sys.stderr, flush=True)
        try:
            profile = load_preset(identifier)
            result = _run_profile(profile, output / identifier)
        except Exception as exc:
            result = {"id": identifier, "execution_status": "FAIL", "compliance_status": "UNVERIFIED",
                      "visual_status": "UNVERIFIED", "issues": [{"code": "profile_load", "type": type(exc).__name__, "message": str(exc)}]}
        result["elapsed_seconds"] = round(time.monotonic() - start, 3)
        report["results"].append(result)
        report["summary"] = {status: sum(item["execution_status"] == status for item in report["results"])
                             for status in ("PASS", "FAIL", "UNVERIFIED")}
        report["summary"].update(selected=len(identifiers), completed=len(report["results"]),
                                 pending=len(identifiers) - len(report["results"]))
        write_json(output / "matrix.json", report)
        print(f"DONE {identifier} {result['execution_status']} {result['elapsed_seconds']}s", file=sys.stderr, flush=True)
    report["finished_at"] = datetime.now(timezone.utc).isoformat()
    report["status"] = _status(item["execution_status"] for item in report["results"])
    report["tested_ids"] = [item["id"] for item in report["results"]]
    report["untested_ids"] = sorted(set(identifiers) - set(report["tested_ids"]))
    write_json(output / "matrix.json", report)
    return report


def merge_matrices(output: Path, matrices: list[Path], supplements: list[Path] | None = None) -> dict:
    """Merge chronological real batches without rerunning unchanged successes.

    Every current catalog identifier remains visible; absent evidence becomes
    UNVERIFIED, and an earlier failure survives until a later actual batch replaces it.
    """
    from .io import read_json
    from .presets import list_presets
    output = Path(output).resolve()
    if output.exists():
        raise FileExistsError("Merged matrix output must be a new directory")
    identifiers = [p["id"] for p in list_presets()]
    latest, histories, sources = {}, {}, []
    for filename in matrices:
        filename = Path(filename).resolve()
        batch = read_json(filename)
        ids = [result["id"] for result in batch["results"]]
        if len(ids) != len(set(ids)):
            raise ValueError(f"Duplicate result ID in batch: {filename}")
        if set(batch.get("tested_ids", ids)) != set(ids):
            raise ValueError(f"Batch tested_ids disagree with actual results: {filename}")
        source = {"matrix": str(filename), "sha256": digest(filename), "result_count": len(ids),
                  "started_at": batch.get("started_at"), "finished_at": batch.get("finished_at")}
        sources.append(source)
        for result in batch["results"]:
            identifier = result["id"]
            if result.get("execution_status") not in {"PASS", "FAIL", "UNVERIFIED"}:
                raise ValueError(f"Invalid execution status: {identifier}")
            histories.setdefault(identifier, []).append({**source, "execution_status": result["execution_status"],
                                                         "directory": result.get("directory"), "profile_version": result.get("profile_version"),
                                                         "source_version": result.get("source_version")})
            latest[identifier] = copy.deepcopy(result)
            latest[identifier]["evidence_matrix"] = source
    supplemental = []
    for filename in supplements or []:
        filename = Path(filename).resolve()
        data = read_json(filename)
        supplemental.append({"path": str(filename), "sha256": digest(filename), "data": data})
    results, untested = [], []
    for identifier in identifiers:
        if identifier in latest:
            result = latest[identifier]
            result["run_history"] = histories[identifier]
        else:
            untested.append(identifier)
            result = {"id": identifier, "execution_status": "UNVERIFIED", "compliance_status": "UNVERIFIED",
                      "visual_status": "UNVERIFIED", "layout_status": "UNVERIFIED",
                      "issues": [{"code": "not_tested", "message": "No actual run evidence exists for this catalog ID."}], "run_history": []}
        result.setdefault("layout_status", "UNVERIFIED")
        results.append(result)
    summary = {status: sum(result["execution_status"] == status for result in results)
               for status in ("PASS", "FAIL", "UNVERIFIED")}
    summary.update(catalog=len(identifiers), tested=len(identifiers) - len(untested), untested=len(untested))
    report = {"version": "1.0.0", "finished_at": datetime.now(timezone.utc).isoformat(),
              "scope": "Latest actual execution evidence per catalog ID; execution, layout, visual review and source compliance are separate.",
              "merge_policy": "Input batches are chronological; retain all run histories and use each ID's latest actual result.",
              "status": _status(result["execution_status"] for result in results),
              "compliance_status": _status(result.get("compliance_status", "UNVERIFIED") for result in results),
              "visual_status": _status(result.get("visual_status", "UNVERIFIED") for result in results),
              "catalog_ids": identifiers, "tested_ids": [i for i in identifiers if i not in untested],
              "untested_ids": untested, "retired_ids": sorted(set(latest) - set(identifiers)),
              "summary": summary, "results": results, "source_batches": sources, "supplements": supplemental}
    output.mkdir(parents=True)
    write_json(output / "matrix.json", report)
    return report
