"""Exercise physical source measurements and preserve editable manuscript objects."""
import json
import runpy
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement

from manuweave.models import StyleProfile
from manuweave.presets import load_preset
from manuweave.templates import validate_template
from manuweave.word import apply_docx, inspect_docx

ROOT = Path(__file__).resolve().parents[1] / 'src/manuweave/data/templates/computer-special'
NAMES = ['comp-jmlr', 'comp-jfp-epi-final', 'comp-jair-2025', 'comp-lmcs']


def test_native_measurements_keep_physical_units_fonts_and_actual_role_variants():
    jmlr = load_preset('comp-jmlr')
    assert jmlr['rules']['body']['size_pt'] == pytest.approx(10.95 * 72 / 72.27)
    assert jmlr['extracted_format']['roles']['body']['size_pt']['source_value']['value'] == 10.95
    assert jmlr['rules']['heading1']['size_pt'] == pytest.approx(12 * 72 / 72.27)
    assert jmlr['extracted_format']['roles']['heading3']['small_caps']['value'] is True
    jfp = load_preset('comp-jfp-epi-final')
    assert jfp['rules']['body']['font'] == 'newtxtext'
    assert jfp['rules']['title']['italic'] is True
    assert jfp['page']['width_mm'] == pytest.approx(176, abs=1e-5)
    assert jfp['page']['height_mm'] == pytest.approx(250, abs=1e-5)
    assert jfp['rules']['reference']['size_pt'] == pytest.approx(8 * 72 / 72.27)
    jair = load_preset('comp-jair-2025')
    assert jair['rules']['body']['font'] == 'Linux Libertine'
    assert jair['rules']['heading1']['font'] == 'Linux Biolinum'
    assert jair['rules']['author']['size_pt'] == pytest.approx(12 * 72 / 72.27)
    assert jair['rules']['title']['size_pt'] == pytest.approx(14.4 * 72 / 72.27)
    lmcs = load_preset('comp-lmcs')
    assert lmcs['rules']['body']['first_line_pt'] == pytest.approx(20 * 72 / 72.27)
    assert lmcs['rules']['abstract']['size_pt'] == pytest.approx(9 * 72 / 72.27)


@pytest.mark.parametrize('identifier', NAMES)
def test_source_binding_and_real_application_preserve_formula_table_and_all_roles(tmp_path, identifier):
    profile = load_preset(identifier)
    StyleProfile.model_validate(profile, strict=True)
    validate_template(profile)
    assert profile['template']['assets'] == [] and profile['latex'] == {}
    assert len(profile['rules']) == 12
    for role, rules in profile['rules'].items():
        assert rules['line_spacing_rule'] == 'atLeast'
        for prop, expected in rules.items():
            fact = profile['extracted_format']['roles'][role][prop]
            assert fact['value'] == expected
            assert fact['unit'] and fact['source']['locator'] and fact['source']['sha256']
    doc = Document()
    roles = {}
    for index, role in enumerate(profile['rules']):
        doc.add_paragraph('Synthetic ' + role + ' constant 0.125')
        roles['p' + str(index)] = role
    math = OxmlElement('m:oMath')
    run, value = OxmlElement('m:r'), OxmlElement('m:t')
    value.text = 'E=mc²'
    run.append(value)
    math.append(run)
    doc.add_paragraph()._p.append(math)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = 'Merged header'
    table.cell(1, 0).text, table.cell(1, 1).text = 'A', '42'
    source, output, again = (tmp_path / name for name in ['source.docx', 'output.docx', 'again.docx'])
    doc.save(source)
    before = inspect_docx(source)
    result = apply_docx(source, output, profile, roles)
    assert result['audit']['content']['status'] == result['audit']['style']['status'] == 'PASS'
    after = inspect_docx(output)
    assert before['inventory'] == after['inventory']
    assert before['objects'] == after['objects']
    second = apply_docx(output, again, profile, roles)
    assert second['audit']['style']['status'] == 'PASS'
    assert inspect_docx(again)['inventory'] == after['inventory']


def test_builder_reproduces_json_without_publisher_assets():
    profiles = runpy.run_path(str(ROOT / 'builder_core_special.py'))['profiles']()
    assert {p['id'] for p in profiles} == set(NAMES)
    for profile in profiles:
        assert profile == json.loads((ROOT / 'profiles' / (profile['id'] + '.json')).read_text(encoding='utf-8'))
