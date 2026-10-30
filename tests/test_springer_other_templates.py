"""Detailed source values execute without publisher classes or downloaded assets."""
import json
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement

from manuweave.models import StyleProfile
from manuweave.presets import load_preset
from manuweave.templates import validate_template
from manuweave.word import apply_docx, inspect_docx

ROOT = Path(__file__).resolve().parents[1] / "src/manuweave/data/templates"
PATHS = sorted(p for family in ("springer", "oup", "computer-other") for p in (ROOT / family / "profiles").glob("*.json"))
FACTS = json.loads((ROOT / "computer-other/source-facts.json").read_text(encoding="utf-8"))
LAYOUT_PROFILES = {}
for path in PATHS:
    profile = json.loads(path.read_text(encoding="utf-8"))
    LAYOUT_PROFILES.setdefault(profile["extracted_format"]["mode"], profile)


@pytest.mark.parametrize("path", PATHS, ids=lambda p: p.stem)
def test_every_extracted_profile_has_confirmed_executable_property_evidence(path):
    profile = json.loads(path.read_text(encoding="utf-8"))
    StyleProfile.model_validate(profile, strict=True)
    validate_template(profile)
    assert profile["template"]["assets"] == []
    assert "remote_archive" not in profile["template"]
    assert "fixture" not in profile["template"]
    assert profile["latex"] == {}
    assert profile["extracted_format"]["mode"] in FACTS["layouts"]
    assert len(profile["extracted_format"]["roles"]) >= 10
    for target, properties in [*profile["rules"].items(), ("page", profile["page"]), ("sections", profile.get("sections", {}))]:
        for prop, expected in properties.items():
            matching = [r for r in profile["template"]["requirements"] if r["target"] == target and r["property"] == prop]
            assert matching and matching[0]["expected"] == expected
            assert matching[0]["certainty"] == "verified"
            assert matching[0]["source"]["sha256"]
            assert matching[0]["source"]["locator"]


def _fixture(path):
    doc = Document()
    doc.add_paragraph("Synthetic Formatting Validation", "Title")
    doc.add_paragraph("Abstract: Synthetic editable manuscript.")
    doc.add_paragraph("Keywords: synthetic, validation")
    doc.add_paragraph("1 Methods", "Heading 1")
    doc.add_paragraph("Protected values are 42 and 12.5; citation [1] remains supplied text.")
    doc.add_paragraph("1.1 Native Objects", "Heading 2")
    p = doc.add_paragraph()
    math = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.text = "E=mc²"
    run.append(text)
    math.append(run)
    p._p.append(math)
    doc.add_paragraph("Table 1 Synthetic values")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Synthetic group"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "42"
    doc.add_paragraph("Figure 1 Synthetic caption")
    doc.add_paragraph("References", "Heading 1")
    doc.add_paragraph("[1] Synthetic Author. Supplied test reference. 2026.")
    doc.save(path)


@pytest.mark.parametrize("mode", sorted(LAYOUT_PROFILES))
def test_every_distinct_layout_formats_docx_and_preserves_semantic_objects(mode, tmp_path):
    profile = LAYOUT_PROFILES[mode]
    source, output = tmp_path / "source.docx", tmp_path / "formatted.docx"
    _fixture(source)
    before = inspect_docx(source)
    result = apply_docx(source, output, profile)
    assert result["audit"]["content"]["status"] == "PASS"
    assert result["audit"]["style"]["status"] == "PASS", result["audit"]["style"]
    after = inspect_docx(output)
    assert before["objects"] == after["objects"]
    assert set(before["inventory"]) == set(after["inventory"])
    assert [b["text"] for b in before["blocks"]] == [b["text"] for b in after["blocks"]]


def test_distinct_publisher_values_and_units_are_retained():
    sn = load_preset("computer-springer-10994")
    assert sn["rules"]["title"]["size_pt"] == 17
    assert sn["rules"]["body"]["size_pt"] == 10
    assert sn["rules"]["body"]["line_spacing_rule"] == "atLeast"
    assert sn["rules"]["body"]["line_spacing_pt"] == 12
    assert sn["rules"]["figure_caption"]["size_pt"] == pytest.approx(8 * 72 / 72.27)
    assert sn["rules"]["table_caption"]["size_pt"] == 8
    comjnl = load_preset("computer-oup-comjnl")
    assert comjnl["rules"]["body"]["font"] == "Computer Modern Roman"
    assert comjnl["rules"]["title"]["font"] == "Computer Modern Sans Serif"
    assert comjnl["rules"]["reference"]["size_pt"] == 7.75
    assert comjnl["page"]["column_gap_mm"] == pytest.approx(6, abs=.0001)
    assert comjnl["sections"]["front_matter_columns"] == 1
    fcs = load_preset("computer-fcs-word-2025")
    assert fcs["rules"]["title"]["size_pt"] == 18
    assert fcs["rules"]["body"]["size_pt"] == 9
    assert fcs["extracted_format"]["page"]["column_widths_mm"]["value"] == pytest.approx([90.4875, 107.8441666667])
