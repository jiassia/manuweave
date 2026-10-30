"""Per-journal checks for source-confirmed engineering submission formats."""
from __future__ import annotations

import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches
from PIL import Image

from manuweave.models import StyleProfile
from manuweave.presets import load_preset
from manuweave.templates import validate_template
from manuweave.word import apply_docx, audit_docx, inspect_docx

DATA = Path(__file__).resolve().parents[1] / "src/manuweave/data/templates/chinese-engineering"
FACTS = json.loads((DATA / "source-facts.json").read_text(encoding="utf-8"))["documents"]
IDS = sorted(FACTS)


def fixture(directory):
    image = directory / "fixture.png"
    Image.new("RGB", (160, 60), "white").save(image)
    document = Document()
    document.add_paragraph("合成格式验证", "Title")
    document.add_paragraph("摘要：仅为软件测试，不报告科研结果。")
    document.add_paragraph("关键词：排版；保真；合成材料")
    document.add_paragraph("1 研究方法", "Heading 1")
    document.add_paragraph("原有数值 12、24 和 0.125 必须完整保留。[1]")
    document.add_paragraph("1.1 处理流程", "Heading 2")
    document.add_paragraph("原有正文内容不允许以排版为由删减。")
    document.add_paragraph("1.1.1 实现步骤", "Heading 3")
    paragraph = document.add_paragraph()
    equation = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.text = "E=mc²"
    run.append(text)
    equation.append(run)
    paragraph._p.append(equation)
    document.add_paragraph("表1 合成数值")
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "合成分组"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "12"
    table.cell(2, 0).text = "B"
    table.cell(2, 1).text = "24"
    document.add_picture(str(image), width=Inches(1.7))
    document.add_paragraph("图1 合成示意图")
    document.add_paragraph("参考文献", "Heading 1")
    document.add_paragraph("[1] 合成测试材料，非真实发表文献。2026。")
    path = directory / "original.docx"
    document.save(path)
    return path


@pytest.mark.parametrize("identity", IDS)
def test_each_journal_has_confirmed_official_sources_for_every_value(identity):
    profile = load_preset(identity)
    StyleProfile.model_validate(profile, strict=True)
    validate_template(profile)
    fact = FACTS[identity]
    assert profile["rules"] == fact["rules"]
    assert profile["page"] == fact["page"]
    assert profile["template"]["engine_support"] == {"docx": "partial", "latex": "unsupported"}
    assert profile["template"]["license"] == "rules_only_source_not_redistributed"
    assert profile["template"]["assets"] == []
    assert len(profile["sources"]) >= 2
    assert profile["rules"]["body"]["east_asia_font"] == "宋体"
    assert profile["rules"]["body"]["size_pt"] > 0
    assert profile["page"]["width_mm"] > 200 and profile["page"]["height_mm"] > 290
    for source in profile["sources"]:
        assert source["url"].startswith(("https://", "http://"))
        assert len(source["sha256"]) == 64 and source["bytes"] > 10000
        assert source["retrieved_at"] == "2026-10-30"
        assert source["original_file_redistributed"] is False
    requirements = profile["template"]["requirements"]
    assert requirements and all(r["certainty"] == "verified" for r in requirements)
    assert len({r["id"] for r in requirements}) == len(requirements)
    formats = [*profile["rules"].items(), ("page", profile["page"])]
    if profile.get("sections"):
        formats.append(("sections", profile["sections"]))
    for target, properties in formats:
        for prop, expected in properties.items():
            assert any(r["target"] == target and r["property"] == prop and r["expected"] == expected
                       and r["checking"] == "xml" for r in requirements)
    assert all(r["source"]["locator"] and r["source"]["url"] in {s["url"] for s in profile["sources"]} for r in requirements)


@pytest.mark.parametrize("identity", IDS)
def test_each_journal_applies_and_audits_native_objects_without_damage(identity, tmp_path):
    profile = load_preset(identity)
    source = fixture(tmp_path)
    output = tmp_path / "formatted.docx"
    result = apply_docx(source, output, profile)
    assert result["audit"]["content"]["status"] == "PASS"
    assert result["audit"]["style"]["status"] == "PASS"
    original, formatted = inspect_docx(source), inspect_docx(output)
    allowed_structural_parts = {"word/document.xml", "word/settings.xml"} if profile.get("sections") else set()
    assert {k: v for k, v in original["inventory"].items() if k not in allowed_structural_parts} == {
        k: v for k, v in formatted["inventory"].items() if k not in allowed_structural_parts}
    assert [b["text"] for b in original["blocks"] if b["text"]] == [b["text"] for b in formatted["blocks"] if b["text"]]
    assert formatted["objects"]["equations"] == formatted["objects"]["images"] == formatted["objects"]["tables"] == 1
    if profile.get("sections"):
        assert len(formatted["page"]) == 2
        assert [p["columns"] for p in formatted["page"]] == [1, 2]
    with ZipFile(source) as before, ZipFile(output) as after:
        assert all(before.read(name) == after.read(name) for name in before.namelist() if name.startswith("word/media/"))
        assert b"gridSpan" in after.read("word/document.xml")


@pytest.mark.parametrize("identity", IDS)
def test_each_journal_is_stable_when_applied_again(identity, tmp_path):
    profile = load_preset(identity)
    source = fixture(tmp_path)
    first, second = tmp_path / "first.docx", tmp_path / "second.docx"
    apply_docx(source, first, profile)
    apply_docx(first, second, profile)
    assert audit_docx(first, second, profile)["content"]["status"] == "PASS"
    assert audit_docx(first, second, profile)["style"]["status"] == "PASS"
    with ZipFile(first) as before, ZipFile(second) as after:
        assert set(before.namelist()) == set(after.namelist())
        assert all(before.read(name) == after.read(name) for name in before.namelist())


def test_journals_keep_their_actual_distinct_requirements():
    assert set(IDS) == {"cn-aroc-official-word", "cn-jcyyy-official-word", "cn-xddzjs-official-word", "cn-emt-official-word"}
    profiles = {identity: load_preset(identity) for identity in IDS}
    assert profiles["cn-aroc-official-word"]["rules"]["body"]["size_pt"] == 9
    assert profiles["cn-jcyyy-official-word"]["rules"]["body"]["line_spacing"] == 1.2
    assert profiles["cn-xddzjs-official-word"]["rules"]["body"]["line_spacing"] == 1.5
    assert profiles["cn-emt-official-word"]["rules"]["body"]["size_pt"] == 9.5
    assert profiles["cn-emt-official-word"]["rules"]["body"]["line_spacing"] == 1
    modern = profiles["cn-xddzjs-official-word"]
    assert all(modern["page"]["margin_" + edge + "_mm"] == 20 for edge in ["top", "bottom", "left", "right"])
    assert modern["sections"]["different_first_page"] is True
    assert modern["sections"]["different_odd_even_headers"] is True
    assert "columns" not in profiles["cn-jcyyy-official-word"]["page"]
    assert "margin_left_mm" not in profiles["cn-jcyyy-official-word"]["page"]
    assert not any(p.suffix.lower() in {".doc", ".docx", ".pdf", ".zip"} for p in DATA.rglob("*"))
