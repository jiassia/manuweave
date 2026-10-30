import zipfile

from docx import Document
from lxml import etree
import pytest

from manuweave.completeness import analyze_missing, apply_completion
from manuweave.word import inspect_docx


def manuscript(tmp_path):
    source = tmp_path / 'source.docx'
    doc = Document()
    doc.add_paragraph('A synthetic manuscript', 'Title')
    doc.add_paragraph('Methods', 'Heading 1')
    doc.add_paragraph('The method preserves native document objects and editable mathematics.')
    doc.add_table(rows=2, cols=2).cell(0, 0).text = 'Observed value 42'
    doc.save(source)
    return source


def block(source, text):
    return next(b for b in inspect_docx(source)['blocks'] if b['text'] == text)


def test_missing_categories(tmp_path):
    source = manuscript(tmp_path)
    report = analyze_missing(inspect_docx(source))
    assert {'missing_abstract', 'missing_keywords', 'missing_toc', 'missing_references'} <= {i['id'] for i in report['issues']}
    assert report['counts']['evidence_draft'] >= 2
    assert report['counts']['needs_user'] >= 1


def test_image_paragraph_and_native_object_share_one_caption(tmp_path):
    from PIL import Image

    image = tmp_path / 'one.png'
    Image.new('RGB', (100, 40), 'white').save(image)
    source = tmp_path / 'figure.docx'
    doc = Document()
    doc.add_paragraph('Synthetic methods', 'Heading 1')
    doc.add_picture(str(image))
    doc.add_paragraph('Figure 1. A synthetic test image.')
    doc.save(source)
    ir = inspect_docx(source)
    assert any(b['role'] == 'image' for b in ir['blocks'])
    assert len([b for b in ir['object_blocks'] if b['role'] == 'image']) == 1
    assert not [item for item in analyze_missing(ir)['issues'] if item['id'].startswith('missing_caption_')]


def test_native_toc_and_numbering_keep_text_and_table(tmp_path):
    source = manuscript(tmp_path)
    before = Document(source)
    report = apply_completion(source, tmp_path / 'completed.docx', {'add_toc': True, 'number_headings': True})
    result = Document(tmp_path / 'completed.docx')
    assert [p.text for p in before.paragraphs] == [p.text for index, p in enumerate(result.paragraphs) if index != 1]
    assert result.tables[0].cell(0, 0).text == 'Observed value 42'
    with zipfile.ZipFile(tmp_path / 'completed.docx') as archive:
        assert b'TOC ' in archive.read('word/document.xml')
        assert b'numPr' in archive.read('word/document.xml')
        etree.fromstring(archive.read('word/numbering.xml'))
    assert report['status'] == 'PASS'
    apply_completion(tmp_path / 'completed.docx', tmp_path / 'twice.docx', {'add_toc': True, 'number_headings': True})
    with zipfile.ZipFile(tmp_path / 'twice.docx') as archive:
        assert archive.read('word/document.xml').count(b'TOC ') == 1


def test_evidence_bound_completion(tmp_path):
    source = manuscript(tmp_path)
    quote = 'The method preserves native document objects and editable mathematics.'
    b = block(source, quote)
    request = {'enabled_content': True, 'drafts': [{'role': 'abstract', 'text': quote,
               'evidence': [{'block_id': b['id'], 'quote': quote}]}]}
    output = tmp_path / 'extractive.docx'
    result = apply_completion(source, output, request)
    assert result['changes'][0]['operation'] == 'insert_extractive_draft'
    assert Document(output).paragraphs[0].text == 'A synthetic manuscript'
    assert Document(output).paragraphs[1].text == 'Abstract: ' + quote
    assert source.read_bytes() != output.read_bytes()
    assert output.with_suffix('.docx.completion.json').is_file()


def test_fabricated_numbers_remain_outside_manuscript(tmp_path):
    source = manuscript(tmp_path)
    quote = 'The method preserves native document objects and editable mathematics.'
    b = block(source, quote)
    output = tmp_path / 'pending.docx'
    result = apply_completion(source, output, {'enabled_content': True, 'drafts': [
        {'role': 'abstract', 'text': 'Our experiment achieved 99.9% accuracy [12].',
         'evidence': [{'block_id': b['id'], 'quote': quote}]}]})
    assert result['pending_drafts'][0]['status'] == 'AWAITING_REVIEW'
    assert '99.9' not in '\n'.join(p.text for p in Document(output).paragraphs)
    assert result['changes'] == []


def test_no_evidence_no_write_and_explicit_mode(tmp_path):
    source = manuscript(tmp_path)
    output = tmp_path / 'rejected.docx'
    with pytest.raises(ValueError):
        apply_completion(source, output, {'drafts': [{'role': 'abstract', 'text': 'New result',
                          'evidence': [{'block_id': 'p2', 'quote': 'No such quote'}]}]})
    assert not output.exists()
    with pytest.raises(ValueError):
        apply_completion(source, output, {'enabled_content': True, 'drafts': [
            {'role': 'abstract', 'text': 'New result', 'evidence': [{'block_id': 'p2', 'quote': 'No such quote'}]}]})
    assert not output.exists()
    with pytest.raises(ValueError):
        apply_completion(source, source, {'add_toc': True})
    with pytest.raises(ValueError):
        apply_completion(source, output, {'execute_python': 'print(1)'})


def test_caption_uses_explicit_text(tmp_path):
    source = manuscript(tmp_path)
    tables = [b for b in inspect_docx(source)['object_blocks'] if b['role'] == 'table']
    assert tables
    output = tmp_path / 'captioned.docx'
    result = apply_completion(source, output, {'captions': [
        {'block_id': tables[0]['id'], 'text': 'Synthetic observations', 'kind': 'table'}]})
    assert result['changes'][0]['text'] == 'Table 1. Synthetic observations'


def test_empty_headings_do_not_count_as_completion():
    report = analyze_missing({'blocks': [
        {'id': 'p0', 'role': 'abstract_heading', 'text': 'Abstract'},
        {'id': 'p1', 'role': 'keywords', 'text': 'Keywords:'},
        {'id': 'p2', 'role': 'references_heading', 'text': 'References'},
    ]})
    ids = {i['id'] for i in report['issues']}
    assert {'missing_abstract', 'missing_keywords', 'missing_references'} <= ids


def test_toc_caches_real_headings_and_marks_correct_outline_nodes(tmp_path):
    source = tmp_path / 'direct.docx'
    doc = Document()
    doc.add_paragraph('Synthetic title', 'Title')
    doc.add_paragraph('1 Methods')
    doc.add_paragraph('Observed data 42 remain unchanged.')
    doc.add_paragraph('1.1 Native objects')
    doc.save(source)
    output = tmp_path / 'toc.docx'
    report = apply_completion(source, output, {'add_toc': True, 'number_headings': True})
    with zipfile.ZipFile(output) as archive:
        root = etree.fromstring(archive.read('word/document.xml'))
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    cache = root.xpath('.//w:fldSimple/w:r/w:t/text()', namespaces=ns)
    assert cache == ['1 Methods', '1.1 Native objects']
    for text, level in [('1 Methods', '0'), ('1.1 Native objects', '1')]:
        node = root.xpath('.//w:p[w:r/w:t=$text]', text=text, namespaces=ns)[0]
        assert node.xpath('w:pPr/w:outlineLvl/@w:val', namespaces=ns) == [level]
    title = root.xpath('.//w:p[w:r/w:t="Synthetic title"]', namespaces=ns)[0]
    assert not title.xpath('w:pPr/w:outlineLvl', namespaces=ns)
    assert report['integrity']['status'] == 'PASS'
    again = tmp_path / 'toc-again.docx'
    second = apply_completion(output, again, {'add_toc': True, 'number_headings': True})
    assert second['changes'] == []
    with zipfile.ZipFile(again) as archive:
        assert archive.read('word/document.xml') == zipfile.ZipFile(output).read('word/document.xml')


def test_empty_abstract_heading_receives_evidence_text(tmp_path):
    source = tmp_path / 'empty.docx'
    doc = Document()
    doc.add_paragraph('Synthetic title', 'Title')
    doc.add_paragraph('Abstract:')
    doc.add_paragraph('Methods', 'Heading 1')
    quote = 'The method preserves native document objects.'
    doc.add_paragraph(quote)
    doc.save(source)
    b = block(source, quote)
    output = tmp_path / 'filled.docx'
    report = apply_completion(source, output, {'enabled_content': True, 'drafts': [
        {'role': 'abstract', 'text': quote, 'evidence': [{'block_id': b['id'], 'quote': quote}]}]})
    paras = [p.text for p in Document(output).paragraphs]
    assert paras[1:3] == ['Abstract:', quote]
    assert report['integrity']['status'] == 'PASS'


def test_completion_audits_native_equation_image_merged_table_and_notes(tmp_path):
    import io
    from PIL import Image
    from docx.oxml import OxmlElement
    source = tmp_path / 'native.docx'
    doc = Document()
    doc.add_paragraph('Native source', 'Title')
    doc.add_paragraph('Methods', 'Heading 1')
    equation = OxmlElement('m:oMath')
    run = OxmlElement('m:r')
    text = OxmlElement('m:t')
    text.text = 'x=42'
    run.append(text)
    equation.append(run)
    doc.add_paragraph()._p.append(equation)
    image = io.BytesIO()
    Image.new('RGB', (32, 32), 'navy').save(image, format='PNG')
    image.seek(0)
    doc.add_picture(image)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = 'Merged 42'
    doc.save(source)
    with zipfile.ZipFile(source, 'a') as archive:
        archive.writestr('word/footnotes.xml', '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:footnote w:id="1"><w:p><w:r><w:t>Note 42</w:t></w:r></w:p></w:footnote></w:footnotes>')
    output = tmp_path / 'native-completed.docx'
    report = apply_completion(source, output, {'add_toc': True})
    assert report['integrity']['status'] == 'PASS'
    assert {x['name'] for x in report['integrity']['checks']} >= {'equations', 'images', 'tables', 'media_footnotes_endnotes_and_other_parts'}
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(output) as after:
        for name in before.namelist():
            if name.startswith('word/media/') or name == 'word/footnotes.xml':
                assert before.read(name) == after.read(name)


def test_toc_refuses_empty_placeholder_without_headings(tmp_path):
    source = tmp_path / 'no-headings.docx'
    doc = Document()
    doc.add_paragraph('Title')
    doc.add_paragraph('A body paragraph without section hierarchy.')
    doc.save(source)
    with pytest.raises(ValueError, match='without headings'):
        apply_completion(source, tmp_path / 'bad.docx', {'add_toc': True})
    assert not (tmp_path / 'bad.docx').exists()


def test_frontmatter_order_is_title_abstract_keywords_toc_body(tmp_path):
    source = manuscript(tmp_path)
    quote = 'The method preserves native document objects and editable mathematics.'
    b = block(source, quote)
    output = tmp_path / 'ordered.docx'
    apply_completion(source, output, {'enabled_content': True, 'add_toc': True, 'drafts': [
        {'role': 'keywords', 'text': 'native document objects', 'evidence': [{'block_id': b['id'], 'quote': 'native document objects'}]},
        {'role': 'abstract', 'text': quote, 'evidence': [{'block_id': b['id'], 'quote': quote}]}]})
    texts = [p.text for p in Document(output).paragraphs]
    assert texts[0] == 'A synthetic manuscript'
    assert texts[1] == 'Abstract: ' + quote
    assert texts[2] == 'Keywords: native document objects'
    with zipfile.ZipFile(output) as archive:
        root = etree.fromstring(archive.read('word/document.xml'))
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    paragraphs = root.findall('w:body/w:p', ns)
    assert paragraphs[3].find('w:fldSimple', ns) is not None
    assert texts[4] == 'Methods'


def test_page_numbers_create_native_footer_for_every_section_and_are_stable(tmp_path):
    from docx.enum.section import WD_SECTION_START
    source = tmp_path / 'sections.docx'
    doc = Document()
    doc.add_paragraph('Section one', 'Heading 1')
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    doc.add_paragraph('Section two', 'Heading 1')
    doc.save(source)
    original = source.read_bytes()
    output = tmp_path / 'numbered.docx'
    result = apply_completion(source, output, {'add_page_numbers': True})
    assert result['integrity']['status'] == 'PASS'
    assert source.read_bytes() == original
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
          'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
    with zipfile.ZipFile(output) as archive:
        root = etree.fromstring(archive.read('word/document.xml'))
        references = root.xpath('.//w:sectPr/w:footerReference[@w:type="default"]/@r:id', namespaces=ns)
        assert len(references) == 2
        assert len(set(references)) == 1
        rels = etree.fromstring(archive.read('word/_rels/document.xml.rels'))
        rel = next(child for child in rels if child.get('Id') == references[0])
        target = 'word/' + rel.get('Target')
        footer = etree.fromstring(archive.read(target))
        assert footer.xpath('.//w:fldSimple/@w:instr', namespaces=ns) == [' PAGE ']
        types = etree.fromstring(archive.read('[Content_Types].xml'))
        assert any(child.get('PartName') == '/' + target for child in types)
    second = tmp_path / 'numbered-again.docx'
    rerun = apply_completion(output, second, {'add_page_numbers': True})
    assert rerun['changes'] == []
    with zipfile.ZipFile(output) as before, zipfile.ZipFile(second) as after:
        assert before.namelist() == after.namelist()
        for part in before.namelist():
            assert before.read(part) == after.read(part)


def test_page_numbers_preserve_existing_footer_and_section_layout(tmp_path):
    from docx.enum.section import WD_SECTION_START
    source = tmp_path / 'existing-footers.docx'
    doc = Document()
    doc.add_paragraph('Original text 42')
    doc.sections[0].footer.paragraphs[0].text = 'Existing footer note 42'
    doc.sections[0].different_first_page_header_footer = True
    doc.sections[0].first_page_footer.paragraphs[0].text = 'First page author notice'
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    section.footer.is_linked_to_previous = False
    section.footer.paragraphs[0].text = 'Second footer 17'
    doc.add_paragraph('Second section 17')
    doc.save(source)
    output = tmp_path / 'with-pages.docx'
    result = apply_completion(source, output, {'add_page_numbers': True})
    assert result['integrity']['status'] == 'PASS'
    after = Document(output)
    assert after.sections[0].footer.paragraphs[0].text == 'Existing footer note 42'
    assert after.sections[0].first_page_footer.paragraphs[0].text == 'First page author notice'
    assert after.sections[1].footer.paragraphs[0].text == 'Second footer 17'
    assert after.sections[0].different_first_page_header_footer
    with zipfile.ZipFile(output) as archive:
        footer_parts = [name for name in archive.namelist() if name.startswith('word/footer') and name.endswith('.xml')]
        assert len(footer_parts) == 3
        assert all(b' PAGE ' in archive.read(name) for name in footer_parts)
    assert all(check['status'] == 'PASS' for check in result['integrity']['checks'])


def test_page_numbers_do_not_duplicate_existing_page_field(tmp_path):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    source = tmp_path / 'already-numbered.docx'
    doc = Document()
    doc.add_paragraph('A source paragraph')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), ' PAGE ')
    doc.sections[0].footer.paragraphs[0]._p.append(field)
    doc.save(source)
    output = tmp_path / 'unchanged-footer.docx'
    result = apply_completion(source, output, {'add_page_numbers': True})
    assert not any(change['operation'] == 'insert_page_field' for change in result['changes'])
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(output) as after:
        assert before.read('word/footer1.xml') == after.read('word/footer1.xml')
