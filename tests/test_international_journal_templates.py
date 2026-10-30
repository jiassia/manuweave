"""Journal-specific, located source rules and conservative native Word execution."""
import json
import re
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches
from PIL import Image

from manuweave.models import StyleProfile
from manuweave.presets import load_preset
from manuweave.templates import audit_requirements, validate_template
from manuweave.word import apply_docx, audit_docx, inspect_docx

DATA = Path(__file__).resolve().parents[1] / "src/manuweave/data/templates/international-journals"
PROFILES = sorted((DATA / "profiles").glob("*.json"))
FACTS = json.loads((DATA / "source-facts.json").read_text(encoding="utf-8"))["documents"]


def fixture(directory, abstract_citation=False):
    image = directory / "synthetic.png"
    Image.new("RGB", (100, 30), "blue").save(image)
    doc = Document()
    doc.add_paragraph("Synthetic international journal fixture", "Title")
    doc.add_paragraph("Synthetic Fixture Author, Synthetic Test Institution")
    doc.add_paragraph("Abstract: This synthetic document validates native formatting and reports no research findings." +
                      (" A prohibited citation appears here [1]." if abstract_citation else ""))
    doc.add_paragraph("Keywords: synthetic; native format; software testing")
    doc.add_paragraph("1 Introduction", "Heading 1")
    doc.add_paragraph("The protected synthetic count is 42, with citation [1].")
    doc.add_paragraph("2 Methods", "Heading 1")
    paragraph = doc.add_paragraph()
    math = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.text = "E=mc²"
    run.append(text)
    math.append(run)
    paragraph._p.append(math)
    doc.add_paragraph("Table 1. Synthetic values.")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Synthetic groups"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "12"
    table.cell(2, 0).text = "B"
    table.cell(2, 1).text = "24"
    doc.add_picture(str(image), width=Inches(1.5))
    doc.add_paragraph("Figure 1. Synthetic image.")
    doc.add_paragraph("3 Conclusions", "Heading 1")
    doc.add_paragraph("No scientific conclusion follows from this software fixture.")
    doc.add_paragraph("References", "Heading 1")
    doc.add_paragraph("[1] Synthetic Fixture Author. Invented software-test reference. 2026.")
    source = directory / "source.docx"
    doc.save(source)
    return source


@pytest.mark.parametrize("path", PROFILES, ids=lambda p: p.stem)
def test_each_journal_profile_has_actual_located_official_source(path):
    profile = json.loads(path.read_text(encoding="utf-8"))
    StyleProfile.model_validate(profile, strict=True)
    validate_template(profile)
    assert load_preset(profile["id"])["id"] == profile["id"]
    facts = FACTS[profile["id"]]
    assert profile["sources"] == facts["sources"]
    assert profile["template"]["engine_support"] == {"docx": "partial", "latex": "unsupported"}
    assert profile["template"]["assets"] == []
    assert profile["template"]["license"] == "rules_only_source_not_redistributed"
    assert all(source["http_status"] == 200 and source["byte_count"] > 10000
               and re.fullmatch(r"[0-9a-f]{64}", source["sha256"])
               and source["source_verified"] is True
               and source["submission_compliance_verified"] is False
               and source["original_file_redistributed"] is False
               and source["retrieved_at"] == "2026-10-30" for source in profile["sources"])
    requirements = profile["template"]["requirements"]
    assert all(rule["certainty"] == "verified" for rule in requirements)
    for rule in requirements:
        assert rule["source"]["url"] in {source["url"] for source in profile["sources"]}
        assert rule["source"]["response_sha256"] in {source["sha256"] for source in profile["sources"]}
        assert rule["source"]["locator"].startswith("/html/")
    for role, properties in profile["rules"].items():
        assert properties == {"line_spacing": facts["spacing"]}
        assert any(rule["target"] == role and rule["property"] == "line_spacing" and rule["expected"] == facts["spacing"]
                   and rule["checking"] == "xml" for rule in requirements)
    assert "font" not in profile["rules"]["body"]
    assert "size_pt" not in profile["rules"]["body"]
    assert not any(key.startswith("margin_") for key in profile["page"])


@pytest.mark.parametrize("path", PROFILES, ids=lambda p: p.stem)
def test_each_profile_native_apply_audit_and_repeat_preserve_all_objects(path, tmp_path):
    profile = json.loads(path.read_text(encoding="utf-8"))
    source = fixture(tmp_path)
    first, second = tmp_path / "formatted.docx", tmp_path / "repeated.docx"
    result = apply_docx(source, first, profile)
    assert result["audit"]["content"]["status"] == result["audit"]["style"]["status"] == "PASS"
    original, formatted = inspect_docx(source), inspect_docx(first)
    assert original["inventory"] == formatted["inventory"]
    assert [block["text"] for block in original["blocks"]] == [block["text"] for block in formatted["blocks"]]
    assert formatted["objects"]["equations"] == formatted["objects"]["tables"] == formatted["objects"]["images"] == 1
    for block in formatted["blocks"]:
        if block["role"] in profile["rules"] and block["text"].strip():
            assert abs(block["format"]["line_spacing"] - FACTS[profile["id"]]["spacing"]) < .01
    apply_docx(first, second, profile)
    assert audit_docx(first, second, profile)["style"]["status"] == "PASS"
    with ZipFile(source) as before, ZipFile(first) as after:
        assert all(before.read(name) == after.read(name) for name in before.namelist() if name.startswith("word/media/"))
        assert b"gridSpan" in after.read("word/document.xml")
    with ZipFile(first) as before, ZipFile(second) as after:
        assert before.read("word/document.xml") == after.read("word/document.xml")


@pytest.mark.parametrize("path", PROFILES, ids=lambda p: p.stem)
def test_each_explicit_abstract_citation_ban_detects_real_violation(path, tmp_path):
    profile = json.loads(path.read_text(encoding="utf-8"))
    source = fixture(tmp_path, abstract_citation=True)
    result = audit_requirements(source, profile)
    citation = next(check for check in result["checks"] if check["id"] == "abstract-citations")
    assert citation["status"] == "FAIL"
    assert result["status"] == "FAIL"


def test_catalog_has_distinct_confirmed_manuscript_requirements_without_source_redistribution():
    assert len(PROFILES) == 4
    assert FACTS["int-frontiers-public-health-word"]["spacing"] == 1
    assert FACTS["int-bmc-public-health-research-word"]["spacing"] == 2
    assert FACTS["int-plos-one-research-word"]["columns"] == 1
    comp = FACTS["int-plos-comp-biol-research-word"]
    assert next(rule for rule in comp["policies"] if rule["id"] == "author-summary")["expected"] == [150, 200]
    assert not any(path.suffix.lower() in {".html", ".htm", ".doc", ".docx", ".pdf"} for path in DATA.rglob("*"))
