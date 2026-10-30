"""Source-backed food science and statistics manuscript-format profiles."""
import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches
from PIL import Image

from manuweave.models import StyleProfile
from manuweave.presets import load_preset
from manuweave.templates import validate_template
from manuweave.word import apply_docx, audit_docx, inspect_docx

DATA = Path(__file__).resolve().parents[1] / "src/manuweave/data/templates/chinese-multidisciplinary"
FACTS = json.loads((DATA / "source-facts.json").read_text(encoding="utf-8"))["documents"]
PROFILES = sorted((DATA / "profiles").glob("*.json"))


def fixture(directory):
    image = directory / "synthetic.png"
    Image.new("RGB", (90, 30), "white").save(image)
    doc = Document()
    doc.add_paragraph("合成排版测试", "Title")
    doc.add_paragraph("摘要：本段为软件测试材料，不表示研究结果。")
    doc.add_paragraph("关键词：排版；验证；合成")
    doc.add_paragraph("1 测试方法", "Heading 1")
    doc.add_paragraph("受保护数值为 12 和 24，引用标记为 [1]。")
    doc.add_paragraph("1.1 对象保留", "Heading 2")
    p = doc.add_paragraph()
    math = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.text = "E=mc²"
    run.append(text)
    math.append(run)
    p._p.append(math)
    doc.add_paragraph("表1 合成数值")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "合成分组"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "12"
    doc.add_picture(str(image), width=Inches(1))
    doc.add_paragraph("图1 合成示意图")
    doc.add_paragraph("参考文献", "Heading 1")
    doc.add_paragraph("[1] 合成作者. 本条为软件测试材料. 2026.")
    path = directory / "source.docx"
    doc.save(path)
    return path


@pytest.mark.parametrize("path", PROFILES, ids=lambda path: path.stem)
def test_each_profile_has_exact_confirmed_sources(path):
    profile = json.loads(path.read_text(encoding="utf-8"))
    StyleProfile.model_validate(profile, strict=True)
    validate_template(profile)
    assert load_preset(profile["id"])["id"] == profile["id"]
    fact = FACTS[profile["id"]]
    assert profile["template"]["source_version"] == fact["source_version"]
    assert profile["template"]["license"] == "rules_only_source_not_redistributed"
    assert profile["template"]["assets"] == []
    assert profile["template"]["engine_support"] == {"docx": "partial", "latex": "unsupported"}
    assert all(len(source["sha256"]) == 64 and source["bytes"] > 10000 for source in profile["sources"])
    assert all(source["retrieved_at"] == "2026-10-30" and not source["original_file_redistributed"] for source in profile["sources"])
    for requirement in profile["template"]["requirements"]:
        assert requirement["certainty"] == "verified"
        assert requirement["source"]["locator"]
        assert requirement["source"]["url"] in {s["url"] for s in profile["sources"]}
    for role, properties in profile["rules"].items():
        for prop, value in properties.items():
            assert any(r["checking"] == "xml" and r["target"] == role and r["property"] == prop and r["expected"] == value
                       for r in profile["template"]["requirements"])


@pytest.mark.parametrize("path", PROFILES, ids=lambda path: path.stem)
def test_each_profile_apply_audit_and_repeat_preserve_content(path, tmp_path):
    profile = json.loads(path.read_text(encoding="utf-8"))
    source = fixture(tmp_path)
    result = apply_docx(source, tmp_path / "first.docx", profile)
    assert result["audit"]["content"]["status"] == result["audit"]["style"]["status"] == "PASS"
    before, after = inspect_docx(source), inspect_docx(tmp_path / "first.docx")
    assert before["inventory"] == after["inventory"]
    assert before["objects"] == after["objects"]
    assert [p["text"] for p in before["blocks"]] == [p["text"] for p in after["blocks"]]
    apply_docx(tmp_path / "first.docx", tmp_path / "second.docx", profile)
    assert audit_docx(tmp_path / "first.docx", tmp_path / "second.docx", profile)["style"]["status"] == "PASS"
    with ZipFile(tmp_path / "first.docx") as first, ZipFile(tmp_path / "second.docx") as second:
        assert all(first.read(name) == second.read(name) for name in first.namelist())
        assert b"gridSpan" in second.read("word/document.xml")


def test_only_explicit_confirmed_rules_are_bundled():
    assert len(PROFILES) == 3
    industry = load_preset("cn-foodindustry-word-confirmed")
    assert not any("size_pt" in rule for rule in industry["rules"].values())
    assert next(r for r in industry["template"]["requirements"] if r["id"] == "font-size-policy")["expected"] is False
    stats = load_preset("cn-statistics-decision-word-confirmed")
    assert stats["rules"]["body"] == {"east_asia_font": "宋体", "size_pt": 12.0, "line_spacing": 1.0}
    assert stats["page"] == {"columns": 1}
    science = load_preset("cn-foodscience-word-2020")
    assert science["rules"]["body"]["size_pt"] == 9.5
    assert science["rules"]["title"]["size_pt"] == 24.0
    assert not any(p.suffix in {".doc", ".docx", ".pdf", ".html"} for p in DATA.rglob("*"))
