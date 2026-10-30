"""Physical conversion and preservation checks for Cambridge and ASL layout facts."""
import pytest
from docx import Document
from manuweave.presets import load_preset
from manuweave.models import StyleProfile
from manuweave.templates import validate_template
from manuweave.word import apply_docx, inspect_docx

NAMES = ['comp-cambridge-mscs', 'comp-cambridge-tplp', 'comp-asl-jsl']


def test_original_classes_retain_distinct_native_typography_and_units():
    msc, tlp, asl = [load_preset(name) for name in NAMES]
    assert msc['rules']['body']['font'] == 'Times Roman'
    assert msc['rules']['heading1']['font'] == 'Helvetica'
    assert msc['rules']['body']['size_pt'] == pytest.approx(10.25 * 72 / 72.27)
    assert msc['rules']['figure_caption']['size_pt'] == pytest.approx(7.75 * 72 / 72.27)
    assert tlp['rules']['body']['font'] == asl['rules']['body']['font'] == 'Computer Modern Roman'
    assert tlp['rules']['title']['italic'] is True
    assert tlp['rules']['title']['size_pt'] == pytest.approx(17.28 * 72 / 72.27)
    assert tlp['rules']['heading1']['alignment'] == 'center'
    assert asl['rules']['abstract']['size_pt'] == pytest.approx(7 * 72 / 72.27)
    assert asl['extracted_format']['roles']['heading2']['run_in']['value'] is True
    assert msc['page']['width_mm'] == pytest.approx(6.85 * 25.4)
    assert msc['extracted_format']['native_page']['PW']['value'] == 0
    assert msc['extracted_format']['measurement']['exit_code'] == 1
    assert tlp['extracted_format']['measurement']['exit_code'] == 0


@pytest.mark.parametrize('identifier', NAMES)
def test_application_keeps_native_objects_and_values_with_minimum_baselines(tmp_path, identifier):
    source = tmp_path / 'original.docx'
    document = Document()
    document.add_paragraph('Original software fixture', 'Title')
    document.add_paragraph('Abstract: Preserved value 0.125.')
    document.add_paragraph('Methods', 'Heading 1')
    document.add_paragraph('Preserved value 42.')
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = 'Merged native table'
    table.cell(1, 0).text = '12'
    document.save(source)
    profile = load_preset(identifier)
    StyleProfile.model_validate(profile, strict=True)
    validate_template(profile)
    assert profile['template']['assets'] == []
    assert all(role['line_spacing_rule'] == 'atLeast' for role in profile['rules'].values())
    output = tmp_path / 'formatted.docx'
    result = apply_docx(source, output, profile)
    assert result['audit']['content']['status'] == result['audit']['style']['status'] == 'PASS'
    before, after = inspect_docx(source), inspect_docx(output)
    assert before['objects'] == after['objects']
    assert before['inventory'] == after['inventory']
    assert b'lineRule="atLeast"' in __import__('zipfile').ZipFile(output).read('word/document.xml')

