import zipfile

import pytest
from docx import Document
from lxml import etree

from manuweave.completeness import apply_completion
from manuweave.fields import _update
from manuweave.word import NS, audit_docx


def test_toc_cache_uses_only_original_headings(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("Synthetic paper", "Title")
    doc.add_paragraph("Methods", "Heading 1")
    doc.add_paragraph("The original value is 42.")
    doc.save(source)
    completed = tmp_path / "completed.docx"
    apply_completion(source, completed, {"add_toc": True})
    with pytest.raises(ValueError, match="not a source heading"):
        _update(completed, tmp_path / "bad.docx", ["Invented conclusion\t1"])
    assert not (tmp_path / "bad.docx").exists()
    _update(completed, tmp_path / "updated.docx", ["Methods\t2\r"])
    assert audit_docx(completed, tmp_path / "updated.docx")["content"]["status"] == "PASS"
    with zipfile.ZipFile(tmp_path / "updated.docx") as archive:
        tree = etree.fromstring(archive.read("word/document.xml"))
    toc = tree.find(".//w:fldSimple", NS)
    assert toc.get("{" + NS["w"] + "}dirty") == "false"
    assert toc.xpath(".//w:t/text()", namespaces=NS) == ["Methods", "2"]
