from pathlib import Path

import pytest
from docx import Document

from manuweave.presets import load_preset
from manuweave.templates import guard_engine, validate_template
from manuweave.word import apply_docx, inspect_docx

FAMILY_ROOT = Path(__file__).resolve().parents[1] / "src" / "manuweave" / "data" / "templates" / "computer-wiley"


@pytest.mark.parametrize("identifier", [path.stem for path in sorted((FAMILY_ROOT / "profiles").glob("*.json"))])
def test_wiley_confirmed_word_rules_apply_without_content_changes(identifier, tmp_path):
    profile = load_preset(identifier)
    validate_template(profile)
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("Synthetic Wiley formatting example", "Title")
    document.add_paragraph("1 Methods", "Heading 1")
    document.add_paragraph("The synthetic source value is 12.5 and must remain unchanged.")
    document.save(source)
    output = tmp_path / "formatted.docx"
    result = apply_docx(source, output, profile)
    assert result["audit"]["content"]["status"] == "PASS"
    assert result["audit"]["style"]["status"] == "PASS"
    if identifier not in {"computer-wiley-spe-manuscript", "computer-wiley-stvr-manuscript", "computer-wiley-jsep-manuscript", "computer-wiley-cavw-manuscript"}:
        assert result["audit"]["style"]["exact_source_sizes"] is False
        assert result["audit"]["style"]["representation_adjustments"]
    parsed = inspect_docx(output)
    body = next(block for block in parsed["blocks"] if block["role"] == "body")
    for prop, expected in profile["rules"]["body"].items():
        if isinstance(expected, (int, float)):
            represented = round(expected * 2) / 2 if prop == "size_pt" else expected
            assert body["format"][prop] == pytest.approx(represented, abs=0.03)
        else:
            assert body["format"][prop] == ("both" if prop == "alignment" and expected == "justify" else expected)
    if "columns" in profile["page"]:
        assert parsed["page"][0]["columns"] == profile["page"]["columns"]
    if identifier == "computer-wiley-jasist-njd-2025":
        assert profile["rules"]["body"]["size_pt"] == 9
        assert parsed["page"][0]["width_mm"] == pytest.approx(210, abs=0.03)
        assert parsed["page"][0]["height_mm"] == pytest.approx(276, abs=0.03)
    with pytest.raises(ValueError, match="no verified latex implementation"):
        guard_engine(profile, "latex")


def test_wiley_sources_do_not_bundle_unlicensed_publisher_classes():
    import json

    from manuweave import templates

    root = templates.ROOT / "computer-wiley"
    facts = json.loads((root / "source-facts.json").read_text(encoding="utf-8"))
    assert len(facts["coverage_gaps"]) == 7
    assert not list(root.rglob("*.cls"))
    assert all(source["url"].startswith("https://") and len(source["sha256"]) == 64 for source in facts["sources"])
    assert all(path.suffix == ".json" for path in root.rglob("*") if path.is_file())


def test_detailed_source_values_preserve_units_and_effective_definitions():
    jasist = load_preset("computer-wiley-jasist-njd-2025")
    extracted = jasist["extracted_format"]
    assert extracted["roles"]["body"]["size_pt"]["value"] == 9
    assert extracted["roles"]["body"]["size_pt"]["unit"] == "bp"
    assert extracted["roles"]["heading1"]["size_pt"]["value"] == 10
    assert extracted["roles"]["heading1"]["size_pt"]["unit"] == "tex_pt"
    assert "1671" in extracted["roles"]["heading1"]["size_pt"]["source"]["locator"]
    assert extracted["roles"]["title"]["size_pt"]["physical_word_pt"] == pytest.approx(18 * 72 / 72.27)
    assert jasist["rules"]["title"]["size_pt"] == pytest.approx(18 * 72 / 72.27)
    assert jasist["rules"]["title"]["size_pt"] != 18
    assert "line_spacing_pt" not in jasist["rules"]["body"]
    assert extracted["roles"]["body"]["load_time_baseline_pt"]["value"] == 13
    assert extracted["page"]["column_gap_mm"]["value"] == 6
    cgf = load_preset("computer-wiley-cgf-regular-submission")
    assert cgf["extracted_format"]["roles"]["body"]["size_pt"]["unit"] == "tex_pt"
    assert cgf["extracted_format"]["roles"]["body"]["size_pt"]["physical_word_pt"] == pytest.approx(9 * 72 / 72.27)
    assert cgf["extracted_format"]["roles"]["caption"]["label_bold"]["value"] is True
    assert cgf["extracted_format"]["roles"]["caption"]["text_italic"]["value"] is True


def test_all_shared_font_sizes_keep_exact_physical_values_and_catalog_bindings():
    import json

    catalog = json.loads((FAMILY_ROOT.parent / "computer-catalog.json").read_text(encoding="utf-8"))
    titles = {journal["id"]: journal["name"] for journal in catalog["journals"]}
    linked = set()
    fractional_sizes = 0
    for path in (FAMILY_ROOT / "profiles").glob("*.json"):
        profile = load_preset(path)
        assert len(profile["catalog_ids"]) == 1
        assert titles[profile["catalog_ids"][0]] == profile["journal_binding"]["name"]
        linked.update(profile["catalog_ids"])
        for role, rule in profile["rules"].items():
            if "size_pt" not in rule:
                continue
            source = profile["extracted_format"]["roles"][role]["size_pt"]
            expected = source["value"] * 72 / 72.27 if source["unit"] == "tex_pt" else source["value"]
            assert rule["size_pt"] == expected
            corresponding = [requirement for requirement in profile["template"]["requirements"]
                             if requirement["target"] == role and requirement["property"] == "size_pt"]
            assert len(corresponding) == 1
            assert corresponding[0]["expected"] == expected
            fractional_sizes += expected != round(expected * 2) / 2
    assert len(linked) == 13
    assert fractional_sizes > 100


def test_source_guard_rejects_a_prequantized_shared_title_size():
    import copy

    profile = copy.deepcopy(load_preset("computer-wiley-jasist-njd-2025"))
    assert profile["rules"]["title"]["size_pt"] == 18 * 72 / 72.27
    profile["rules"]["title"]["size_pt"] = 18
    with pytest.raises(ValueError, match=r"Unconfirmed journal formatting rule: title.size_pt"):
        validate_template(profile)


def test_optional_recommended_design_is_not_a_mandatory_journal_layout():
    profile = load_preset("computer-wiley-iet-cvi-njd-choice")
    assert profile["journal_binding"]["strength"] == "recommended_optional_choice"
    assert "optional" in profile["scope"]
    assert profile["extracted_format"]["citations"]["journal_reference_family"]["value"] == "Chicago"
    assert not profile["latex"]
    assert not profile["template"]["assets"]
    assert "remote_archive" not in profile["template"]


@pytest.mark.parametrize("identifier", ["computer-wiley-jasist-njd-2025", "computer-wiley-cgf-regular-submission"])
def test_all_declared_roles_apply_stably_with_explicit_host_role_map(identifier, tmp_path):
    import zipfile

    profile = load_preset(identifier)
    role_names = list(profile["rules"])
    source = tmp_path / "all-roles.docx"
    document = Document()
    for role in role_names:
        document.add_paragraph(f"Synthetic {role} block retains 12.5 and its exact supplied words.")
    document.save(source)
    blocks = inspect_docx(source)["blocks"]
    assert len(blocks) == len(role_names)
    role_map = {block["id"]: role for block, role in zip(blocks, role_names)}
    first, second = tmp_path / "first.docx", tmp_path / "second.docx"
    result = apply_docx(source, first, profile, roles=role_map)
    assert result["audit"]["content"]["status"] == result["audit"]["style"]["status"] == "PASS"
    repeated = apply_docx(first, second, profile, roles=role_map)
    assert repeated["audit"]["content"]["status"] == repeated["audit"]["style"]["status"] == "PASS"
    with zipfile.ZipFile(first) as a, zipfile.ZipFile(second) as b:
        assert a.read("word/document.xml") == b.read("word/document.xml")
