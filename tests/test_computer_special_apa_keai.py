"""Journal linkage, bounded author choices, and real editable Word application."""
import json
import runpy
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement

from manuweave.models import StyleProfile
from manuweave.presets import load_preset
from manuweave.templates import validate_template
from manuweave.word import apply_docx, inspect_docx

ROOT = Path(__file__).resolve().parents[1] / 'src/manuweave/data/templates/computer-special'
FACTS = json.loads((ROOT / 'apa-keai-source-facts.json').read_text(encoding='utf-8'))['verified_layouts']
NAMES = [item['id'] for item in FACTS]


@pytest.mark.parametrize('identifier', NAMES)
def test_executable_values_have_specific_source_locators(identifier):
    profile = load_preset(identifier)
    StyleProfile.model_validate(profile, strict=True)
    validate_template(profile)
    assert profile['catalog_journal_id'].startswith('ccf-2026-')
    assert profile['journal_title'] in profile['name']
    assert profile['template']['assets'] == []
    assert profile['latex'] == {}
    assert 'fixture' not in profile['template']
    for target, values in [*profile['rules'].items(), ('page', profile['page'])]:
        group = profile['extracted_format']['page'] if target == 'page' else profile['extracted_format']['roles'][target]
        for prop, expected in values.items():
            assert group[prop]['value'] == expected
            assert group[prop]['source']['locator']
            assert group[prop]['unit']
    assert all(s['verified'] for s in profile['sources'])


def test_igi_final_manuscript_apa_referral_keeps_permitted_font_choice_explicit():
    for identifier in [n for n in NAMES if n.startswith('comp-igi-')]:
        profile = load_preset(identifier)
        assert profile['stage'] == 'final_author_manuscript'
        assert 'titleid=' in profile['sources'][0]['url']
        assert profile['rules']['body']['size_pt'] == 12
        assert profile['rules']['body']['line_spacing'] == 2
        assert profile['rules']['body']['first_line_pt'] == 36
        assert profile['rules']['heading3']['italic'] is True
        assert profile['rules']['reference']['first_line_pt'] == -36
        assert set(profile['page'].values()) == {25.4}
        detailed = profile['extracted_format']
        assert 'Author-selected' in detailed['roles']['body']['font']['condition']
        assert detailed['roles']['heading4']['run_in']['value'] is True
        assert len(detailed['font_options']['value']) == 7
        assert not any(s.get('sha256', '').startswith('d0203228') for s in profile['sources'])


def test_jslhr_does_not_turn_reference_apa_referral_into_body_geometry():
    profile = load_preset('comp-jslhr-word')
    assert profile['page'] == {}
    assert profile['rules']['body'] == {'line_spacing_rule': 'auto', 'line_spacing': 2}
    assert profile['rules']['reference']['first_line_pt'] == -36
    assert profile['extracted_format']['captions']['minimum_figure_dpi']['value'] == 300
    assert all('font' not in rule and 'size_pt' not in rule for rule in profile['rules'].values())


def test_vrih_preserves_source_font_variants_and_main_body_section():
    profile = load_preset('comp-vrih-word')
    assert profile['rules']['author']['font'] == 'Arial'
    assert profile['rules']['correspondence']['size_pt'] == 9
    assert profile['rules']['heading1']['bold'] is True
    assert profile['rules']['body']['size_pt'] == 10.5
    assert profile['rules']['title']['size_pt'] == 18
    assert profile['rules']['affiliation']['size_pt'] == 8
    assert profile['rules']['affiliation']['italic'] is True
    assert profile['rules']['body_start']['first_line_pt'] == 0
    assert profile['rules']['body']['first_line_pt'] == 10.1
    assert profile['page']['columns'] == 1
    assert profile['page']['margin_top_mm'] == pytest.approx(1008 * 25.4 / 1440)
    assert profile['page']['column_gap_mm'] == pytest.approx(288 * 25.4 / 1440)
    assert 'font' not in profile['rules'].get('equation', {})


def test_jasa_uses_actual_author_guideline_dimensions_and_heading_hierarchy():
    profile = load_preset('comp-jasa-word')
    assert profile['page']['width_mm'] == 215.9
    assert profile['page']['height_mm'] == 279.4
    assert profile['page']['margin_left_mm'] == 25.4
    assert profile['rules']['body']['size_pt'] == 12
    assert profile['rules']['heading3']['bold'] is profile['rules']['heading3']['italic'] is True
    assert profile['extracted_format']['equations']['number_alignment']['value'] == 'right'
    assert '2023-06-01' in profile['extracted_format']['references']['style']['value']


@pytest.mark.parametrize('identifier', NAMES)
def test_applies_every_selected_role_and_preserves_native_objects(tmp_path, identifier):
    profile = load_preset(identifier)
    doc = Document()
    roles = {}
    for index, role in enumerate(profile['rules']):
        doc.add_paragraph('Synthetic ' + role + ' with constant 0.125')
        roles['p' + str(index)] = role
    math = OxmlElement('m:oMath')
    run = OxmlElement('m:r')
    value = OxmlElement('m:t')
    value.text = 'E=mc²'
    run.append(value)
    math.append(run)
    doc.add_paragraph()._p.append(math)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = 'Synthetic merged header'
    table.cell(1, 0).text = 'A'
    table.cell(1, 1).text = '12.5'
    source, output, again = (tmp_path / name for name in ['source.docx', 'output.docx', 'again.docx'])
    doc.save(source)
    before = inspect_docx(source)
    result = apply_docx(source, output, profile, roles)
    assert result['audit']['content']['status'] == 'PASS'
    assert result['audit']['style']['status'] == 'PASS'
    after = inspect_docx(output)
    assert before['inventory'] == after['inventory']
    assert before['objects'] == after['objects']
    second = apply_docx(output, again, profile, roles)
    assert second['audit']['style']['status'] == 'PASS'
    assert inspect_docx(again)['inventory'] == after['inventory']


def test_builder_regenerates_checked_in_format_facts_without_original_assets():
    expected = runpy.run_path(str(ROOT / 'builder_apa_keai.py'))['profiles']()
    assert {p['id'] for p in expected} == set(NAMES)
    for profile in expected:
        actual = json.loads((ROOT / 'profiles' / (profile['id'] + '.json')).read_text(encoding='utf-8'))
        assert profile == actual
