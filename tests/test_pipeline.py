import copy
import json
from pathlib import Path

import pytest
from docx import Document

from manuweave import pipeline
from manuweave.cli import main
from manuweave.io import read_json
from manuweave.models import EditPlan


@pytest.fixture
def manuscript(tmp_path):
    path = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("A synthetic research paper", "Title")
    doc.add_paragraph("1 Methods", "Heading 1")
    doc.add_paragraph("The synthetic dataset contains 42 observations.")
    doc.save(path)
    return path


def test_plan_bound_to_input_and_whitelist(manuscript, tmp_path):
    profile = pipeline.resolve_profile("english-single")
    edit_plan = pipeline.plan(manuscript, profile)
    malicious = copy.deepcopy(edit_plan)
    malicious["operations"].append({"op": "execute_python", "target": "document", "code": "pass"})
    with pytest.raises(ValueError):
        EditPlan.model_validate(malicious)
    malicious = copy.deepcopy(edit_plan)
    malicious["operations"].insert(0, {"op": "assign_role", "target": "missing", "role": "body"})
    with pytest.raises(ValueError, match="unknown block"):
        pipeline.apply(malicious, tmp_path / "bad.docx")
    doc = Document(manuscript)
    doc.add_paragraph("Changed after planning")
    doc.save(manuscript)
    with pytest.raises(ValueError, match="Source changed"):
        pipeline.apply(edit_plan, tmp_path / "stale.docx")


def test_run_and_resume_never_claim_visual_pass(manuscript, tmp_path, monkeypatch):
    # Isolate the visual/resume gate from fonts installed on the test host.
    monkeypatch.setattr("manuweave.fonts.check_fonts", lambda _: {"status": "PASS", "scope": "Synthetic font inventory"})
    monkeypatch.setattr("manuweave.render.render", lambda *_: {"status": "UNVERIFIED", "reason": "No renderer"})
    result = pipeline.run(manuscript, tmp_path / "run", pipeline.resolve_profile("english-single"))
    assert result["status"] == "UNVERIFIED"
    audit = read_json(result["report"])
    assert audit["content"]["status"] == "PASS"
    assert audit["style"]["status"] == "PASS"
    assert pipeline.run(state_path=result["state"]) == result
    output = Path(result["output"])
    output.write_bytes(b"tampered")
    with pytest.raises(ValueError, match="Native output changed"):
        pipeline.run(state_path=result["state"])


def test_resume_after_renderer_exception(manuscript, tmp_path, monkeypatch):
    def fail(*_):
        raise RuntimeError("renderer unavailable")
    monkeypatch.setattr("manuweave.render.render", fail)
    output = tmp_path / "run"
    with pytest.raises(RuntimeError):
        pipeline.run(manuscript, output, pipeline.resolve_profile("english-single"))
    state = read_json(output / "state.json")
    assert state["stage"] == "RENDER"
    monkeypatch.setattr("manuweave.render.render", lambda *_: {"status": "UNVERIFIED"})
    result = pipeline.run(state_path=output / "state.json")
    assert Path(result["output"]).exists()


def test_cli_schema_and_json_errors(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr("sys.argv", ["manuweave", "schema", "--output", str(tmp_path / "schemas")])
    assert main() == 0
    assert json.loads(capsys.readouterr().out)["output"]
    assert (tmp_path / "schemas" / "EditPlan.schema.json").exists()
    monkeypatch.setattr("sys.argv", ["manuweave", "inspect", str(tmp_path / "absent.docx")])
    assert main() == 2
    assert json.loads(capsys.readouterr().out)["status"] == "FAIL"


def test_latex_plan_hash_covers_subfiles(tmp_path):
    source = tmp_path / "latex"
    source.mkdir()
    main = source / "main.tex"
    main.write_text(r"\documentclass{article}\begin{document}\input{body}\end{document}")
    (source / "body.tex").write_text("The result is 12.")
    edit_plan = pipeline.plan(main, pipeline.resolve_profile("english-single"))
    (source / "body.tex").write_text("The result is 24.")
    with pytest.raises(ValueError, match="Source changed"):
        pipeline.apply(edit_plan, tmp_path / "changed")


def test_resume_adopts_only_revalidated_interrupted_candidate(manuscript, tmp_path, monkeypatch):
    # Recovery semantics must not depend on proprietary host font availability.
    monkeypatch.setattr("manuweave.fonts.check_fonts", lambda _: {"status": "PASS", "scope": "Synthetic font inventory"})
    original_artifact = pipeline._artifact

    def interrupted(state, name, data):
        if name == "apply.json":
            raise RuntimeError("simulated power loss before apply journal commit")
        return original_artifact(state, name, data)

    monkeypatch.setattr(pipeline, "_artifact", interrupted)
    output = tmp_path / "interrupted"
    with pytest.raises(RuntimeError):
        pipeline.run(manuscript, output, pipeline.resolve_profile("english-single"))
    assert (output / "editable.docx").exists()
    monkeypatch.setattr(pipeline, "_artifact", original_artifact)
    monkeypatch.setattr("manuweave.render.render", lambda *_: {"status": "UNVERIFIED"})
    result = pipeline.run(state_path=output / "state.json")
    assert result["status"] == "UNVERIFIED"
    assert read_json(output / "apply.json")["changes"][0]["operation"] == "recovered_verified_candidate"
