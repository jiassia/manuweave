import zipfile

import pytest
from docx import Document
from PIL import Image

from manuweave.latex import convert


@pytest.fixture
def pandoc_available():
    import shutil
    if shutil.which("pandoc"):
        return
    try:
        import pypandoc
        pypandoc.get_pandoc_path()
    except (ImportError, OSError):
        pytest.skip("Optional Pandoc adapter is not installed")


def test_docx_to_latex_extracts_media(tmp_path, pandoc_available):
    image = tmp_path / "source.png"
    Image.new("RGB", (40, 30), "blue").save(image)
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("The protected synthetic value is 42.")
    doc.add_picture(str(image))
    doc.save(source)
    result = convert(source, tmp_path / "converted.tex")
    assert result["status"] == "PASS"
    assert result["fidelity"] == "UNVERIFIED"
    assert "42" in (tmp_path / "converted.tex").read_text(encoding="utf-8")
    assert list((tmp_path / "converted-media").rglob("*.png"))


def test_multifile_latex_to_docx_embeds_media(tmp_path, pandoc_available):
    project = tmp_path / "latex"
    project.mkdir()
    (project / "main.tex").write_text(r"\documentclass{article}\usepackage{graphicx}\begin{document}\input{body}\end{document}")
    (project / "body.tex").write_text(r"Synthetic value 42. \includegraphics{plot.png}")
    Image.new("RGB", (40, 30), "blue").save(project / "plot.png")
    result = convert(project / "main.tex", tmp_path / "converted.docx")
    assert result["status"] == "PASS"
    assert "42" in " ".join(p.text for p in Document(tmp_path / "converted.docx").paragraphs)
    with zipfile.ZipFile(tmp_path / "converted.docx") as archive:
        assert any(name.startswith("word/media/") for name in archive.namelist())


def test_conversion_blocks_unsafe_input(tmp_path, pandoc_available):
    (tmp_path / "main.tex").write_text(r"\documentclass{article}\begin{document}\input{../private}\end{document}")
    with pytest.raises(ValueError):
        convert(tmp_path / "main.tex", tmp_path / "unsafe.docx")
