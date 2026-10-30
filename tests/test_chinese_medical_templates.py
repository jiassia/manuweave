"""Confirmed medical manuscript instructions and native Word preservation."""
from __future__ import annotations

import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches
from PIL import Image

from manuweave.models import StyleProfile
from manuweave.presets import load_preset
from manuweave.templates import audit_requirements, validate_template
from manuweave.word import apply_docx, audit_docx, inspect_docx

DATA = Path(__file__).resolve().parents[1] / "src/manuweave/data/templates/chinese-medical"
PROFILES = sorted((DATA / "profiles").glob("*.json"))
FACTS = json.loads((DATA / "source-facts.json").read_text(encoding="utf-8"))["documents"]


def _fixture(directory):
    image = directory / "synthetic.png"
    Image.new("RGB", (120, 40), "white").save(image)
    doc = Document()
    doc.add_paragraph("合成文档格式验证", "Title")
    doc.add_paragraph("摘要：本段仅为软件测试材料，无医学结论。")
    doc.add_paragraph("关键词：文档格式；合成材料；软件验证")
    doc.add_paragraph("1 资料与方法", "Heading 1")
    doc.add_paragraph("受保护的合成数值为 12。测试参考标记为 [1]。")
    doc.add_paragraph("1.1 处理流程", "Heading 2")
    doc.add_paragraph("第二个合成数值为 24。")
    para = doc.add_paragraph()
    math = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.text = "E=mc²"
    run.append(text)
    math.append(run)
    para._p.append(math)
    doc.add_paragraph("表1 合成数值")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "合成分组"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "12"
    table.cell(2, 0).text = "B"
    table.cell(2, 1).text = "24"
    doc.add_picture(str(image), width=Inches(2))
    doc.add_paragraph("图1 合成示意图")
    doc.add_paragraph("参考文献", "Heading 1")
    doc.add_paragraph("[1] 合成作者. 2026. 本条为虚构的软件测试材料。")
    source = directory / "synthetic.docx"
    doc.save(source)
    return source


@pytest.mark.parametrize("profile_file", PROFILES, ids=lambda p: p.stem)
def test_each_medical_profile_has_explicit_currently_linked_source_facts(profile_file):
    profile = json.loads(profile_file.read_text(encoding="utf-8"))
    StyleProfile.model_validate(profile, strict=True)
    validate_template(profile)
    fact = FACTS[profile["id"]]
    assert load_preset(profile["id"])["id"] == profile["id"]
    assert profile["template"]["source_version"] == fact["source_version"]
    assert profile["template"]["engine_support"] == {"docx": "partial", "latex": "unsupported"}
    assert profile["template"]["license"] == "rules_only_source_not_redistributed"
    assert profile["template"]["assets"] == []
    assert all(len(s["sha256"]) == 64 and s["byte_length"] > 1000 for s in profile["sources"])
    assert all(s["retrieved_at"] == "2026-10-30" and s["source_verified"] for s in profile["sources"])
    assert [s["sha256"] for s in profile["sources"]] == [s["sha256"] for s in fact["sources"]]
    requirements = profile["template"]["requirements"]
    assert len({r["id"] for r in requirements}) == len(requirements)
    assert all(r["certainty"] == "verified" for r in requirements)
    assert all(r["source"]["locator"] and r["source"]["url"] in {s["url"] for s in fact["sources"]} for r in requirements)
    for role, rule in profile["rules"].items():
        for prop, value in rule.items():
            assert any(r["target"] == role and r["property"] == prop and r["expected"] == value
                       and r["checking"] == "xml" and r["applies_to"] == ["docx"] for r in requirements)
    for prop, value in profile["page"].items():
        assert any(r["target"] == "page" and r["property"] == prop and r["expected"] == value
                   and r["checking"] == "xml" for r in requirements)


@pytest.mark.parametrize("profile_file", PROFILES, ids=lambda p: p.stem)
def test_each_medical_profile_executes_and_preserves_native_document(profile_file, tmp_path):
    profile = json.loads(profile_file.read_text(encoding="utf-8"))
    source = _fixture(tmp_path)
    output = tmp_path / "formatted.docx"
    result = apply_docx(source, output, profile)
    assert result["audit"]["content"]["status"] == "PASS"
    assert result["audit"]["style"]["status"] == "PASS"
    before, after = inspect_docx(source), inspect_docx(output)
    assert before["inventory"] == after["inventory"]
    assert [b["text"] for b in before["blocks"]] == [b["text"] for b in after["blocks"]]
    assert after["objects"]["equations"] == after["objects"]["images"] == after["objects"]["tables"] == 1
    by_role = {b["role"]: b for b in after["blocks"]}
    for role, rule in profile["rules"].items():
        assert role in by_role
        for prop, expected in rule.items():
            if prop == "alignment" and expected == "justify":
                expected = "both"
            assert by_role[role]["format"].get(prop) == expected
    for prop, expected in profile["page"].items():
        assert abs(after["page"][0][prop] - expected) < .03
    repeat = tmp_path / "repeat.docx"
    apply_docx(output, repeat, profile)
    assert audit_docx(output, repeat, profile)["style"]["status"] == "PASS"
    assert inspect_docx(repeat)["inventory"] == after["inventory"]
    with ZipFile(source) as original, ZipFile(output) as formatted:
        assert all(original.read(n) == formatted.read(n) for n in original.namelist() if n.startswith("word/media/"))
        assert b"gridSpan" in formatted.read("word/document.xml")
    requirements = audit_requirements(output, profile)
    assert requirements["status"] == "UNVERIFIED"
    assert any(r["status"] == "UNVERIFIED" for r in requirements["checks"])


def test_explicit_point_values_and_manuscript_scopes_do_not_gain_defaults():
    cgp = load_preset("cn-cgp-research-word-2023")
    hospital = load_preset("cn-cjhp-research-word-2026")
    assert len(PROFILES) == 2
    assert cgp["rules"]["title"]["size_pt"] == 14
    assert cgp["rules"]["body"]["size_pt"] == 10.5
    assert cgp["rules"]["body"]["line_spacing_pt"] == 15
    assert cgp["page"] == dict.fromkeys(["margin_top_mm", "margin_bottom_mm", "margin_left_mm", "margin_right_mm"], 12.7)
    assert hospital["rules"] == {"body": {"font": "Times New Roman", "east_asia_font": "宋体", "size_pt": 10.5,
                                           "line_spacing_rule": "auto", "line_spacing": 1.5}}
    assert hospital["page"] == {}
    image_text = next(r for r in cgp["template"]["requirements"] if r["id"] == "figure-internal-text")
    assert image_text["expected"]["sizes_pt"] == [6, 8]
    assert image_text["checking"] == "host"
    assert "figure_caption" not in cgp["rules"]
    assert not any(p.suffix in {".doc", ".docx", ".pdf"} for p in DATA.rglob("*"))
