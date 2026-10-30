"""Per-template checks against measured official Word source facts and native edits."""
from __future__ import annotations

import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image

from manuweave.models import StyleProfile
from manuweave.templates import validate_template
from manuweave.word import apply_docx, audit_docx, inspect_docx

DATA = Path(__file__).resolve().parents[1] / "src/manuweave/data/templates/official-word"
PROFILES = sorted((DATA / "profiles").glob("*.json"))
FACTS = json.loads((DATA / "source-facts.json").read_text(encoding="utf-8"))["documents"]


def _sample(directory):
    source = directory / "synthetic-source.docx"
    image = directory / "synthetic-plot.png"
    Image.new("RGB", (100, 40), "white").save(image)
    doc = Document()
    doc.add_paragraph("Synthetic template verification", "Title")
    doc.add_paragraph("Abstract: This is a synthetic software fixture with no research claims.")
    doc.add_paragraph("Keywords: native document; template verification")
    doc.add_paragraph("Introduction", "Heading 1")
    doc.add_paragraph("The protected synthetic value is 12. A fixture reference is [1].")
    doc.add_paragraph("Methods", "Heading 2")
    doc.add_paragraph("The second protected synthetic value is 24.")
    doc.add_paragraph("Implementation", "Heading 3")
    equation = doc.add_paragraph()
    math = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.text = "E=mc²"
    run.append(text)
    math.append(run)
    equation._p.append(math)
    doc.add_paragraph("Table 1: Synthetic values.")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Synthetic grouped values"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "12"
    table.cell(2, 0).text = "B"
    table.cell(2, 1).text = "24"
    doc.add_picture(str(image), width=Inches(2))
    doc.add_paragraph("Figure 1: Synthetic diagram.")
    doc.add_paragraph("References", "Heading 1")
    doc.add_paragraph("[1] Synthetic Fixture. 2026. Invented test material, not a real source.")
    doc.save(source)
    return source


@pytest.mark.parametrize("profile_file", PROFILES, ids=lambda p: p.stem)
def test_every_official_word_profile_has_pinned_rule_sources(profile_file):
    profile = json.loads(profile_file.read_text(encoding="utf-8"))
    StyleProfile.model_validate(profile, strict=True)
    validate_template(profile)
    fact = FACTS[profile["id"]]
    assert profile["sources"][0]["sha256"] == fact["sha256"]
    assert len(fact["sha256"]) == 64
    assert profile["template"]["source_version"] == fact["source_modified"]
    assert profile["template"]["license"] == "rules_only_source_not_redistributed"
    assert profile["template"]["assets"] == []
    assert profile["template"]["engine_support"] == {"docx": "partial", "latex": "unsupported"}
    requirements = profile["template"]["requirements"]
    assert len(requirements) > 80
    assert len({r["id"] for r in requirements}) == len(requirements)
    assert all(r["source"]["url"] == fact["url"] and r["source"]["locator"] for r in requirements)
    for role, properties in profile["rules"].items():
        for prop, value in properties.items():
            assert any(r["target"] == role and r["property"] == prop and r["expected"] == value
                       and r["certainty"] == "verified" and r["checking"] == "xml" for r in requirements)
    assert all(r["certainty"] != "unknown" for r in requirements)


@pytest.mark.parametrize("profile_file", PROFILES, ids=lambda p: p.stem)
def test_every_official_word_profile_executes_native_formatting(profile_file, tmp_path):
    profile = json.loads(profile_file.read_text(encoding="utf-8"))
    source = _sample(tmp_path)
    output = tmp_path / "formatted.docx"
    result = apply_docx(source, output, profile)
    assert result["audit"]["content"]["status"] == "PASS"
    assert result["audit"]["style"]["status"] == "PASS"
    ir = inspect_docx(output)
    assert ir["objects"]["equations"] == 1
    assert ir["objects"]["images"] == 1
    assert ir["objects"]["tables"] == 1
    assert any(b["role"] == "figure_caption" for b in ir["blocks"])
    assert any(b["role"] == "table_caption" for b in ir["blocks"])
    by_role = {b["role"]: b for b in ir["blocks"] if b["text"].strip()}
    assert by_role["body"]["format"]["size_pt"] == profile["rules"]["body"]["size_pt"]
    assert by_role["figure_caption"]["format"]["space_before_pt"] == profile["rules"]["figure_caption"]["space_before_pt"]
    assert by_role["table_caption"]["format"]["space_before_pt"] == profile["rules"]["table_caption"]["space_before_pt"]
    assert by_role["reference"]["format"]["line_spacing_rule"] == profile["rules"]["reference"]["line_spacing_rule"]
    assert by_role["reference"]["format"]["first_line_pt"] == -18
    assert by_role["reference"]["format"]["left_indent_pt"] == 18
    again = tmp_path / "formatted-again.docx"
    apply_docx(output, again, profile)
    assert audit_docx(output, again, profile)["style"]["status"] == "PASS"
    assert inspect_docx(output)["inventory"] == inspect_docx(again)["inventory"]
    with ZipFile(source) as original, ZipFile(output) as formatted:
        media = [n for n in original.namelist() if n.startswith("word/media/")]
        assert all(original.read(n) == formatted.read(n) for n in media)


def test_word_profiles_are_different_observed_templates_not_relabelled_defaults():
    ieee = json.loads((DATA / "profiles/ieee-word-a4-2019.json").read_text(encoding="utf-8"))
    letter = json.loads((DATA / "profiles/ieee-word-letter-2019.json").read_text(encoding="utf-8"))
    acm = json.loads((DATA / "profiles/acm-word-submission-2023.json").read_text(encoding="utf-8"))
    assert ieee["rules"]["body"]["font"] == "Times New Roman"
    assert ieee["rules"]["body"]["size_pt"] == 10
    assert ieee["rules"]["body"]["line_spacing"] == .95
    assert ieee["rules"]["title"]["size_pt"] == 24
    assert ieee["rules"]["reference"]["line_spacing_rule"] == "exact"
    assert ieee["rules"]["reference"]["line_spacing_pt"] == 9
    assert abs(ieee["page"]["width_mm"] - 210) < .02
    assert letter["page"]["width_mm"] == 215.9
    assert ieee["page"]["columns"] == letter["page"]["columns"] == 2
    assert acm["page"]["columns"] == 1
    assert acm["rules"]["body"]["font"] == "Linux Libertine O"
    assert acm["rules"]["body"]["size_pt"] == 9
    assert acm["rules"]["body"]["line_spacing_rule"] == "atLeast"
    assert acm["rules"]["body"]["line_spacing_pt"] == 13.5
    assert acm["rules"]["title"]["font"] == "Linux Biolinum O"
    assert acm["rules"]["abstract"]["size_pt"] == 8
    assert not list(DATA.rglob("*.docx")) and not list(DATA.rglob("*.doc"))


def test_conditional_anonymity_does_not_become_a_blanket_citation_ban():
    acm = json.loads((DATA / "profiles/acm-word-submission-2023.json").read_text(encoding="utf-8"))
    requirements = {r["id"]: r for r in acm["template"]["requirements"]}
    assert requirements["double-blind-identity"]["certainty"] == "conditional"
    assert requirements["citation-venue-selection"]["expected"] == ["numeric", "author_year"]
    assert not any(r["property"] == "citations_allowed" and r["expected"] is False for r in requirements.values())
    assert requirements["figure-alternative-description"]["expected"] is True
    assert not any(r["property"] == "maximum_description_characters" for r in requirements.values())


@pytest.mark.parametrize("identifier", ["ieee-word-a4-2019", "ieee-word-letter-2019"])
def test_ieee_preserves_equation_layout_without_captured_source_measurements(identifier, tmp_path):
    profile = json.loads((DATA / "profiles" / f"{identifier}.json").read_text(encoding="utf-8"))
    assert "equation" not in profile["rules"]
    assert "equation" not in FACTS[identifier]["styles"]
    assert not any(r["target"] == "equation" and r["checking"] == "xml" for r in profile["template"]["requirements"])
    assert not any(c["role"] == "equation" for c in profile["checks"])
    source = _sample(tmp_path)
    document = Document(source)
    equation = next(p for p in document.paragraphs if p._p.xpath(".//m:oMath"))
    equation.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    equation.paragraph_format.space_before = Pt(7)
    equation.paragraph_format.space_after = Pt(5)
    equation.paragraph_format.line_spacing = Pt(20)
    document.save(source)
    original = next(b for b in inspect_docx(source)["blocks"] if b["role"] == "equation")
    output = tmp_path / "formatted.docx"
    result = apply_docx(source, output, profile)
    formatted = next(b for b in inspect_docx(output)["blocks"] if b["role"] == "equation")
    assert formatted["format"] == original["format"]
    assert result["audit"]["content"]["status"] == "PASS"
    assert result["audit"]["objects_after"]["equations"] == 1


@pytest.mark.parametrize("identifier", ["ieee-word-a4-2019", "ieee-word-letter-2019"])
def test_ieee_host_typography_uses_only_captured_original_xml_values(identifier):
    profile = json.loads((DATA / "profiles" / f"{identifier}.json").read_text(encoding="utf-8"))
    requirements = {r["id"]: r for r in profile["template"]["requirements"]}
    for observation_key, rule_id in (("heading1_small_caps", "heading1-small-caps"),
                                     ("heading1_numbering", "heading-numbering")):
        observed = FACTS[identifier].get("xml_observations", {}).get(observation_key)
        if observed is None:
            assert rule_id not in requirements
        else:
            assert observed["source_sha256"] == FACTS[identifier]["sha256"]
            assert observed["xml_namespace"] == "http://purl.oclc.org/ooxml/wordprocessingml/main"
            assert requirements[rule_id]["expected"] == observed["value"]
            assert requirements[rule_id]["source"]["locator"] == observed["locator"]


def test_builder_does_not_restore_host_claims_without_original_xml_observations(tmp_path, monkeypatch):
    import importlib.util
    spec = importlib.util.spec_from_file_location("official_word_builder", DATA / "build_profiles.py")
    builder = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(builder)
    facts = json.loads((DATA / "source-facts.json").read_text(encoding="utf-8"))
    for fact in facts["documents"].values():
        fact.pop("xml_observations", None)
    (tmp_path / "source-facts.json").write_text(json.dumps(facts), encoding="utf-8")
    (tmp_path / "profiles").mkdir()
    monkeypatch.setattr(builder, "HERE", tmp_path)
    builder.build()
    for path in (tmp_path / "profiles").glob("ieee*.json"):
        generated = json.loads(path.read_text(encoding="utf-8"))
        assert not {"heading1-small-caps", "heading-numbering"} & {r["id"] for r in generated["template"]["requirements"]}
