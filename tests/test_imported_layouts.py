"""Only audited, stage-specific Word layout facts become executable profiles."""
from __future__ import annotations

import copy
import importlib.util
import json
from pathlib import Path
from zipfile import ZipFile

import pytest

from manuweave.matrix import _word_fixture
from manuweave.models import StyleProfile
from manuweave.presets import load_preset, validate_profile
from manuweave.templates import guard_engine
from manuweave.word import apply_docx, inspect_docx

DATA = Path(__file__).resolve().parents[1] / "src/manuweave/data/templates/imported-journals"
FACTS = json.loads((DATA / "source-facts.json").read_text(encoding="utf-8"))
IDS = ("cn-cims-initial-word", "cn-cims-revised-word", "int-ncomms-initial-word")


def profile(identifier):
    return json.loads((DATA / "profiles" / (identifier + ".json")).read_text(encoding="utf-8"))


@pytest.mark.parametrize("identifier", IDS)
def test_imported_profile_has_verified_evidence_for_every_executable_property(identifier):
    item = profile(identifier)
    StyleProfile.model_validate(item, strict=True)
    validate_profile(item)
    guard_engine(item, "docx")
    assert load_preset(identifier)["id"] == identifier
    assert item["template"]["engine_support"] == {"docx": "partial", "latex": "unsupported"}
    assert item["template"]["assets"] == []
    assert item["manuscript_format"] == "Microsoft Word"
    assert item["stage"] in {"initial_submission", "revised_manuscript"}
    sources = {source["url"]: source for source in item["sources"]}
    requirements = item["template"]["requirements"]
    for target, properties in [("page", item["page"]), *item["rules"].items()]:
        for prop, expected in properties.items():
            matches = [rule for rule in requirements if rule["target"] == target and rule["property"] == prop and rule["expected"] == expected]
            assert len(matches) == 1
            rule = matches[0]
            assert rule["certainty"] == "verified" and rule["checking"] == "xml"
            assert rule["applies_to"] == ["docx"]
            source = sources[rule["source"]["url"]]
            assert source["source_verified"] is True
            assert source["original_file_redistributed"] is False
            assert source["submission_compliance_verified"] is False
            assert rule["source"]["response_sha256"] == source["sha256"]
            assert rule["source"]["locator"]
    assert "temporary_original" not in json.dumps(item)
    assert all(path.suffix not in {".doc", ".docx", ".html", ".pdf"} for path in DATA.rglob("*"))


def test_stage_specific_limits_and_nature_source_modal_strength_are_preserved():
    initial, revised, nature = [profile(identifier) for identifier in IDS]
    assert initial["rules"] == {} and initial["page"] == {"columns": 1}
    assert sum(len(properties) for properties in revised["rules"].values()) + len(revised["page"]) == 6
    assert revised["rules"] == {"body": {"font": "Times New Roman", "east_asia_font": "宋体", "size_pt": 10.5, "line_spacing_rule": "auto", "line_spacing": 1.5}}
    assert revised["page"] == {"columns": 1}
    assert nature["page"] == {"columns": 1}
    assert all(properties == {"line_spacing": 2.0} for properties in nature["rules"].values())
    assert all(rule["strength"] == "should" for rule in nature["template"]["requirements"])
    assert all(rule["conditions"]["condition"] == "microsoft_word_manuscript" for rule in nature["template"]["requirements"])
    assert len(next(entry for entry in FACTS["entries"] if entry["metadata"]["id"] == IDS[2])["format_facts"]) == 2
    source = nature["sources"][0]
    assert source["url"] == "https://www.nature.com/ncomms/submit/how-to-submit"
    assert source["http_status"] == 200 and source["byte_count"] == 171997
    assert source["sha256"] == "4bdf7208d86f6e2b6b4be7a3cc99faf35ffbbaa258d1248fc0d056aaa5a07adb"
    assert source["retrieved_at"] == "2026-10-30T11:46:22.372797+00:00"
    assert "?" not in source["url"]
    assert all(rule["source"]["locator"] == "/html/body/div[2]/div/div/div[2]/div[2]/p[26]" for rule in nature["template"]["requirements"])
    assert not any("anonym" in rule["property"].lower() or "author" in rule["target"].lower() for item in (initial, revised, nature) for rule in item["template"]["requirements"])


@pytest.mark.parametrize("identifier", IDS)
def test_each_imported_layout_applies_audits_and_repeats_without_content_loss(identifier, tmp_path):
    item = profile(identifier)
    source, roles, _ = _word_fixture(tmp_path / "input")
    original = inspect_docx(source)
    first, second = tmp_path / "first.docx", tmp_path / "second.docx"
    result = apply_docx(source, first, item, roles)
    assert result["audit"]["content"]["status"] == "PASS"
    assert result["audit"]["style"]["status"] == "PASS"
    after = inspect_docx(first)
    assert original["inventory"] == after["inventory"]
    assert [block["text"] for block in original["blocks"]] == [block["text"] for block in after["blocks"]]
    assert after["objects"]["equations"] == after["objects"]["tables"] == after["objects"]["images"] == 1
    assert all(section["columns"] == 1 for section in after["page"])
    for block in after["blocks"]:
        role = roles.get(block["id"], block["role"])
        for prop, expected in item["rules"].get(role, {}).items():
            actual = block["format"].get(prop)
            assert actual == expected or isinstance(actual, (int, float)) and isinstance(expected, (int, float)) and abs(actual - expected) < 0.06
    if identifier == IDS[0]:
        assert [block["format"] for block in original["blocks"]] == [block["format"] for block in after["blocks"]]
    again = apply_docx(first, second, item, roles)
    assert again["audit"]["content"]["status"] == again["audit"]["style"]["status"] == "PASS"
    with ZipFile(first) as a, ZipFile(second) as b:
        assert a.namelist() == b.namelist()
        assert all(a.read(name) == b.read(name) for name in a.namelist())


@pytest.mark.parametrize("identifier", IDS)
def test_imported_layout_does_not_remove_existing_author_identity(identifier, tmp_path):
    from docx import Document
    source, _, _ = _word_fixture(tmp_path / "input")
    document = Document(source)
    author = document.add_paragraph("Synthetic Fixture Author; Synthetic Test Institution; contact@example.invalid")
    document.paragraphs[0]._p.addnext(author._p)
    document.save(source)
    original = source.read_bytes()
    output = tmp_path / "formatted.docx"
    result = apply_docx(source, output, profile(identifier))
    assert result["audit"]["content"]["status"] == "PASS"
    assert source.read_bytes() == original
    assert any(block["text"] == author.text for block in inspect_docx(output)["blocks"])


def test_ncomms_unestablished_fixed_font_is_rejected_and_rebuild_is_deterministic():
    item = copy.deepcopy(profile(IDS[2]))
    item["rules"]["body"]["font"] = "Times New Roman"
    with pytest.raises(ValueError, match="Unconfirmed journal formatting rule"):
        validate_profile(item)
    spec = importlib.util.spec_from_file_location("mw_imported_layout_builder", DATA / "build_profiles.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    rebuilt = {item["id"]: item for item in module.build_profiles()}
    assert set(rebuilt) == set(IDS)
    assert all(rebuilt[identifier] == profile(identifier) for identifier in IDS)
