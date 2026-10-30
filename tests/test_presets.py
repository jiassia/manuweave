import json

import pytest

from manuweave.presets import AmbiguousPresetError, export_preset, import_preset, list_presets, load_preset, template_catalog, validate_profile


def test_all_profiles_have_real_rules_and_checks():
    entries = list_presets()
    assert entries
    assert len({p["id"] for p in entries}) == len(entries)
    for item in entries:
        profile = load_preset(item["id"])
        if "template" in profile:
            assert all(rule["certainty"] != "unknown" for rule in profile["template"]["requirements"])
            assert item["template_family"] == profile["template"]["family"]
            assert item["source_version"] == profile["template"]["source_version"]
            assert item["requirements_count"] == len(profile["template"]["requirements"])
        else:
            assert profile["compliance"] == "authored_style"
            assert profile["authored_style"]["author"] == "ManuWeave"
            assert "ManuWeave" in profile["name"]
            assert "custom" in profile["name"].lower() or "自定义" in profile["name"]
            assert profile["checks"] and profile["examples"]
            assert profile["sources"] == []
            assert all(check["property"] in profile["rules"][check["role"]] for check in profile["checks"])
            assert profile["page"]["width_mm"] >= 210
            assert profile["page"]["height_mm"] >= 279
            assert profile["tables"]["borders"]["top"]["style"] == "single"
            assert profile["tables"]["borders"]["insideV"]["style"] == "nil"


def test_paper_dimensions():
    for name in ("zh-undergraduate", "zh-master", "zh-phd"):
        profile = load_preset(name)
        assert (profile["page"]["width_mm"], profile["page"]["height_mm"]) == (210, 297)
    conference = load_preset("en-conference")
    assert (conference["page"]["width_mm"], conference["page"]["height_mm"]) == (215.9, 279.4)


def test_natural_language_and_ambiguity():
    assert load_preset("请使用自定义中文双栏排版")["page"]["columns"] == 2
    with pytest.raises(AmbiguousPresetError) as exc:
        load_preset("自定义中文")
    assert set(exc.value.candidates) == {"zh-journal-single", "zh-journal-double", "zh-undergraduate", "zh-master", "zh-phd"}


def test_unverified_brand_layouts_are_not_builtin():
    ids = {entry["id"] for entry in list_presets()}
    assert not ids.intersection({"acm-general", "apa7-general", "elsevier-general", "ieee-general", "springer-general"})
    with pytest.raises(AmbiguousPresetError) as choices:
        load_preset("APA 7")
    assert choices.value.candidates
    assert all(load_preset(identifier).get("template") for identifier in choices.value.candidates)
    for entry in list_presets():
        if entry["compliance"] == "authored_style":
            assert not any(brand in alias.lower() for alias in entry["aliases"] for brand in ("acm", "apa", "elsevier", "ieee", "springer"))


def test_roundtrip_and_existing_file_protection(tmp_path):
    original = load_preset("en-conference")
    original["id"] = "custom-conference"
    original["rules"]["body"]["size_pt"] = 11
    path = export_preset(original, tmp_path / "custom.json")
    imported = import_preset(path, tmp_path / "library")
    assert load_preset(imported) == original
    with pytest.raises(FileExistsError):
        export_preset(original, path)


@pytest.mark.parametrize("mutation", [
    lambda p: p.update(id="../../escape"),
    lambda p: p["rules"]["body"].update(execute="rm"),
    lambda p: p["rules"]["body"].update(size_pt=-1),
    lambda p: p["rules"]["body"].update(line_spacing=float("nan")),
    lambda p: p["page"].update(columns=99),
    lambda p: p["page"].update(columns=True),
    lambda p: p["page"].update(width_mm=20),
    lambda p: p["page"].update(height_mm=float("inf")),
    lambda p: p["tables"]["borders"]["top"].update(color="../../"),
    lambda p: p["tables"]["borders"]["top"].update(size_pt=100),
    lambda p: p["tables"]["borders"].update(diagonal={"style": "single", "size_pt": 1, "color": "auto"}),
    lambda p: p.update(running={"header": {"execute": "unknown"}}),
    lambda p: p.update(running={"unapproved": {"font": "Arial"}}),
    lambda p: p.update(running={"footer": {"size_pt": "12"}}),
])
def test_reject_invalid_profile(mutation):
    profile = load_preset("en-article-single")
    mutation(profile)
    with pytest.raises(ValueError):
        validate_profile(profile)


def test_portable_data_location(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    assert load_preset("en-article-single")["latex"]["documentclass"] == "article"


def test_corrupt_custom_profile(tmp_path):
    path = tmp_path / "bad.json"
    path.write_text(json.dumps({"id": "bad"}), encoding="utf-8")
    with pytest.raises(ValueError):
        load_preset(path)


def test_catalog_preserves_sources_and_explicit_unknown():
    entries = template_catalog()
    assert {entry["id"] for entry in entries} == {entry["id"] for entry in list_presets()}
    for entry in entries:
        assert entry["requirements_count"] == len(entry["requirements"])
        assert "sources" in entry
        assert entry["rules"] or entry["page"] or (
            entry["latex"].get("documentclass")
            and entry["engine_support"].get("latex") == "native"
            and entry["requirements"]
        )
        assert entry["coverage"]["evaluation"] == "not_executed"
        expected_unknown = [requirement for requirement in entry["requirements"] if requirement["certainty"] == "unknown"]
        assert entry["unknown"] == expected_unknown
    generic = next(entry for entry in entries if entry["id"] == "en-article-single")
    assert generic["compliance"] == "authored_style"
    assert generic["template_family"] is None
    assert generic["coverage"]["coverage_notes"]


def test_merge_nested_template_profiles_and_duplicate_ids(tmp_path, monkeypatch):
    import manuweave.presets as presets_module

    generic = load_preset("en-article-single")
    generic_dir = tmp_path / "presets"
    templates_dir = tmp_path / "templates"
    export_preset(generic, generic_dir / "generic.json")
    distinct = {**generic, "id": "nested-profile", "name": "Nested custom layout", "aliases": ["nested unique"]}
    nested = templates_dir / "family" / "profiles" / "nested.json"
    export_preset(distinct, nested)
    monkeypatch.setattr(presets_module, "DATA", generic_dir)
    monkeypatch.setattr(presets_module, "TEMPLATES", templates_dir)
    assert [profile["id"] for profile in list_presets()] == ["en-article-single", "nested-profile"]
    assert load_preset("use nested unique layout")["id"] == "nested-profile"
    with pytest.raises(AmbiguousPresetError) as error:
        load_preset("custom")
    assert error.value.candidates == sorted(error.value.candidates)
    nested.write_text(json.dumps(generic), encoding="utf-8")
    with pytest.raises(ValueError, match="Duplicate preset id"):
        list_presets()


def test_template_validation_uses_owning_contract(monkeypatch):
    import sys
    import types

    profile = load_preset("en-article-single")
    calls = []
    contract = types.ModuleType("manuweave.templates")
    contract.validate_template = lambda profile: calls.append(profile["id"])
    monkeypatch.setitem(sys.modules, "manuweave.templates", contract)
    profile["template"] = {"test_routing_only": True}
    validate_profile(profile)
    assert calls == ["en-article-single"]


def test_selected_source_requirements_are_rechecked_after_metadata_cache(tmp_path, monkeypatch):
    import manuweave.presets as presets_module
    profile = load_preset("ieee-word-a4-2019")
    profile.update(id="selected-json", aliases=["selected extracted layout"])
    root = tmp_path / "templates"
    directory = root / "test" / "profiles"
    directory.mkdir(parents=True)
    selected = directory / "selected.json"
    selected.write_text(json.dumps(profile), encoding="utf-8")
    unrelated = json.loads(json.dumps(profile))
    unrelated.update(id="unrelated-invalid", aliases=["unrelated invalid layout"])
    unrelated["rules"]["body"]["size_pt"] = 99
    (directory / "unrelated.json").write_text(json.dumps(unrelated), encoding="utf-8")
    monkeypatch.setattr(presets_module, "DATA", tmp_path / "generic")
    monkeypatch.setattr(presets_module, "TEMPLATES", root)
    assert load_preset("selected-json")["id"] == "selected-json"
    profile["rules"]["body"]["size_pt"] = 99
    selected.write_text(json.dumps(profile), encoding="utf-8")
    with pytest.raises(ValueError, match="Unconfirmed journal formatting rule"):
        load_preset("selected-json")
    with pytest.raises(ValueError, match="Unconfirmed journal formatting rule"):
        list_presets()


def test_extended_native_roles_and_empty_rules(tmp_path):
    profile = load_preset("en-article-single")
    profile["rules"] = {
        role: {} for role in (
            "heading1", "heading2", "heading3", "heading4", "heading5", "heading6",
            "abstract_heading", "references_heading", "table_body", "image", "legacy_image", "toc", "unknown",
        )
    }
    profile["running"] = {"header": {"font": "Arial", "size_pt": 9}, "footer": {}}
    path = export_preset(profile, tmp_path / "extended.json")
    assert load_preset(path) == profile
    profile["rules"] = {}
    assert validate_profile(profile)["rules"] == {}


def test_learned_profile_export_import_and_real_apply(tmp_path):
    from docx import Document

    from manuweave.word import apply_docx, learn_docx

    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("Synthetic preset integration", "Title")
    doc.add_paragraph("Abstract")
    doc.add_paragraph("This synthetic experiment contains 12 measurements.")
    doc.add_paragraph("Keywords: synthetic; integration")
    doc.add_paragraph("1 Introduction")
    doc.add_paragraph("The retained source value equals 12.5.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Variable"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "x"
    table.cell(1, 1).text = "12.5"
    doc.sections[0].header.paragraphs[0].text = "Synthetic running header"
    doc.save(source)
    reference = tmp_path / "reference.docx"
    reference_profile = load_preset("zh-master")
    reference_profile["running"] = {"header": {"font": "Arial", "size_pt": 9}}
    apply_docx(source, reference, reference_profile)
    learned = learn_docx(reference)
    assert "abstract_heading" in learned["rules"]
    assert learned["page"]["width_mm"] == pytest.approx(210, abs=0.1)
    assert learned["tables"]["header_bottom"]["style"] == "single"
    path = export_preset(learned, tmp_path / "learned.json")
    imported = import_preset(path, tmp_path / "library")
    reloaded = load_preset(imported)
    assert reloaded == learned
    result = apply_docx(source, tmp_path / "final.docx", reloaded)
    assert result["audit"]["content"]["status"] == "PASS"
    assert result["audit"]["style"]["status"] == "PASS"


def test_preset_lookup_observes_changed_aliases_and_new_duplicate_ids(tmp_path, monkeypatch):
    from manuweave import presets
    profile = load_preset("en-article-single")
    profile.update(id="fixture-format", name="Fixture format", aliases=["alpha alias one"])
    library = tmp_path / "library"
    library.mkdir()
    path = library / "fixture-format.json"
    path.write_text(json.dumps(profile), encoding="utf-8")
    monkeypatch.setattr(presets, "DATA", library)
    monkeypatch.setattr(presets, "TEMPLATES", tmp_path / "empty-templates")
    assert load_preset("alpha alias one")["id"] == "fixture-format"
    profile["aliases"] = ["beta alias different"]
    profile["rules"]["body"]["size_pt"] = 13
    path.write_text(json.dumps(profile), encoding="utf-8")
    assert load_preset("beta alias different")["rules"]["body"]["size_pt"] == 13
    with pytest.raises(ValueError, match="Unknown preset"):
        load_preset("alpha alias one")
    (library / "other-name.json").write_text(json.dumps(profile), encoding="utf-8")
    with pytest.raises(ValueError, match="Duplicate preset id"):
        load_preset("fixture-format")
