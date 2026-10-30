"""Source binding and actual Word preservation for the selected T&F/HCI layouts."""
import json
import runpy
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement

from manuweave.models import StyleProfile
from manuweave.presets import load_preset
from manuweave.templates import validate_template
from manuweave.word import apply_docx, inspect_docx

ROOT = Path(__file__).resolve().parents[1] / 'src/manuweave/data/templates/computer-special'
FACTS = json.loads((ROOT / 'tf-ws-source-facts.json').read_text(encoding='utf-8'))['verified_layouts']
NAMES = [item['id'] for item in FACTS]

def synthetic_fixture(destination):
    """Create original software-test content, never copy an author template's text."""
    document = Document()
    document.add_paragraph('Synthetic Authoring Layout Verification', 'Title')
    document.add_paragraph('Abstract: This software fixture checks native editable objects and typography.')
    document.add_paragraph('Keywords: synthetic; formatting; native objects')
    document.add_paragraph('1 Methods', 'Heading 1')
    document.add_paragraph('This synthetic paragraph contains the fixed test number 0.125. No scientific claim is made.')
    document.add_paragraph('1.1 Objects', 'Heading 2')
    document.add_paragraph('The original Word equation and merged table remain editable.')
    math = OxmlElement('m:oMath')
    run = OxmlElement('m:r')
    value = OxmlElement('m:t')
    value.text = 'E=mc²'
    run.append(value)
    math.append(run)
    document.add_paragraph()._p.append(math)
    document.add_paragraph('Table 1. Synthetic test values.')
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = 'Synthetic merged header'
    table.cell(1, 0).text = 'A'
    table.cell(1, 1).text = '12'
    table.cell(2, 0).text = 'B'
    table.cell(2, 1).text = '24'
    document.add_paragraph('Figure 1. Synthetic object caption.')
    document.add_paragraph('2 Results', 'Heading 1')
    document.add_paragraph('The fixture contains numbers 12, 24, and 0.125 for preservation checks.')
    document.add_paragraph('References')
    document.add_paragraph('[1] Synthetic software fixture reference, 2026.')
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)



@pytest.mark.parametrize('identifier', NAMES)
def test_each_layout_has_specific_journal_linkage_and_only_evidenced_operations(identifier):
    profile = load_preset(identifier)
    StyleProfile.model_validate(profile, strict=True)
    validate_template(profile)
    assert profile['catalog_journal_id'].startswith('ccf-2026-')
    assert profile['journal_title'] in profile['name']
    assert profile['template']['engine_support'] == {'docx': 'partial', 'latex': 'unsupported'}
    assert profile['template']['assets'] == []
    assert 'fixture' not in profile['template']
    assert profile['latex'] == {}
    assert profile['sources'][0]['url'].startswith('https://www.tandfonline.com/action/authorSubmission?')
    assert 'journalCode=' in profile['sources'][0]['url']
    assert profile['stage'].endswith('authoring_template')
    rules = {(r['target'], r['property']): r for r in profile['template']['requirements']}
    for target, properties in [*profile['rules'].items(), ('page', profile['page'])]:
        for prop, value in properties.items():
            requirement = rules[(target, prop)]
            assert requirement['expected'] == value
            assert requirement['certainty'] == 'verified'
            assert requirement['checking'] == 'xml'
            assert requirement['applies_to'] == ['docx']
            assert requirement['source']['url'] in {source['url'] for source in profile['sources']}
    assert not any(r['certainty'] == 'unknown' for r in rules.values())


def test_tf_word_source_is_an_actual_authoring_template_with_native_units():
    for identifier in ['comp-tf-tgis20-word', 'comp-tf-ucis20-word', 'comp-tf-ugit20-word']:
        profile = load_preset(identifier)
        assert profile['rules']['body']['size_pt'] == 12
        assert profile['rules']['body']['line_spacing'] == 2
        assert profile['rules']['title']['size_pt'] == 14
        assert profile['rules']['abstract']['size_pt'] == 11
        assert profile['rules']['heading2']['italic'] is True
        assert profile['rules']['reference']['first_line_pt'] == -36
        assert profile['page']['width_mm'] == pytest.approx(11901 * 25.4 / 1440, abs=.0001)
        assert profile['page']['margin_top_mm'] == pytest.approx(1418 * 25.4 / 1440, abs=.0001)
        assert profile['sources'][-1]['archive_member'] == 'TF_Template_Word_Windows_2016.dotx'
        assert profile['sources'][-1]['sha256'] != profile['sources'][-1]['member_sha256']
        assert profile['sources'][-1]['bytes'] == 25178


def test_hci_physical_tex_units_do_not_become_unconverted_word_points():
    profile = load_preset('comp-tf-hhci20-word')
    assert profile['rules']['body']['size_pt'] == pytest.approx(10 * 72 / 72.27, abs=1e-6)
    assert profile['page']['margin_top_mm'] == pytest.approx(70 * 25.4 / 72.27, abs=1e-6)
    assert profile['page']['margin_bottom_mm'] == pytest.approx(180 * 25.4 / 72.27, abs=1e-6)
    assert profile['page']['width_mm'] == 215.9
    assert profile['page']['height_mm'] == 279.4
    assert profile['page']['columns'] == 1
    assert not {'margin_left_mm', 'margin_right_mm'} & profile['page'].keys()
    assert 'font' not in profile['rules']['body']


@pytest.mark.parametrize('identifier', NAMES)
def test_native_application_preserves_objects_values_and_is_stable(tmp_path, identifier):
    source = tmp_path / 'synthetic.docx'
    synthetic_fixture(source)
    before = source.read_bytes()
    profile = load_preset(identifier)
    output = tmp_path / 'formatted.docx'
    result = apply_docx(source, output, profile)
    assert result['audit']['content']['status'] == 'PASS'
    assert result['audit']['style']['status'] == 'PASS'
    original = inspect_docx(source)
    after = inspect_docx(output)
    assert original['objects'] == after['objects']
    assert original['objects']['equations'] == original['objects']['tables'] == 1
    assert original['inventory'] == after['inventory']
    again = tmp_path / 'again.docx'
    second = apply_docx(output, again, profile)
    assert second['audit']['content']['status'] == second['audit']['style']['status'] == 'PASS'
    assert source.read_bytes() == before
    assert inspect_docx(again)['inventory'] == after['inventory']


def test_maintainer_builder_matches_every_checked_in_profile():
    builder = runpy.run_path(str(ROOT / 'builder_tf_ws.py'))
    expected = {item['id']: item for item in builder['profiles']()}
    assert set(expected) == set(NAMES)
    for identifier, profile in expected.items():
        actual = json.loads((ROOT / 'profiles' / (identifier + '.json')).read_text(encoding='utf-8'))
        assert actual == profile


def test_detailed_json_retains_author_equation_and_run_in_heading_values():
    for identifier in ['comp-tf-tgis20-word', 'comp-tf-ucis20-word', 'comp-tf-ugit20-word']:
        profile = load_preset(identifier)
        detailed = profile['extracted_format']
        assert detailed['roles']['author']['size_pt']['value'] == 14
        assert detailed['roles']['affiliation']['italic']['value'] is True
        assert detailed['roles']['correspondence']['size_pt']['value'] == 12
        assert detailed['equations']['alignment']['value'] == 'center'
        assert detailed['equations']['line_spacing']['value'] == 2
        assert detailed['roles']['heading4']['run_in']['value'] is True
        assert detailed['roles']['heading4']['heading_span_italic']['value'] is True
        assert detailed['roles']['heading4']['following_text_italic']['value'] is False
        assert detailed['roles']['body_subsequent_paragraph']['first_line_pt']['value'] == 36
        assert detailed['page']['header_distance_mm']['value'] == pytest.approx(709 * 25.4 / 1440, abs=1e-6)
        assert detailed['page']['footer_distance_mm']['value'] == detailed['page']['header_distance_mm']['value']
        assert detailed['page']['column_gap_mm']['value'] == pytest.approx(708 * 25.4 / 1440, abs=1e-6)
        for role in detailed['roles'].values():
            for item in role.values():
                assert item['value'] is not None and item['unit']
                assert item['source']['url'] and item['source']['locator']
    hci = load_preset('comp-tf-hhci20-word')['extracted_format']
    for heading in range(1, 6):
        assert hci['roles'][f'heading{heading}']['size_pt']['value'] > 0
        assert hci['roles'][f'heading{heading}']['size_pt']['unit'] == 'TeX pt'
    assert hci['roles']['title']['size_pt']['value'] == 20.74
    assert hci['roles']['body']['font']['value'] == 'Palatino'
    assert hci['roles']['author']['size_pt']['condition'].startswith('separate title-page document')


def test_new_author_roles_apply_only_to_the_host_selected_source_blocks(tmp_path):
    from docx import Document

    source = tmp_path / 'frontmatter.docx'
    document = Document()
    document.add_paragraph('Synthetic authoring manuscript', 'Title')
    document.add_paragraph('Synthetic Author')
    document.add_paragraph('Synthetic Institute')
    document.add_paragraph('Synthetic correspondence details')
    document.save(source)
    profile = load_preset('comp-tf-tgis20-word')
    output = tmp_path / 'frontmatter-formatted.docx'
    result = apply_docx(source, output, profile, {'p1': 'author', 'p2': 'affiliation', 'p3': 'correspondence'})
    assert result['audit']['content']['status'] == result['audit']['style']['status'] == 'PASS'
    blocks = {item['id']: item for item in inspect_docx(output)['blocks']}
    assert blocks['p1']['format']['size_pt'] == 14
    assert blocks['p2']['format']['size_pt'] == 12
    assert blocks['p2']['format']['italic'] is True
    assert blocks['p3']['format']['size_pt'] == 12
