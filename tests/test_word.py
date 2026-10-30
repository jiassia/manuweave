import io
import zipfile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Pt
from PIL import Image
from lxml import etree as ET

from manuweave.word import NS, W, apply_docx, audit_docx, inspect_docx, learn_docx


def fixture(path, styled=False):
    doc = Document()
    doc.add_paragraph('Synthetic study of signal stability', 'Title')
    doc.add_paragraph('Abstract')
    doc.add_paragraph('We observed 12 synthetic measurements and report only fixture data.')
    doc.add_paragraph('Keywords: signal; stability')
    doc.add_paragraph('1 Introduction')
    doc.add_paragraph('The synthetic measurement is 12.5 and remains unchanged.')
    doc.add_paragraph('2 Methods')
    p = doc.add_paragraph()
    math = OxmlElement('m:oMath')
    run = OxmlElement('m:r')
    text = OxmlElement('m:t')
    text.text = 'x=12.5'
    run.append(text)
    math.append(run)
    p._p.append(math)
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).text = 'Merged header'
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(1, 1).text = '12.5'
    table.cell(1, 1).merge(table.cell(2, 1))
    image = io.BytesIO()
    Image.new('RGB', (40, 40), 'blue').save(image, format='PNG')
    image.seek(0)
    doc.add_picture(image)
    p = doc.add_paragraph('Figure 1 Synthetic signal')
    bookmark = OxmlElement('w:bookmarkStart')
    bookmark.set(W+'id', '9')
    bookmark.set(W+'name', 'fig1')
    p._p.append(bookmark)
    end = OxmlElement('w:bookmarkEnd')
    end.set(W+'id', '9')
    p._p.append(end)
    p = doc.add_paragraph('See ')
    f = OxmlElement('w:fldSimple')
    f.set(W+'instr', ' REF fig1 \\h ')
    p._p.append(f)
    doc.add_paragraph('References')
    doc.add_paragraph('[1] Synthetic reference, 2026.')
    doc.sections[0].header.paragraphs[0].text = 'Fixture header'
    if styled:
        # Reference relies on direct formatting, intentionally avoiding pStyle.
        for index, p in enumerate(doc.paragraphs):
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(6)
            for r in p.runs:
                r.font.name = 'Arial'
                r.font.size = Pt(11)
        for p in doc.paragraphs:
            if p.text.startswith(('1 ', '2 ')):
                for r in p.runs:
                    r.font.size = Pt(16)
                    r.bold = True
    doc.save(path)
    return path


def rewrite(source, out, edit):
    with zipfile.ZipFile(source) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    root = ET.fromstring(parts['word/document.xml'])
    edit(root)
    parts['word/document.xml'] = ET.tostring(root)
    with zipfile.ZipFile(out, 'w') as z:
        for n, data in parts.items():
            z.writestr(n, data)


def test_learn_apply_complex_fidelity(tmp_path):
    a = fixture(tmp_path/'source.docx')
    b = fixture(tmp_path/'reference.docx', True)
    profile = learn_docx(b)
    assert profile['rules']['heading1']['size_pt'] == 16
    assert profile['rules']['body']['font'] == 'Arial'
    result = apply_docx(a, tmp_path/'out.docx', profile)
    assert result['audit']['content']['status'] == 'PASS'
    assert result['audit']['style']['status'] == 'PASS'
    assert result['audit']['objects_after'] == result['audit']['objects_before']
    assert result['audit']['objects_after']['equations'] == 1
    with zipfile.ZipFile(a) as src, zipfile.ZipFile(tmp_path/'out.docx') as dst:
        for name in src.namelist():
            if name != 'word/document.xml' and not name.startswith(('word/header', 'word/footer')):
                assert src.read(name) == dst.read(name)


def test_reference_text_not_copied(tmp_path):
    a = fixture(tmp_path/'source.docx')
    doc = Document()
    p = doc.add_paragraph('UNRELATED REFERENCE TITLE')
    p.runs[0].font.size = Pt(22)
    doc.add_paragraph('UNRELATED TEMPLATE CONTENT')
    b = tmp_path/'reference.docx'
    doc.save(b)
    apply_docx(a, tmp_path/'out.docx', learn_docx(b))
    assert all('UNRELATED' not in b['text'] for b in inspect_docx(tmp_path/'out.docx')['blocks'])


def test_data_tampering_detected(tmp_path):
    a = fixture(tmp_path/'source.docx')
    def edit(root):
        for text in root.findall('.//w:t', NS):
            if '12.5' in (text.text or ''):
                text.text = text.text.replace('12.5', '99.9')
    rewrite(a, tmp_path/'bad.docx', edit)
    assert audit_docx(a, tmp_path/'bad.docx')['content']['status'] == 'FAIL'


def test_equation_deletion_detected(tmp_path):
    a = fixture(tmp_path/'source.docx')
    def edit(root):
        el = root.find('.//m:oMath', NS)
        el.getparent().remove(el)
    rewrite(a, tmp_path/'bad.docx', edit)
    assert audit_docx(a, tmp_path/'bad.docx')['content']['status'] == 'FAIL'


def test_stability_and_builtin_profile(tmp_path):
    a = fixture(tmp_path/'source.docx')
    profile = {'rules': {'body': {'font':'Times New Roman','size_pt':12,'line_spacing':1.5,'alignment':'both','first_line_pt':24}}, 'page': {'margin_top_mm':25.4,'columns':1}}
    apply_docx(a, tmp_path/'first.docx', profile)
    apply_docx(tmp_path/'first.docx', tmp_path/'second.docx', profile)
    with zipfile.ZipFile(tmp_path/'first.docx') as first, zipfile.ZipFile(tmp_path/'second.docx') as second:
        assert first.read('word/document.xml') == second.read('word/document.xml')


def test_inherited_style_and_direct_overrides(tmp_path):
    doc = Document()
    doc.styles['Normal'].font.name = 'Arial'
    doc.styles['Normal'].font.size = Pt(10)
    doc.add_paragraph('Title')
    p = doc.add_paragraph('This paragraph uses inherited font and a direct size.')
    p.runs[0].font.size = Pt(13)
    path = tmp_path/'inherited.docx'
    doc.save(path)
    f = inspect_docx(path)['blocks'][1]['format']
    assert f['font'] == 'Arial'
    assert f['size_pt'] == 13


def test_no_overwrite(tmp_path):
    a = fixture(tmp_path/'source.docx')
    original = a.read_bytes()
    with pytest.raises(ValueError):
        apply_docx(a, a, {})
    assert a.read_bytes() == original
    out = tmp_path/'out.docx'
    out.write_text('Existing file')
    with pytest.raises(ValueError):
        apply_docx(a, out, {})
    assert out.read_text() == 'Existing file'


@pytest.mark.parametrize('member,payload', [('../escape', b'x'), ('word/vbaProject.bin', b'x'), ('word/document.xml', b'<!DOCTYPE x [<!ENTITY xx SYSTEM "file:///etc/passwd">]><x>&xx;</x>')])
def test_malicious_packages_rejected(tmp_path, member, payload):
    path = tmp_path/'bad.docx'
    with zipfile.ZipFile(path, 'w') as z:
        z.writestr(member, payload)
    with pytest.raises(ValueError):
        inspect_docx(path)


def test_unknown_plan_property_rejected(tmp_path):
    a = fixture(tmp_path/'source.docx')
    with pytest.raises(ValueError):
        apply_docx(a, tmp_path/'out.docx', {'rules':{'body':{'execute_python':'evil'}}})

def test_semantic_paragraph_properties_tampering_detected(tmp_path):
    a = fixture(tmp_path/'source.docx')
    def edit(root):
        p = root.find('.//w:body/w:p', NS)
        ppr = p.find('w:pPr', NS)
        num = ET.SubElement(ppr, W+'numPr')
        ET.SubElement(num, W+'numId').set(W+'val','42')
    rewrite(a, tmp_path/'bad.docx', edit)
    assert audit_docx(a, tmp_path/'bad.docx')['content']['status'] == 'FAIL'


def test_theme_font_resolved_and_justify_alias(tmp_path):
    doc = Document()
    doc.add_paragraph('Title')
    doc.add_paragraph('Body inherits the default theme font.')
    source = tmp_path/'theme.docx'
    doc.save(source)
    assert inspect_docx(source)['blocks'][1]['format']['font'] == 'Cambria'
    result = apply_docx(source, tmp_path/'out.docx', {'rules':{'body':{'alignment':'justify'}}})
    assert result['audit']['style']['status'] == 'PASS'


@pytest.mark.parametrize('properties', [{'size_pt':float('nan')},{'bold':'false'},{'line_spacing':-1}])
def test_invalid_measurements_rejected(tmp_path, properties):
    a = fixture(tmp_path/'source.docx')
    with pytest.raises(ValueError):
        apply_docx(a, tmp_path/'out.docx', {'rules':{'body':properties}})

def test_table_running_and_page_learning(tmp_path):
    from docx.shared import Mm
    source = fixture(tmp_path/'source.docx')
    reference = fixture(tmp_path/'reference.docx', True)
    doc = Document(reference)
    doc.sections[0].page_width = Mm(210)
    doc.sections[0].page_height = Mm(297)
    doc.sections[0].header.paragraphs[0].text = 'UNRELATED REFERENCE HEADER'
    doc.sections[0].footer.paragraphs[0].text = 'UNRELATED REFERENCE FOOTER'
    for part in (doc.sections[0].header, doc.sections[0].footer):
        p = part.paragraphs[0]
        p.alignment = 1
        p.runs[0].font.name = 'Arial'
        p.runs[0].font.size = Pt(9)
    table = doc.tables[0]
    borders = OxmlElement('w:tblBorders')
    for edge in ('top','bottom','left','right','insideH','insideV'):
        border = OxmlElement('w:'+edge)
        border.set(W+'val', 'single' if edge in ('top','bottom') else 'nil')
        border.set(W+'sz','8')
        border.set(W+'color','000000')
        borders.append(border)
    table._tbl.tblPr.append(borders)
    for cell in table.rows[0].cells:
        props = cell._tc.get_or_add_tcPr()
        old = props.find('w:tcBorders', NS)
        if old is None:
            old = OxmlElement('w:tcBorders')
            props.append(old)
        bottom = OxmlElement('w:bottom')
        bottom.set(W+'val','single')
        bottom.set(W+'sz','4')
        bottom.set(W+'color','000000')
        old.append(bottom)
    doc.save(reference)
    doc = Document(source)
    doc.sections[0].footer.paragraphs[0].text = 'SOURCE FOOTER'
    doc.save(source)
    profile = learn_docx(reference)
    assert profile['tables']['borders']['top']['size_pt'] == 1
    assert profile['tables']['header_bottom']['size_pt'] == 0.5
    assert profile['page']['width_mm'] == pytest.approx(210,abs=0.03)
    assert profile['running']['header']['size_pt'] == 9
    out = tmp_path/'out.docx'
    result = apply_docx(source,out,profile)
    assert result['audit']['content']['status'] == 'PASS'
    assert result['audit']['style']['status'] == 'PASS'
    inspection = inspect_docx(out)
    assert inspection['running']['header'][0]['text'] == 'Fixture header'
    assert inspection['running']['footer'][0]['text'] == 'SOURCE FOOTER'
    assert inspection['running']['header'][0]['format']['size_pt'] == 9
    assert inspection['objects'] == inspect_docx(source)['objects']
    assert inspection['page'][-1]['height_mm'] == pytest.approx(297,abs=0.03)


def test_fit_columns_preserves_image_and_table_content(tmp_path):
    from docx.shared import Mm
    source = fixture(tmp_path/'source.docx')
    doc = Document(source)
    doc.inline_shapes[0].width = Mm(150)
    doc.inline_shapes[0].height = Mm(75)
    doc.save(source)
    profile = {'rules':{},'page':{'width_mm':210,'height_mm':297,'margin_left_mm':25,'margin_right_mm':25,'columns':2}}
    out = tmp_path/'out.docx'
    result = apply_docx(source,out,profile)
    assert result['audit']['content']['status'] == 'PASS'
    assert result['audit']['style']['status'] == 'PASS'
    inspection = inspect_docx(out)
    assert inspection['layout_risks'] == []
    image = next(o for o in inspection['layout_objects'] if o['role']=='image')
    assert image['width_mm'] == pytest.approx((160-12.7)/2,abs=0.03)
    assert image['width_mm']/image['height_mm'] == pytest.approx(2,abs=0.01)
    assert any(c.get('operation') == 'fit_to_column' for c in result['changes'])
    with zipfile.ZipFile(source) as a,zipfile.ZipFile(out) as b:
        for name in a.namelist():
            if name.startswith('word/media/') or name.endswith('.rels'):
                assert a.read(name)==b.read(name)
    second = tmp_path/'second.docx'
    apply_docx(out,second,profile)
    with zipfile.ZipFile(out) as a,zipfile.ZipFile(second) as b:
        assert a.read('word/document.xml')==b.read('word/document.xml')


def test_overflow_audit_detects_unfitted_shape(tmp_path):
    from docx.shared import Mm
    source = fixture(tmp_path/'source.docx')
    doc = Document(source)
    doc.inline_shapes[0].width = Mm(300)
    doc.save(source)
    report = audit_docx(source,source,{})
    assert report['style']['status']=='FAIL'
    assert any(i['property']=='column_overflow' for i in report['style']['issues'])


def test_running_text_tampering_detected(tmp_path):
    source=fixture(tmp_path/'source.docx')
    doc=Document(source)
    doc.sections[0].header.paragraphs[0].text='Unexpected altered author'
    out=tmp_path/'bad.docx'
    doc.save(out)
    assert audit_docx(source,out)['content']['status']=='FAIL'

def _toc_fixture(path):
    doc=Document()
    doc.add_paragraph('Title')
    p=doc.add_paragraph()
    field=OxmlElement('w:fldSimple')
    field.set(W+'instr',' TOC \\o "1-3" ')
    field.set(W+'dirty','true')
    run=OxmlElement('w:r')
    text=OxmlElement('w:t')
    text.text='1 Introduction'
    run.append(text)
    field.append(run)
    p._p.append(field)
    doc.add_paragraph('1 Introduction')
    doc.add_paragraph('The original measurement is 12.5.')
    doc.save(path)
    return path


def test_toc_cache_refresh_is_format_change(tmp_path):
    source=_toc_fixture(tmp_path/'source.docx')
    def edit(root):
        field=root.find('.//w:fldSimple',NS)
        field.set(W+'dirty','false')
        for child in list(field):
            field.remove(child)
        run=ET.SubElement(field,W+'r')
        ET.SubElement(run,W+'t').text='1 Introduction'
        ET.SubElement(run,W+'tab')
        ET.SubElement(run,W+'t').text='2'
        ET.SubElement(run,W+'br')
    out=tmp_path/'refreshed.docx'
    rewrite(source,out,edit)
    assert inspect_docx(out)['blocks'][1]['role']=='toc'
    assert audit_docx(source,out)['content']['status']=='PASS'


def test_toc_refresh_does_not_hide_body_tampering(tmp_path):
    source=_toc_fixture(tmp_path/'source.docx')
    def edit(root):
        for text in root.findall('.//w:t',NS):
            if '12.5' in (text.text or ''):
                text.text=text.text.replace('12.5','999')
    out=tmp_path/'bad.docx'
    rewrite(source,out,edit)
    assert audit_docx(source,out)['content']['status']=='FAIL'


def test_toc_instruction_and_objects_remain_audited(tmp_path):
    source=_toc_fixture(tmp_path/'source.docx')
    def instruction(root):
        root.find('.//w:fldSimple',NS).set(W+'instr',' REF unrelated ')
    out=tmp_path/'bad.docx'
    rewrite(source,out,instruction)
    assert audit_docx(source,out)['content']['status']=='FAIL'
    def object_added(root):
        field=root.find('.//w:fldSimple',NS)
        math=ET.SubElement(field,'{'+NS['m']+'}oMath')
        ET.SubElement(math,'{'+NS['m']+'}r')
    out=tmp_path/'object.docx'
    rewrite(source,out,object_added)
    assert audit_docx(source,out)['content']['status']=='FAIL'


def test_unrelated_field_cache_remains_audited(tmp_path):
    source=_toc_fixture(tmp_path/'source.docx')
    def non_toc(root):
        root.find('.//w:fldSimple',NS).set(W+'instr',' TOCevil ')
    first=tmp_path/'first.docx'
    rewrite(source,first,non_toc)
    def mutate(root):
        root.find('.//w:fldSimple/w:r/w:t',NS).text='Changed ordinary field result'
    second=tmp_path/'second.docx'
    rewrite(first,second,mutate)
    assert audit_docx(first,second)['content']['status']=='FAIL'

def test_heading_keep_next_audit_and_repair(tmp_path):
    source=fixture(tmp_path/'source.docx')
    report=audit_docx(source,source,{'rules':{}})
    assert report['style']['status']=='FAIL'
    assert any(i['property']=='keepNext' for i in report['style']['issues'])
    out=tmp_path/'out.docx'
    report=apply_docx(source,out,{'rules':{}})['audit']
    assert report['content']['status']=='PASS'
    assert report['style']['status']=='PASS'
    with zipfile.ZipFile(out) as z:
        root=ET.fromstring(z.read('word/document.xml'))
    for block in inspect_docx(out)['blocks']:
        if block['role'] in ('heading1','abstract_heading','references_heading'):
            p=root.findall('.//w:body//w:p',NS)[int(block['id'][1:])]
            assert p.find('w:pPr/w:keepNext',NS).get(W+'val')=='1'
            assert p.find('w:pPr/w:keepLines',NS).get(W+'val')=='1'
    second=tmp_path/'second.docx'
    apply_docx(out,second,{'rules':{}})
    with zipfile.ZipFile(out) as a,zipfile.ZipFile(second) as b:
        assert a.read('word/document.xml')==b.read('word/document.xml')


def test_heading_keep_rules_respect_role_overrides(tmp_path):
    doc=Document()
    doc.add_paragraph('Title')
    doc.add_paragraph('A custom section label')
    source=tmp_path/'source.docx'
    doc.save(source)
    roles={'p1':'heading2'}
    assert audit_docx(source,source,{},roles=roles)['style']['status']=='FAIL'
    out=tmp_path/'out.docx'
    report=apply_docx(source,out,{},roles=roles)['audit']
    assert report['style']['status']=='PASS'
    with zipfile.ZipFile(out) as z:
        root=ET.fromstring(z.read('word/document.xml'))
    assert root.findall('.//w:body//w:p',NS)[1].find('w:pPr/w:keepNext',NS) is not None

def test_strict_ooxml_is_rejected_without_empty_success(tmp_path):
    source=fixture(tmp_path/'transitional.docx')
    strict=tmp_path/'strict.docx'
    with zipfile.ZipFile(source) as archive, zipfile.ZipFile(strict,'w') as out:
        for name in archive.namelist():
            payload=archive.read(name)
            if name.endswith(('.xml','.rels')):
                payload=payload.replace(NS['w'].encode(),b'http://purl.oclc.org/ooxml/wordprocessingml/main')
            out.writestr(name,payload)
    original=strict.read_bytes()
    for action in (lambda:inspect_docx(strict),lambda:learn_docx(strict),lambda:audit_docx(strict,strict,{}),lambda:apply_docx(strict,tmp_path/'out.docx',{})):
        with pytest.raises(ValueError,match='Strict OOXML.*Transitional'):
            action()
    assert strict.read_bytes()==original
    assert not (tmp_path/'out.docx').exists()


@pytest.mark.parametrize('xml',[
    '<w:document xmlns:w="'+NS['w']+'"><w:p/></w:document>',
    '<w:styles xmlns:w="'+NS['w']+'"><w:body/></w:styles>',
    '<document><body/></document>',
    '<w:document xmlns:w="'+NS['w']+'"><w:body/><w:body/></w:document>',
])
def test_fake_word_without_valid_body_rejected(tmp_path,xml):
    path=tmp_path/'fake.docx'
    with zipfile.ZipFile(path,'w') as archive:
        archive.writestr('word/document.xml',xml)
    with pytest.raises(ValueError,match='Invalid Transitional Word document'):
        inspect_docx(path)

@pytest.mark.parametrize('rule,points',[('exact',9),('atLeast',13.5),('atLeast',11.25),('atLeast',8.4)])
def test_fixed_minimum_spacing_and_hanging_learning(tmp_path,rule,points):
    from docx.enum.text import WD_LINE_SPACING
    doc=Document()
    doc.styles['Normal'].paragraph_format.line_spacing=1.5
    doc.add_paragraph('Title')
    p=doc.add_paragraph('The reference body has a controlled baseline grid.')
    p.paragraph_format.line_spacing=Pt(points)
    p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.EXACTLY if rule=='exact' else WD_LINE_SPACING.AT_LEAST
    p.paragraph_format.left_indent=Pt(18)
    p.paragraph_format.right_indent=Pt(6)
    p.paragraph_format.first_line_indent=Pt(-18)
    p.runs[0].font.color.rgb=__import__('docx').shared.RGBColor.from_string('C00000')
    reference=tmp_path/'reference.docx'
    doc.save(reference)
    block=inspect_docx(reference)['blocks'][1]
    assert block['format']['line_spacing_rule']==rule
    assert block['format']['line_spacing_pt']==pytest.approx(points,abs=0.01)
    assert 'line_spacing' not in block['format']
    assert block['format']['first_line_pt']==-18
    profile=learn_docx(reference)
    assert profile['rules']['body']['line_spacing_pt']==pytest.approx(points,abs=0.01)
    assert 'line_spacing' not in profile['rules']['body']
    source=fixture(tmp_path/'source.docx')
    out=tmp_path/'out.docx'
    result=apply_docx(source,out,profile)
    assert result['audit']['content']['status']=='PASS'
    assert result['audit']['style']['status']=='PASS'
    body=next(b for b in inspect_docx(out)['blocks'] if b['role']=='body')
    assert body['format']['first_line_pt']==-18
    assert body['format']['left_indent_pt']==18
    assert body['format']['right_indent_pt']==6
    assert body['format']['font_color']=='C00000'
    assert body['format']['line_spacing_pt']==pytest.approx(points,abs=0.01)
    assert 'line_spacing' not in body['format']
    second=tmp_path/'second.docx'
    apply_docx(out,second,profile)
    with zipfile.ZipFile(out) as a,zipfile.ZipFile(second) as b:
        assert a.read('word/document.xml')==b.read('word/document.xml')
    auto=tmp_path/'auto.docx'
    applied=apply_docx(out,auto,{'rules':{'body':{'line_spacing':1.5}}})
    assert applied['audit']['style']['status']=='PASS'
    body=next(b for b in inspect_docx(auto)['blocks'] if b['role']=='body')
    assert body['format']['line_spacing']==1.5
    assert body['format']['line_spacing_rule']=='auto'
    assert 'line_spacing_pt' not in body['format']


def test_spacing_raw_line_inherits_previous_rule(tmp_path):
    from docx.enum.text import WD_LINE_SPACING
    doc=Document()
    doc.styles['Normal'].paragraph_format.line_spacing=Pt(11)
    doc.styles['Normal'].paragraph_format.line_spacing_rule=WD_LINE_SPACING.AT_LEAST
    doc.add_paragraph('Title')
    p=doc.add_paragraph('Paragraph overrides only the raw line amount.')
    spacing=OxmlElement('w:spacing')
    spacing.set(W+'line','270')
    p._p.get_or_add_pPr().append(spacing)
    path=tmp_path/'inherited.docx'
    doc.save(path)
    fmt=inspect_docx(path)['blocks'][1]['format']
    assert fmt['line_spacing_rule']=='atLeast'
    assert fmt['line_spacing_pt']==13.5
    assert 'line_spacing' not in fmt


def test_mixed_line_modes_learn_consistent_profile(tmp_path):
    from docx.enum.text import WD_LINE_SPACING
    doc=Document()
    doc.add_paragraph('Title')
    for index in range(3):
        p=doc.add_paragraph('Paragraph '+str(index))
        if index==2:
            p.paragraph_format.line_spacing=Pt(9)
            p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.EXACTLY
        else:
            p.paragraph_format.line_spacing=1.5
    reference=tmp_path/'mixed.docx'
    doc.save(reference)
    rules=learn_docx(reference)['rules']['body']
    assert rules['line_spacing_rule']=='auto'
    assert rules['line_spacing']==1.5
    assert 'line_spacing_pt' not in rules


def test_caption_subroles_and_legacy_fallback(tmp_path):
    doc=Document()
    doc.add_paragraph('Title')
    doc.add_paragraph('Figure 1 Signal trace')
    doc.add_paragraph('Table 1 Measurement summary')
    doc.add_paragraph('Caption without a label','Caption')
    source=tmp_path/'source.docx'
    doc.save(source)
    roles=[b['role'] for b in inspect_docx(source)['blocks']]
    assert roles==['title','figure_caption','table_caption','caption']
    legacy=tmp_path/'legacy.docx'
    assert apply_docx(source,legacy,{'rules':{'caption':{'size_pt':9}}})['audit']['style']['status']=='PASS'
    assert all(b['format']['size_pt']==9 for b in inspect_docx(legacy)['blocks'][1:])
    distinct=tmp_path/'distinct.docx'
    profile={'rules':{'caption':{'size_pt':10},'figure_caption':{'size_pt':8,'space_before_pt':6},'table_caption':{'size_pt':9,'space_after_pt':3}}}
    assert apply_docx(source,distinct,profile)['audit']['style']['status']=='PASS'
    blocks=inspect_docx(distinct)['blocks']
    assert [b['format']['size_pt'] for b in blocks[1:]]==[8,9,10]
    learned=learn_docx(distinct)['rules']
    assert learned['figure_caption']['space_before_pt']==6
    assert learned['table_caption']['space_after_pt']==3

def test_style_raw_line_inherits_document_default_rule(tmp_path):
    doc=Document()
    doc.add_paragraph('Title')
    doc.add_paragraph('Body style inherits minimum line spacing from document defaults.')
    source=tmp_path/'source.docx'
    doc.save(source)
    out=tmp_path/'default-rule.docx'
    with zipfile.ZipFile(source) as archive:
        parts={name:archive.read(name) for name in archive.namelist()}
    styles=ET.fromstring(parts['word/styles.xml'])
    defaults=styles.find('w:docDefaults/w:pPrDefault/w:pPr',NS)
    spacing=defaults.find('w:spacing',NS)
    spacing.set(W+'lineRule','atLeast')
    spacing.set(W+'line','220')
    normal=styles.find('w:style[@w:styleId="Normal"]',NS)
    props=normal.find('w:pPr',NS)
    if props is None:
        props=ET.SubElement(normal,W+'pPr')
    spacing=ET.SubElement(props,W+'spacing')
    spacing.set(W+'line','270')
    parts['word/styles.xml']=ET.tostring(styles)
    with zipfile.ZipFile(out,'w') as archive:
        for name,data in parts.items():
            archive.writestr(name,data)
    fmt=inspect_docx(out)['blocks'][1]['format']
    assert fmt['line_spacing_rule']=='atLeast'
    assert fmt['line_spacing_pt']==13.5
    assert 'line_spacing' not in fmt


@pytest.mark.parametrize('properties',[
    {'line_spacing_pt':9},
    {'line_spacing_rule':'exact','line_spacing_pt':9,'line_spacing':1.5},
    {'line_spacing_rule':'atLeast'},
    {'font_color':'javascript:bad'},
])
def test_incompatible_spacing_and_invalid_color_rejected(tmp_path,properties):
    source=fixture(tmp_path/'source.docx')
    with pytest.raises(ValueError):
        apply_docx(source,tmp_path/'out.docx',{'rules':{'body':properties}})

@pytest.mark.parametrize('label,role',[('TABLE I Summary','table_caption'),('图一 信号分布','figure_caption'),('Table data remains factual body text.','body')])
def test_caption_labels_support_roman_and_chinese_numbers(tmp_path,label,role):
    doc=Document()
    doc.add_paragraph('Title')
    doc.add_paragraph(label)
    source=tmp_path/'caption.docx'
    doc.save(source)
    assert inspect_docx(source)['blocks'][1]['role']==role

SECTION_RULE={'front_matter_columns':1,'body_columns':2,'body_starts_at':'first_heading1'}


def add_footnote(path):
    with zipfile.ZipFile(path) as z:
        parts={name:z.read(name) for name in z.namelist()}
    root=ET.fromstring(parts['word/document.xml'])
    paragraph=root.findall('.//w:body/w:p',NS)[5]
    run=ET.SubElement(paragraph,W+'r')
    ET.SubElement(run,W+'footnoteReference',{W+'id':'2'})
    parts['word/document.xml']=ET.tostring(root)
    footnotes=ET.Element(W+'footnotes',nsmap={'w':NS['w']})
    note=ET.SubElement(footnotes,W+'footnote',{W+'id':'2'})
    p=ET.SubElement(note,W+'p')
    r=ET.SubElement(p,W+'r')
    ET.SubElement(r,W+'t').text='Footnote contains protected measurement 12.5.'
    parts['word/footnotes.xml']=ET.tostring(footnotes)
    rel_ns='http://schemas.openxmlformats.org/package/2006/relationships'
    rels=ET.fromstring(parts['word/_rels/document.xml.rels'])
    ET.SubElement(rels,'{'+rel_ns+'}Relationship',Id='rIdManuWeaveFootnoteFixture',Type=NS['r']+'/footnotes',Target='footnotes.xml')
    parts['word/_rels/document.xml.rels']=ET.tostring(rels)
    ct_ns='http://schemas.openxmlformats.org/package/2006/content-types'
    content_types=ET.fromstring(parts['[Content_Types].xml'])
    ET.SubElement(content_types,'{'+ct_ns+'}Override',PartName='/word/footnotes.xml',ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml')
    parts['[Content_Types].xml']=ET.tostring(content_types)
    with zipfile.ZipFile(path,'w') as z:
        for name,data in parts.items():
            z.writestr(name,data)


@pytest.mark.parametrize('flags',[{}, {'different_first_page':True,'different_odd_even_headers':True}, {'different_first_page':False,'different_odd_even_headers':False}])
@pytest.mark.parametrize('notes',[False,True])
def test_two_region_sections_integrity_and_idempotency(tmp_path,flags,notes):
    source=fixture(tmp_path/'source.docx')
    if notes:
        add_footnote(source)
    doc=Document(source)
    page_number=OxmlElement('w:pgNumType')
    page_number.set(W+'start','7')
    doc.sections[0]._sectPr.append(page_number)
    doc.save(source)
    profile={'rules':{},'page':{'width_mm':210,'height_mm':297,'margin_left_mm':20,'margin_right_mm':20,'columns':2},'sections':{**SECTION_RULE,'body_column_gap_mm':7,**flags}}
    out=tmp_path/'out.docx'
    result=apply_docx(source,out,profile)
    assert result['audit']['content']['status']=='PASS',result['audit']
    assert result['audit']['style']['status']=='PASS',result['audit']
    assert result['audit']['objects_before']==result['audit']['objects_after']
    with zipfile.ZipFile(source) as a,zipfile.ZipFile(out) as b:
        for name in a.namelist():
            if name not in ('word/document.xml','word/settings.xml'):
                assert a.read(name)==b.read(name),name
        root=ET.fromstring(b.read('word/document.xml'))
        sections=root.findall('.//w:sectPr',NS)
        assert len(sections)==2
        assert sections[0].find('w:cols',NS).get(W+'num')=='1'
        assert sections[1].find('w:cols',NS).get(W+'num')=='2'
        assert sections[1].find('w:type',NS).get(W+'val')=='continuous'
        assert sections[0].find('w:pgNumType',NS).get(W+'start')=='7'
        assert sections[1].find('w:pgNumType',NS).get(W+'start') is None
        assert [n.attrib for n in sections[0].findall('w:headerReference',NS)]==[n.attrib for n in sections[1].findall('w:headerReference',NS)]
    second=tmp_path/'second.docx'
    again=apply_docx(out,second,profile)
    assert again['audit']['style']['status']=='PASS'
    with zipfile.ZipFile(out) as a,zipfile.ZipFile(second) as b:
        for name in a.namelist():
            assert a.read(name)==b.read(name),name


def test_two_region_sections_use_explicit_role_boundary(tmp_path):
    source=fixture(tmp_path/'source.docx')
    # Override the first inferred heading so the second heading starts the body.
    roles={'p4':'body','p6':'heading1'}
    profile={'rules':{},'sections':SECTION_RULE}
    out=tmp_path/'out.docx'
    assert apply_docx(source,out,profile,roles)['audit']['style']['status']=='PASS'
    with zipfile.ZipFile(out) as z:
        root=ET.fromstring(z.read('word/document.xml'))
    assert root.findall('.//w:body//w:p',NS)[5].find('w:pPr/w:sectPr',NS) is not None


def test_two_region_sections_reject_existing_complex_boundaries(tmp_path):
    from docx.enum.section import WD_SECTION_START
    source=fixture(tmp_path/'source.docx')
    doc=Document(source)
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    doc.add_paragraph('A separate appendix section must survive.')
    doc.save(source)
    original=source.read_bytes()
    out=tmp_path/'out.docx'
    with pytest.raises(ValueError,match='complex section'):
        apply_docx(source,out,{'rules':{},'sections':SECTION_RULE})
    assert source.read_bytes()==original
    assert not out.exists()


def test_two_region_audit_detects_style_and_semantic_damage(tmp_path):
    source=fixture(tmp_path/'source.docx')
    profile={'rules':{},'sections':{**SECTION_RULE,'different_first_page':True,'different_odd_even_headers':True}}
    out=tmp_path/'out.docx'
    apply_docx(source,out,profile)
    def damage_columns(root):
        root.find('w:body/w:sectPr/w:cols',NS).set(W+'num','1')
    bad=tmp_path/'bad-columns.docx'
    rewrite(out,bad,damage_columns)
    audit=audit_docx(source,bad,profile)
    assert audit['content']['status']=='PASS'
    assert audit['style']['status']=='FAIL'
    def damage_reference(root):
        front=root.find('.//w:pPr/w:sectPr',NS)
        front.find('w:headerReference',NS).set('{'+NS['r']+'}id','unauthorized-header')
    bad=tmp_path/'bad-reference.docx'
    rewrite(out,bad,damage_reference)
    assert audit_docx(source,bad,profile)['content']['status']=='FAIL'
    def damage_data(root):
        for node in root.findall('.//w:t',NS):
            if '12.5' in (node.text or ''):
                node.text=node.text.replace('12.5','99.9')
    bad=tmp_path/'bad-data.docx'
    rewrite(out,bad,damage_data)
    assert audit_docx(source,bad,profile)['content']['status']=='FAIL'


@pytest.mark.parametrize('rule',[{**SECTION_RULE,'front_matter_columns':True},{**SECTION_RULE,'different_first_page':'true'},{**SECTION_RULE,'body_column_gap_mm':'7'},{**SECTION_RULE,'execute':'unknown'}])
def test_two_region_schema_is_strict(tmp_path,rule):
    from pydantic import ValidationError
    source=fixture(tmp_path/'source.docx')
    with pytest.raises(ValidationError):
        apply_docx(source,tmp_path/'out.docx',{'rules':{},'sections':rule})

@pytest.mark.parametrize('override_body',[False,True])
def test_pure_image_does_not_inherit_text_exact_line_height(tmp_path,override_body):
    from docx.enum.text import WD_LINE_SPACING
    source=fixture(tmp_path/'source.docx')
    doc=Document(source)
    doc.styles['Normal'].paragraph_format.line_spacing=Pt(15)
    doc.styles['Normal'].paragraph_format.line_spacing_rule=WD_LINE_SPACING.EXACTLY
    doc.save(source)
    before=inspect_docx(source)
    image=next(b for b in before['blocks'] if b['role']=='image')
    assert before['linebox_risks']
    roles={image['id']:'body'} if override_body else None
    profile={'rules':{'body':{'font':'Arial','size_pt':10,'line_spacing_rule':'exact','line_spacing_pt':15}}}
    out=tmp_path/'out.docx'
    result=apply_docx(source,out,profile,roles)
    assert result['audit']['content']['status']=='PASS'
    assert result['audit']['style']['status']=='PASS'
    after=inspect_docx(out)
    image=next(b for b in after['blocks'] if b['id']==image['id'])
    assert image['role']=='image'
    assert image['format']['line_spacing_rule']=='atLeast'
    assert after['linebox_risks']==[]
    assert all(b['format']['line_spacing_rule']=='exact' and b['format']['line_spacing_pt']==15 for b in after['blocks'] if b['role']=='body' and b['text'].strip())
    assert any(c.get('operation')=='fit_inline_image_linebox' for c in result['changes'])
    second=tmp_path/'second.docx'
    apply_docx(out,second,profile,roles)
    with zipfile.ZipFile(out) as a,zipfile.ZipFile(second) as b:
        assert a.read('word/document.xml')==b.read('word/document.xml')


def test_image_linebox_uses_resized_extent_and_mixed_text_is_reported(tmp_path):
    from docx.shared import Mm
    source=fixture(tmp_path/'source.docx')
    doc=Document(source)
    doc.inline_shapes[0].width=Mm(160)
    doc.inline_shapes[0].height=Mm(80)
    doc.save(source)
    profile={'rules':{'body':{'line_spacing_rule':'exact','line_spacing_pt':15},'image':{'line_spacing_rule':'exact','line_spacing_pt':15}},'page':{'width_mm':210,'margin_left_mm':25,'margin_right_mm':25,'columns':2}}
    out=tmp_path/'out.docx'
    # An explicitly unsafe image rule remains an audit failure; the actual
    # object still receives a non-clipping linebox instead of losing pixels.
    result=apply_docx(source,out,profile)
    assert result['audit']['content']['status']=='PASS'
    assert result['audit']['style']['status']=='FAIL'
    after=inspect_docx(out)
    image=next(b for b in after['blocks'] if b['role']=='image')
    drawing=next(b for b in after['layout_objects'] if b['role']=='image')
    assert image['format']['line_spacing_pt']==pytest.approx(drawing['height_mm']*72/25.4+2,abs=0.06)
    assert drawing['height_mm']<40
    doc=Document(source)
    image_paragraph=next(p for p in doc.paragraphs if p._p.find('.//w:drawing',NS) is not None)
    image_paragraph.add_run('This image shares a paragraph with protected body text 12.5.')
    mixed=tmp_path/'mixed.docx'
    doc.save(mixed)
    mixed_out=tmp_path/'mixed-out.docx'
    result=apply_docx(mixed,mixed_out,{'rules':{'body':{'line_spacing_rule':'exact','line_spacing_pt':15}}})
    assert result['audit']['content']['status']=='PASS'
    assert result['audit']['style']['status']=='FAIL'
    issues=[i for i in result['audit']['style']['issues'] if i['property']=='inline_image_linebox']
    assert issues and not issues[0]['pure_image']
    paragraph=next(b for b in inspect_docx(mixed_out)['blocks'] if 'shares a paragraph' in b['text'])
    assert paragraph['format']['line_spacing_rule']=='exact'
    assert paragraph['format']['line_spacing_pt']==15


def test_native_word_exact_line_image_is_not_clipped(tmp_path):
    from docx.enum.text import WD_LINE_SPACING
    from docx.shared import Mm,RGBColor
    from PIL import ImageDraw
    from manuweave.render import doctor,render
    if not doctor()['renderers']['word_automation']:
        pytest.skip('Native Microsoft Word automation is not installed')
    doc=Document()
    normal=doc.styles['Normal']
    normal.font.name='Arial'
    normal.font.size=Pt(10)
    normal.font.color.rgb=RGBColor(0,0,0)
    normal.paragraph_format.line_spacing=Pt(15)
    normal.paragraph_format.line_spacing_rule=WD_LINE_SPACING.EXACTLY
    doc.add_paragraph('Synthetic exact-line-height image test')
    doc.add_paragraph('This ordinary body paragraph must retain exact 15-point line spacing.')
    picture=Image.new('RGB',(420,180),'white')
    draw=ImageDraw.Draw(picture)
    draw.line((20,160,400,160),fill='black',width=3)
    draw.line((20,15,20,160),fill='black',width=3)
    points=[(30,145),(110,120),(190,90),(270,65),(350,35)]
    draw.line(points,fill='blue',width=3)
    for x,y in points:
        draw.ellipse((x-5,y-5,x+5,y+5),fill='blue')
    stream=io.BytesIO()
    picture.save(stream,format='PNG')
    stream.seek(0)
    doc.add_picture(stream,width=Mm(120))
    doc.add_paragraph('Figure 1 Synthetic five-point plot')
    source=tmp_path/'source.docx'
    doc.save(source)
    bad=render(source,tmp_path/'clipped-render',renderer='word')
    assert bad['status']=='PASS',bad
    profile={'rules':{'body':{'font':'Arial','size_pt':10,'line_spacing_rule':'exact','line_spacing_pt':15}}}
    output=tmp_path/'fixed.docx'
    result=apply_docx(source,output,profile)
    assert result['audit']['style']['status']=='PASS'
    repaired=render(output,tmp_path/'repaired-render',renderer='word')
    assert repaired['status']=='PASS',repaired
    def blue_height(path):
        with Image.open(path) as png:
            image=png.convert('RGB')
            pixels=image.load()
            rows=[y for y in range(image.height) if any(b>150 and r<100 and g<120 for r,g,b in [pixels[x,y] for x in range(image.width)])]
        return max(rows)-min(rows) if rows else 0
    bad_height=blue_height(tmp_path/'clipped-render'/'page-1.png')
    full_height=blue_height(tmp_path/'repaired-render'/'page-1.png')
    assert full_height>80
    assert full_height>bad_height+50


def test_bilingual_journal_formats_preserve_distinct_abstracts_and_captions(tmp_path):
    source = tmp_path / "bilingual.docx"
    document = Document()
    for text in ("中文标题", "English title", "摘要：中文内容", "Abstract: English summary", "关键词：测试", "Keywords: testing", "图1 中文图题", "Figure 1 English caption", "表1 中文表题", "Table 1 English caption"):
        document.add_paragraph(text)
    document.save(source)
    sizes = {"title_zh": 22, "title_en": 14, "abstract_zh": 9, "abstract_en": 10.5,
             "keywords_zh": 9, "keywords_en": 10.5, "figure_caption_zh": 8,
             "figure_caption_en": 9, "table_caption_zh": 8, "table_caption_en": 9}
    profile = {"id": "bilingual-journal", "name": "Bilingual journal fixture", "version": "1.0.0",
               "page": {}, "rules": {role: {"size_pt": size} for role, size in sizes.items()}}
    from manuweave.models import StyleProfile
    StyleProfile.model_validate(profile, strict=True)
    output = tmp_path / "formatted.docx"
    roles = {"p1": "title_en"}
    report = apply_docx(source, output, profile, roles)
    assert report["audit"]["content"]["status"] == "PASS"
    assert report["audit"]["style"]["status"] == "PASS"
    blocks = inspect_docx(output)["blocks"]
    assert [block["format"]["size_pt"] for block in blocks] == [22, 14, 9, 10.5, 9, 10.5, 8, 9, 8, 9]
    # A wrong English abstract size must fail style audit even when Chinese is correct.
    changed = Document(output)
    for run in changed.paragraphs[3].runs:
        run.font.size = Pt(9)
    damaged = tmp_path / "wrong-english-size.docx"
    changed.save(damaged)
    audit = audit_docx(source, damaged, profile, roles)
    assert audit["content"]["status"] == "PASS"
    assert audit["style"]["status"] == "FAIL"
    assert any(issue["id"] == "p3" and issue["property"] == "size_pt" for issue in audit["style"]["issues"])


def test_explicit_language_role_uses_generic_rule_when_profile_has_no_override(tmp_path):
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("Title")
    document.add_paragraph("Abstract: a generic journal abstract")
    document.add_paragraph("Figure 1 Caption")
    document.save(source)
    profile = {"rules": {"abstract": {"size_pt": 11}, "caption": {"size_pt": 8}}}
    output = tmp_path / "output.docx"
    roles = {"p1": "abstract_en", "p2": "figure_caption_en"}
    assert apply_docx(source, output, profile, roles)["audit"]["style"]["status"] == "PASS"
    assert [b["format"]["size_pt"] for b in inspect_docx(output)["blocks"][1:]] == [11, 8]


def test_source_geometry_grid_and_mirrored_margins_are_applied_and_audited(tmp_path):
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("Synthetic title")
    document.add_paragraph("Main text 12.5 remains unchanged.")
    document.save(source)
    profile = {"rules": {"body": {"size_pt": 7.75}}, "page": {
        "width_mm": 201.4185, "height_mm": 275.5371, "margin_top_mm": 4.1981,
        "margin_bottom_mm": 4.004, "margin_left_mm": 20, "margin_right_mm": 18,
        "columns": 2, "column_gap_mm": 11.0067, "header_distance_mm": 2,
        "footer_distance_mm": 2.5, "mirror_margins": True,
        "grid_type": "lines", "grid_line_pitch_pt": 15.6}}
    from manuweave.models import PageRule
    PageRule.model_validate(profile["page"], strict=True)
    output = tmp_path / "output.docx"
    report = apply_docx(source, output, profile)
    assert report["audit"]["content"]["status"] == "PASS"
    assert report["audit"]["style"]["status"] == "PASS"
    assert not report["audit"]["style"]["exact_source_sizes"]
    assert report["audit"]["style"]["representation_adjustments"][0]["source_value"] == 7.75
    page = inspect_docx(output)["page"][0]
    assert page["column_gap_mm"] == pytest.approx(11.0067, abs=0.02)
    assert page["grid_line_pitch_pt"] == 15.6
    assert page["grid_type"] == "lines"
    with zipfile.ZipFile(output) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    settings = ET.fromstring(parts["word/settings.xml"])
    settings.remove(settings.find("w:mirrorMargins", NS))
    parts["word/settings.xml"] = ET.tostring(settings)
    wrong = tmp_path / "wrong.docx"
    with zipfile.ZipFile(wrong, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    audit = audit_docx(source, wrong, profile)
    assert audit["content"]["status"] == "PASS"
    assert audit["style"]["status"] == "FAIL"
    assert any(issue["property"] == "mirror_margins" for issue in audit["style"]["issues"])


def test_explicit_body_boundary_starts_columns_before_unnumbered_introduction(tmp_path):
    source = tmp_path / "source.docx"
    document = Document()
    for text in ("Title", "Keywords: fixture", "DOI: synthetic", "Unnumbered introduction body", "1 Later section"):
        document.add_paragraph(text)
    document.save(source)
    profile = {"rules": {"body": {"size_pt": 10.5}}, "sections": {
        "front_matter_columns": 1, "body_columns": 2, "body_starts_at": "marked_body_start"}}
    output = tmp_path / "output.docx"
    report = apply_docx(source, output, profile, {"p3": "body_start"})
    assert report["audit"]["content"]["status"] == "PASS"
    assert report["audit"]["style"]["status"] == "PASS"
    with zipfile.ZipFile(output) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    paragraphs = root.findall(".//w:body/w:p", NS)
    assert paragraphs[2].find("w:pPr/w:sectPr/w:cols", NS).get(W + "num") == "1"
    assert root.find("w:body/w:sectPr/w:cols", NS).get(W + "num") == "2"
    assert inspect_docx(output)["blocks"][3]["format"]["size_pt"] == 10.5


def test_paragraph_mark_font_does_not_override_visible_text_inheritance(tmp_path):
    source = tmp_path / "paragraph-mark.docx"
    document = Document()
    style = document.styles.add_style("FixtureAuthor", 1)
    style.font.name = "Arial"
    style.font.size = Pt(14)
    style.font.bold = True
    paragraph = document.add_paragraph(style=style)
    run = paragraph.add_run("Fixture Author")
    run.font.size = Pt(12)
    mark = ET.SubElement(paragraph._p.get_or_add_pPr(), W + "rPr")
    ET.SubElement(mark, W + "rFonts", {W + "ascii": "NEU-BZ", W + "hAnsi": "NEU-BZ"})
    ET.SubElement(mark, W + "sz", {W + "val": "21"})
    ET.SubElement(mark, W + "b", {W + "val": "0"})
    empty = document.add_paragraph()
    empty_mark = ET.SubElement(empty._p.get_or_add_pPr(), W + "rPr")
    ET.SubElement(empty_mark, W + "sz", {W + "val": "21"})
    document.save(source)
    blocks = inspect_docx(source)["blocks"]
    assert blocks[0]["format"]["font"] == "Arial"
    assert blocks[0]["format"]["size_pt"] == 12
    assert blocks[0]["format"]["bold"] is True
    assert blocks[1]["format"]["size_pt"] == 10.5


def test_section_specific_pages_preserve_grid_and_unequal_body_columns(tmp_path):
    source = tmp_path / "source.docx"
    document = Document()
    for text in ("Title", "Abstract: fixture", "Keywords: fixture", "Unnumbered body 12.5", "1 Next section"):
        document.add_paragraph(text)
    document.save(source)
    profile = {"rules": {}, "page": {"width_mm": 210, "height_mm": 297}, "sections": {
        "front_matter_columns": 1, "body_columns": 2, "body_starts_at": "marked_body_start",
        "front_matter_page": {"margin_top_mm": 18, "grid_type": "linesAndChars", "grid_line_pitch_pt": 12},
        "body_page": {"margin_top_mm": 22, "grid_type": "lines", "grid_line_pitch_pt": 15.6,
                      "column_gap_mm": 5, "column_widths_mm": [80.0, 90.0]}}}
    from manuweave.models import StyleProfile
    StyleProfile.model_validate({"id": "section-pages", "name": "Section pages", "version": "1.0.0", **profile}, strict=True)
    output = tmp_path / "output.docx"
    report = apply_docx(source, output, profile, {"p3": "body_start"})
    assert report["audit"]["content"]["status"] == "PASS"
    assert report["audit"]["style"]["status"] == "PASS"
    pages = inspect_docx(output)["page"]
    assert pages[0]["margin_top_mm"] == pytest.approx(18, abs=0.02)
    assert pages[1]["margin_top_mm"] == pytest.approx(22, abs=0.02)
    assert pages[0]["grid_type"] == "linesAndChars"
    assert pages[1]["grid_type"] == "lines"
    assert pages[1]["column_widths_mm"] == pytest.approx([80, 90], abs=0.02)
    repeated = tmp_path / "repeated.docx"
    assert apply_docx(output, repeated, profile, {"p3": "body_start"})["audit"]["content"]["status"] == "PASS"
    with zipfile.ZipFile(output) as first, zipfile.ZipFile(repeated) as second:
        assert {name: first.read(name) for name in first.namelist()} == {name: second.read(name) for name in second.namelist()}
