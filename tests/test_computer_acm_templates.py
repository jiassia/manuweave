"""Source-free ACM formatting facts and generic-engine execution checks."""
import json
from pathlib import Path

import pytest

from manuweave.latex import apply_latex
from manuweave.models import StyleProfile
from manuweave.templates import validate_template

ROOT = Path(__file__).resolve().parents[1] / "src/manuweave/data/templates"
JOURNALS = json.loads((ROOT / "computer-acm/sources/journals.json").read_text(encoding="utf-8"))["journals"]
POINT_FACTOR = 72 / 72.27


def read(identifier, family="computer-acm"):
    return json.loads((ROOT / family / "profiles" / (identifier + ".json")).read_text(encoding="utf-8"))


@pytest.mark.parametrize("journal", JOURNALS, ids=lambda item: item["code"])
def test_each_journal_has_detailed_executable_values_without_vendor_dependencies(journal):
    item = read("acm-" + journal["code"].lower() + "-article-publication")
    StyleProfile.model_validate(item, strict=True)
    assert item["catalog_ids"] == journal["catalog_ids"]
    assert item["journal"]["stage"] == "publication"
    assert item["latex"]["documentclass"] == "article"
    assert "bibliography_style" not in item["latex"]
    validate_template(item)
    contract = item["template"]
    assert contract["assets"] == []
    assert contract["license"] == "rules_only_source_not_redistributed"
    assert contract["engine_support"] == {"docx": "partial", "latex": "partial"}
    assert "fixture" not in contract
    assert "remote_archive" not in item
    extracted = item["extracted_format"]
    assert extracted["mode"] == journal["format"]
    assert {"title", "author", "affiliation", "abstract", "body", "heading1", "heading2", "heading3", "heading4", "reference", "figure_caption", "table_caption"} <= set(extracted["roles"])
    for role, rule in item["rules"].items():
        if "size_pt" in rule:
            source = extracted["roles"][role]["size_pt"]
            assert source["unit"] == "tex_pt"
            assert rule["size_pt"] == pytest.approx(source["value"] * POINT_FACTOR, abs=1e-7)
        if "line_spacing_pt" in rule:
            assert rule["line_spacing_rule"] == "atLeast"
            assert rule["line_spacing_pt"] == pytest.approx(extracted["roles"][role]["line_spacing_pt"]["value"] * POINT_FACTOR, abs=1e-7)
    for group in [*extracted["roles"].values(), extracted["page"], extracted["captions"], extracted["equations"], extracted["references"]]:
        for fact in group.values():
            assert {"value", "unit", "source"} <= set(fact)
            assert fact["source"]["url"].startswith("https://")
            assert fact["source"]["locator"]
            assert "anchor" not in fact["source"]
            assert "quote" not in fact["source"]
            assert not Path(fact["source"]["file"]).is_absolute()


@pytest.mark.parametrize("identifier,body,title,author,abstract,columns,text_height", [
    ("acm-jacm-article-publication", 10, 14.4, 10.95, 9, 1, 574),
    ("acm-dlt-article-publication", 10, 14.4, 12, 9, 1, 550),
    ("acm-tog-article-publication", 9, 17.28, 12, 8, 2, 615),
])
def test_actual_ams_inherited_sizes_and_rounded_geometry(identifier, body, title, author, abstract, columns, text_height):
    item = read(identifier)
    roles = item["extracted_format"]["roles"]
    assert [roles[r]["size_pt"]["value"] for r in ["body", "title", "author", "abstract"]] == [body, title, author, abstract]
    assert item["page"]["columns"] == columns
    assert item["extracted_format"]["page"]["text_height_pt"]["value"] == text_height
    assert item["page"]["height_mm"] - item["page"]["margin_top_mm"] - item["page"]["margin_bottom_mm"] == pytest.approx(text_height * 25.4 / 72.27)


@pytest.mark.parametrize("identifier", ["acm-jacm-article-publication", "acm-dlt-article-publication", "acm-tog-article-publication"])
def test_generic_application_preserves_manuscript_and_installs_no_publisher_files(identifier, tmp_path):
    item = read(identifier)
    (tmp_path / "source").mkdir()
    source = tmp_path / "source/main.tex"
    source.write_text(r"""\documentclass{article}
\title{Synthetic title}
\author{Protected Author One \and Protected Author Two}
\begin{document}
\maketitle
\begin{abstract}Synthetic abstract with value 12.5.\end{abstract}
\section{Methods}
Protected body value 42. A citation~\cite{test}.
\begin{equation}E=mc^2\end{equation}
\begin{thebibliography}{1}\bibitem{test}Protected reference text.\end{thebibliography}
\end{document}
""", encoding="utf-8")
    before = source.read_bytes()
    result = apply_latex(source, tmp_path / "formatted", item)
    assert result["audit"]["content"] == "PASS"
    assert source.read_bytes() == before
    assert not any(p.suffix in {".cls", ".sty", ".bst"} for p in (tmp_path / "formatted").iterdir())
    output = (tmp_path / "formatted/main.tex").read_text(encoding="utf-8")
    assert "Protected Author One" in output and "Protected Author Two" in output
    assert "acmart" not in output
    layout = (tmp_path / "formatted/manuweave-layout.tex").read_text(encoding="utf-8")
    assert "fontsize" in layout and "geometry" in layout


def test_manuscript_spacing_and_conditional_rules_are_retained_without_author_rewriting():
    item = read("acm-manuscript-review-anonymous", "acm")
    assert item["rules"]["body"]["line_spacing_pt"] == pytest.approx(13.75 * POINT_FACTOR)
    assert item["rules"]["body"]["size_pt"] == pytest.approx(9 * POINT_FACTOR)
    assert item["latex"]["options"] == ["10pt"]
    assert item["extracted_format"]["roles"]["heading3"]["run_in_gap_pt"]["value"] == 3.5
    assert item["extracted_format"]["captions"]["multiline_alignment"]["value"] == "justify"
    assert "case_transform" not in item["rules"]["author"]
    assert not item["latex"].get("pre_class_packages")


def test_all_ten_general_acm_modes_are_declarative():
    profiles = list((ROOT / "acm/profiles").glob("*.json"))
    assert len(profiles) == 10
    for path in profiles:
        item = json.loads(path.read_text(encoding="utf-8"))
        StyleProfile.model_validate(item, strict=True)
        validate_template(item)
        assert item["template"]["assets"] == []
        assert item["latex"]["documentclass"] == "article"
        assert item["rules"]["body"]["font"] == "Linux Libertine"
        assert item["extracted_format"]["references"]["article_title_case"]["value"] == "preserve input case"


@pytest.mark.parametrize("identifier", ["acm-jacm-article-publication", "acm-dlt-article-publication", "acm-tog-article-publication"])
def test_extracted_values_apply_to_word_with_minimum_line_boxes(identifier, tmp_path):
    from docx import Document
    from manuweave.word import apply_docx, inspect_docx

    item = read(identifier)
    roles = ["title", "subtitle", "author", "affiliation", "correspondence", "abstract", "keywords",
             "heading1", "body", "heading2", "heading3", "heading4", "figure_caption", "table_caption",
             "table_body", "equation", "references_heading", "reference"]
    source = tmp_path / "source.docx"
    document = Document()
    for role in roles:
        document.add_paragraph("Protected " + role + " text with value 12.5.")
    document.save(source)
    before = inspect_docx(source)
    role_map = {block["id"]: role for block, role in zip(before["blocks"], roles)}
    output = tmp_path / "formatted.docx"
    result = apply_docx(source, output, item, roles=role_map)
    assert result["audit"]["content"]["status"] == "PASS"
    assert result["audit"]["style"]["status"] == "PASS"
    after = inspect_docx(output)
    assert [block["text"] for block in after["blocks"]] == [block["text"] for block in before["blocks"]]
    body = next(block for block in after["blocks"] if "Protected body text" in block["text"])
    assert body["format"]["line_spacing_rule"] == "atLeast"
    assert body["format"]["line_spacing_pt"] == pytest.approx(item["rules"]["body"]["line_spacing_pt"], abs=.03)
    assert body["format"]["size_pt"] == pytest.approx(item["rules"]["body"]["size_pt"], abs=.25)


def test_builder_reproduces_every_profile_from_json_without_publisher_files(tmp_path):
    import runpy
    import shutil

    source_dir = tmp_path / "computer-acm/sources"
    source_dir.mkdir(parents=True)
    (tmp_path / "acm/profiles").mkdir(parents=True)
    for filename in ("journals.json", "extracted-layouts.json"):
        shutil.copyfile(ROOT / "computer-acm/sources" / filename, source_dir / filename)
    build = runpy.run_path(str(ROOT / "computer-acm/sources/build_profiles.py"))["build"]
    build.__globals__["ROOT"] = tmp_path
    build.__globals__["FAMILY"] = tmp_path / "computer-acm"
    build()
    outputs = [*tmp_path.glob("computer-acm/profiles/*.json"), *tmp_path.glob("acm/profiles/*.json")]
    assert len(outputs) == 51
    for path in outputs:
        assert path.read_bytes() == (ROOT / path.relative_to(tmp_path)).read_bytes()
    assert not list(tmp_path.rglob("*.cls"))
    assert not list(tmp_path.rglob("*.sty"))
    assert not list(tmp_path.rglob("*.bst"))


def test_contract_binds_all_executed_values_and_preserves_source_point_precision():
    for family in ("computer-acm", "acm"):
        for path in (ROOT / family / "profiles").glob("*.json"):
            item = json.loads(path.read_text(encoding="utf-8"))
            requirements = item["template"]["requirements"]
            indexed = {(rule["target"], rule["property"]): rule for rule in requirements}
            groups = [*item["rules"].items(), ("page", item["page"]),
                      *item["running"].items(), ("sections", item.get("sections", {}))]
            for target, properties in groups:
                for prop, expected in properties.items():
                    rule = indexed[target, prop]
                    assert rule["expected"] == expected
                    assert rule["certainty"] == "verified" and rule["checking"] == "xml"
                    assert rule["applies_to"] == ["docx", "latex"]
                    assert "source_projection" in rule
            body_size = indexed["body", "size_pt"]
            assert body_size["source_projection"]["source_value"] == item["extracted_format"]["roles"]["body"]["size_pt"]["value"]
            assert body_size["source_projection"]["source_unit"] == "tex_pt"
            assert body_size["source_projection"]["operation"] == "tex_pt_to_bp"
            baseline = indexed["body", "line_spacing_rule"]
            assert baseline["expected"] == "atLeast"
            assert baseline["source_projection"]["source_unit"] == "tex_pt"
            assert baseline["source_projection"]["source_value"] == item["extracted_format"]["roles"]["body"]["line_spacing_pt"]["value"]
            assert baseline["source_projection"]["operation"] == "tex_natural_baseline_to_word_minimum"


@pytest.mark.parametrize("engine", ["docx", "latex"])
@pytest.mark.parametrize("location", ["role", "page", "running", "sections"])
def test_plan_cannot_apply_modified_acm_values_without_matching_source_contract(engine, location, tmp_path):
    from manuweave import pipeline

    profile = read("acm-tog-article-publication")
    if engine == "docx":
        from docx import Document
        source = tmp_path / "source.docx"
        document = Document()
        document.add_paragraph("Protected source content, number 12.5.")
        document.save(source)
        output = tmp_path / "blocked.docx"
    else:
        (tmp_path / "source").mkdir()
        source = tmp_path / "source/main.tex"
        source.write_text(r"\documentclass{article}\begin{document}Protected source number 12.5.\end{document}", encoding="utf-8")
        output = tmp_path / "blocked"
    before = source.read_bytes()
    plan = pipeline.plan(source, profile)
    if location == "role":
        plan["profile"]["rules"]["body"]["size_pt"] += 1
    elif location == "page":
        plan["profile"]["page"]["width_mm"] += 1
    elif location == "running":
        plan["profile"]["running"]["header"]["size_pt"] += 1
    else:
        plan["profile"]["sections"]["body_column_gap_mm"] += 1
    with pytest.raises(ValueError, match="Unconfirmed journal formatting rule"):
        pipeline.apply(plan, output)
    assert not output.exists()
    assert source.read_bytes() == before
