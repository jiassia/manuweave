import json
import runpy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement

from manuweave.presets import load_preset
from manuweave.templates import validate_template
from manuweave.word import apply_docx, inspect_docx
from manuweave.latex import apply_latex

ROOT = Path(__file__).resolve().parents[1] / 'src/manuweave/data/templates/computer-special'


def test_ijoc_profile_is_current_numeric_layout_without_upstream_dependencies():
    profile = load_preset('comp-informs-ijoc-submission')
    validate_template(profile)
    assert profile['catalog_journal_id'] == 'ccf-2026-informs-journal-on-computing'
    assert profile['current_coverage'] is True
    assert profile['stage'] == 'review_submission'
    assert profile['latex'] == {} and profile['template']['assets'] == []
    assert 'fixture' not in profile['template']
    assert profile['page'] == {'columns': 1, **{f'margin_{edge}_mm': 25.4 for edge in ['top', 'bottom', 'left', 'right']}}
    for role in ('body', 'abstract', 'reference'):
        assert profile['rules'][role] == {'size_pt': 12, 'line_spacing': 1.5}
        assert 'font' not in profile['rules'][role]
    assert not {'width_mm', 'height_mm'} & profile['page'].keys()
    assert 'sha256' not in profile['sources'][0]
    assert 'Raw HTTP response bytes were not obtained' in profile['sources'][0]['source_verification_scope']
    assert profile['extracted_format']['references']['ordering']['value'] == 'alphabetical_by_author'


def test_ijoc_word_application_preserves_original_numbers_math_and_table(tmp_path):
    source = tmp_path / 'source.docx'
    doc = Document()
    doc.add_paragraph('Synthetic IJOC fixture', 'Title')
    doc.add_paragraph('Abstract: An original software test fixture.')
    doc.add_paragraph('1 Methods', 'Heading 1')
    doc.add_paragraph('This original test sentence retains 0.125.')
    math = OxmlElement('m:oMath')
    run = OxmlElement('m:r')
    text = OxmlElement('m:t')
    text.text = 'E=mc²'
    run.append(text)
    math.append(run)
    doc.add_paragraph()._p.append(math)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = '12'
    doc.add_paragraph('References')
    doc.add_paragraph('[1] Synthetic local reference, 2026.')
    doc.save(source)
    profile = load_preset('comp-informs-ijoc-submission')
    out = tmp_path / 'formatted.docx'
    result = apply_docx(source, out, profile)
    assert result['audit']['content']['status'] == result['audit']['style']['status'] == 'PASS'
    assert inspect_docx(source)['inventory'] == inspect_docx(out)['inventory']
    assert inspect_docx(source)['objects'] == inspect_docx(out)['objects']


def test_ijoc_latex_has_bp_dimensions_and_does_not_require_a_publisher_class(tmp_path):
    source = tmp_path / 'source'
    source.mkdir()
    (source / 'main.tex').write_text(r'\documentclass[12pt]{article}\begin{document}\begin{abstract}Synthetic abstract.\end{abstract}\section{Methods}Value 0.125.\begin{equation}E=mc^2\end{equation}\begin{thebibliography}{1}\bibitem{x}Synthetic reference.\end{thebibliography}\end{document}', encoding='utf-8')
    result = apply_latex(source / 'main.tex', tmp_path / 'formatted', load_preset('comp-informs-ijoc-submission'))
    assert result['audit']['content'] == result['audit']['style'] == 'PASS'
    layout = (tmp_path / 'formatted/manuweave-layout.tex').read_text(encoding='utf-8')
    assert '12bp' in layout and '25.4mm' in layout
    assert not list((tmp_path / 'formatted').glob('*.cls'))


def test_ijoc_builder_matches_the_checked_in_source_facts():
    module = runpy.run_path(str(ROOT / 'builder_siam_informs.py'))
    profiles = module['profiles']()
    assert len(profiles) == 1
    assert profiles[0] == json.loads((ROOT / 'profiles/comp-informs-ijoc-submission.json').read_text(encoding='utf-8'))
