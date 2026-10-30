"""Physical source units, branch differences, and editable output acceptance."""
import json
import runpy
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement

from manuweave.latex import apply_latex, inspect_applied_rules, inspect_latex
from manuweave.models import StyleProfile
from manuweave.presets import load_preset
from manuweave.templates import validate_template
from manuweave.word import apply_docx, inspect_docx

ROOT = Path(__file__).resolve().parents[1] / 'src/manuweave/data/templates/computer-special'
FACTS = json.loads((ROOT / 'tacl-fi-source-facts.json').read_text(encoding='utf-8'))['verified_layouts']
NAMES = [p['id'] for p in FACTS]
BP = 72 / 72.27


@pytest.mark.parametrize('identifier', NAMES)
def test_explicit_physical_values_have_their_own_source_metadata(identifier):
    profile = load_preset(identifier)
    StyleProfile.model_validate(profile, strict=True)
    validate_template(profile)
    assert profile['template']['assets'] == []
    assert profile['latex']['documentclass'] == 'article'
    assert len(profile['rules']) >= 10
    assert len(profile['template']['requirements']) >= 90
    for target, values in [*profile['rules'].items(), ('page', profile['page'])]:
        detailed = profile['extracted_format']['page'] if target == 'page' else profile['extracted_format']['roles'][target]
        for prop, value in values.items():
            assert detailed[prop]['value'] == value
            assert detailed[prop]['unit']
            assert detailed[prop]['source']['locator']
            assert len(detailed[prop]['source']['sha256']) == 64
    for role, values in profile['rules'].items():
        detailed = profile['extracted_format']['roles'][role]
        assert values['size_pt'] == pytest.approx(detailed['source_size_tex_pt']['value'] * BP)
        assert values['line_spacing_pt'] == pytest.approx(detailed['source_baseline_tex_pt']['value'] * BP)
        assert values['line_spacing_rule'] == 'atLeast'


def test_tacl_records_real_template_geometry_separately_from_nominal_prose():
    submission = load_preset('comp-tacl-2021-submission')
    accepted = load_preset('comp-tacl-2021-accepted')
    assert submission['rules']['title']['size_pt'] == pytest.approx(14.4 * BP)
    assert submission['rules']['body']['size_pt'] == pytest.approx(10.95 * BP)
    assert submission['rules']['heading1']['size_pt'] == pytest.approx(12 * BP)
    assert submission['rules']['heading2']['size_pt'] == pytest.approx(10.95 * BP)
    assert submission['page']['margin_left_mm'] == 25.4
    assert submission['page']['margin_right_mm'] == 24.6
    assert submission['page']['margin_top_mm'] == pytest.approx(25.4 - 5 + 5 * 25.4 / 72.27)
    ex = submission['extracted_format']
    assert ex['nominal_guide_values']['title_size_tex_pt']['value'] == 15
    assert ex['nominal_guide_values']['all_margins_mm']['value'] == 25
    assert ex['page']['full_width_title_box_height_mm']['value'] == 56.35
    assert accepted['extracted_format']['page']['full_width_title_box_height_mm']['value'] == 50
    assert 'author' not in submission['rules'] and 'author' in accepted['rules']
    assert 'anonymous' in ex['roles']['author']['font']['condition']
    assert ex['roles']['heading4']['run_in']['value'] is True
    assert ex['references']['style']['value'].startswith('unnumbered author-date')


def test_fundam_includehead_geometry_and_author_row_sizes_are_not_flattened():
    profile = load_preset('comp-fundamenta-informaticae-34')
    assert profile['page']['width_mm'] == 192
    assert profile['page']['height_mm'] == 262
    assert profile['page']['columns'] == 1
    assert profile['page']['margin_top_mm'] == pytest.approx(28 + 37 * 25.4 / 72.27)
    assert profile['rules']['author']['size_pt'] == pytest.approx(10.95 * BP)
    assert profile['rules']['affiliation']['size_pt'] == pytest.approx(10 * BP)
    assert profile['rules']['affiliation']['italic'] is True
    assert profile['rules']['affiliation']['line_spacing_pt'] == pytest.approx(16.32 * BP)
    assert profile['rules']['heading1']['size_pt'] == pytest.approx(14.4 * BP)
    assert profile['rules']['heading2']['size_pt'] == pytest.approx(12 * BP)
    assert profile['rules']['reference']['size_pt'] == pytest.approx(10 * BP)
    assert 'left_indent_pt' not in profile['rules']['reference']
    assert profile['extracted_format']['roles']['abstract_heading']['inline_label']['value'] is True
    assert profile['extracted_format']['table_row_mechanics']['arraystretch']['value'] == 1.2


@pytest.mark.parametrize('identifier', NAMES)
def test_word_formats_every_declared_role_and_preserves_equation_and_table(tmp_path, identifier):
    profile = load_preset(identifier)
    document = Document()
    mapping = {}
    for i, role in enumerate(profile['rules']):
        document.add_paragraph('Synthetic ' + role + ' with constant 12.5')
        mapping['p' + str(i)] = role
    math = OxmlElement('m:oMath')
    run, value = OxmlElement('m:r'), OxmlElement('m:t')
    value.text = 'E=mc²'
    run.append(value)
    math.append(run)
    document.add_paragraph()._p.append(math)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = 'Merged synthetic heading'
    table.cell(1, 0).text = 'A'
    table.cell(1, 1).text = '12.5'
    source, output, repeat = (tmp_path / name for name in ['source.docx', 'output.docx', 'repeat.docx'])
    document.save(source)
    before = inspect_docx(source)
    result = apply_docx(source, output, profile, mapping)
    assert result['audit']['content']['status'] == 'PASS'
    assert result['audit']['style']['status'] == 'PASS'
    after = inspect_docx(output)
    assert before['objects'] == after['objects']
    assert before['inventory'] == after['inventory']
    again = apply_docx(output, repeat, profile, mapping)
    assert again['audit']['style']['status'] == 'PASS'


def synthetic_tex(profile):
    affiliation = r'\affiliation{Synthetic Institution}' if 'affiliation' in profile['rules'] else ''
    keywords = r'\keywords{synthetic, formatting}' if 'keywords' in profile['rules'] else ''
    return r'''\documentclass{article}
\usepackage{amsmath}
\title{Synthetic Format Study}
\author{Synthetic Author}
''' + affiliation + r'''
\begin{document}
\maketitle
\begin{abstract}Synthetic editable abstract with constant 12.5.\end{abstract}
''' + keywords + r'''
\section{Methods}
Synthetic body text with constant 12.5.
\subsection{Objects}
Synthetic subsection text.
\subsubsection{Preservation}
Synthetic third level text.
\begin{equation}E=mc^2\end{equation}
\begin{table}[ht]\centering\begin{tabular}{lr}Item&Value\\A&12.5\end{tabular}\caption{Synthetic table}\end{table}
\begin{figure}[ht]\centering\fbox{Synthetic diagram}\caption{Synthetic figure}\end{figure}
\begin{thebibliography}{9}\bibitem{a}Synthetic Author. Synthetic Reference. 2026.\end{thebibliography}
\end{document}
'''


@pytest.mark.parametrize('identifier', NAMES)
def test_latex_applies_numeric_rules_without_original_publisher_programs(tmp_path, identifier):
    profile = load_preset(identifier)
    source_directory = tmp_path / 'source'
    source_directory.mkdir()
    source = source_directory / 'source.tex'
    source.write_text(synthetic_tex(profile), encoding='utf-8')
    before = inspect_latex(source)
    output = tmp_path / 'formatted'
    report = apply_latex(source, output, profile)
    assert report['audit']['content'] == 'PASS'
    assert report['audit']['style'] == 'UNVERIFIED'
    checks = inspect_applied_rules(output, profile)
    assert not any(check['status'] == 'FAIL' for check in checks['checks'])
    unresolved = {check['reason'] for check in checks['checks'] if check['status'] == 'UNVERIFIED'}
    assert unresolved <= {
        'heading_label_hanging_requires_structural_rule',
        'caption_hanging_requires_label_specific_rule',
        'list_or_table_spacing_requires_structural_rule',
    }
    after = inspect_latex(output)
    assert before['inventory']['source.tex']['content_sha256'] == after['inventory']['source.tex']['content_sha256']
    assert not list(output.glob('*.cls')) and not list(output.glob('*.sty')) and not list(output.glob('*.bst'))
    layout = (output / 'manuweave-layout.tex').read_text(encoding='utf-8')
    assert 'ptm' in layout and 'bp' in layout


def test_builder_reproduces_all_retained_profiles():
    expected = runpy.run_path(str(ROOT / 'builder_tacl_fi.py'))['profiles']()
    assert {p['id'] for p in expected} == set(NAMES)
    for profile in expected:
        actual = json.loads((ROOT / 'profiles' / (profile['id'] + '.json')).read_text(encoding='utf-8'))
        assert profile == actual
