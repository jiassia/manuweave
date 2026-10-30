from pathlib import Path
import subprocess

from docx import Document
import pytest

from manuweave import libreoffice


def make_source(path):
    doc = Document()
    doc.add_paragraph('Synthetic private manuscript')
    doc.add_paragraph('Original measurement 12.5 must remain unchanged.')
    doc.save(path)
    return path


def test_unavailable_never_downloads_or_creates_output(tmp_path, monkeypatch):
    monkeypatch.delenv('MANUWEAVE_SOFFICE', raising=False)
    monkeypatch.setattr(libreoffice.shutil, 'which', lambda name: None)
    monkeypatch.setattr(libreoffice.subprocess, 'run', lambda *a, **k: pytest.fail('No process should start without an available renderer'))
    assert libreoffice.available()['status'] == 'UNVERIFIED'
    source = make_source(tmp_path/'source.docx')
    original = source.read_bytes()
    output = tmp_path/'render'
    report = libreoffice.export_pdf(source,output)
    assert report['status'] == 'UNVERIFIED'
    assert not output.exists()
    assert source.read_bytes() == original


def test_existing_output_is_never_overwritten(tmp_path):
    source = make_source(tmp_path/'source.docx')
    output = tmp_path/'render'
    output.mkdir()
    marker = output/'keep.txt'
    marker.write_text('Existing result')
    with pytest.raises(FileExistsError):
        libreoffice.export_pdf(source,output)
    assert marker.read_text() == 'Existing result'


def test_external_resources_are_rejected_before_backend(tmp_path, monkeypatch):
    from docx.opc.constants import RELATIONSHIP_TYPE
    doc = Document()
    doc.add_paragraph('Synthetic unsafe external image relationship')
    doc.part.relate_to('https://example.invalid/manuscript-leak',RELATIONSHIP_TYPE.IMAGE,is_external=True)
    source = tmp_path/'external.docx'
    doc.save(source)
    monkeypatch.setattr(libreoffice,'available',lambda:pytest.fail('Unsafe source must fail before renderer selection'))
    output = tmp_path/'render'
    report = libreoffice.export_pdf(source,output)
    assert report['status'] == 'FAIL'
    assert 'External' in report['log']
    assert not output.exists()


def test_invalid_docx_fails_without_starting_renderer(tmp_path, monkeypatch):
    source = tmp_path/'bad.docx'
    source.write_bytes(b'not a ZIP or Word document')
    monkeypatch.setattr(libreoffice,'available',lambda:pytest.fail('Malformed source must fail before renderer selection'))
    assert libreoffice.export_pdf(source,tmp_path/'render')['status']=='FAIL'
    assert not (tmp_path/'render').exists()


def test_available_checks_only_existing_local_programs(tmp_path, monkeypatch):
    monkeypatch.delenv('MANUWEAVE_SOFFICE', raising=False)
    queried=[]
    def which(name):
        queried.append(name)
        return None
    monkeypatch.setattr(libreoffice.shutil,'which',which)
    assert libreoffice.available()['status']=='UNVERIFIED'
    assert queried==['soffice','libreoffice']
    explicit=tmp_path/'soffice.exe'
    explicit.write_bytes(b'local path fixture, not executed')
    monkeypatch.setenv('MANUWEAVE_SOFFICE',str(explicit))
    backend=libreoffice.available()
    assert backend['mode']=='local'
    assert backend['executable']==str(explicit.resolve())
    assert queried==['soffice','libreoffice']


def test_local_timeout_profile_and_source_protection(tmp_path, monkeypatch):
    from urllib.parse import urlparse
    from urllib.request import url2pathname
    source=make_source(tmp_path/'source.docx')
    original=source.read_bytes()
    executable=tmp_path/'soffice.exe'
    monkeypatch.setattr(libreoffice,'available',lambda:{'status':'PASS','mode':'local','executable':str(executable),'renderer':'LibreOffice local'})
    profiles=[]
    def call(command,timeout):
        assert command[0]==str(executable)
        assert timeout==90
        assert '--headless' in command and '--nofirststartwizard' in command
        assert command[command.index('--convert-to')+1]=='pdf:writer_pdf_Export'
        uri=next(a.split('=',1)[1] for a in command if a.startswith('-env:UserInstallation='))
        profile=Path(url2pathname(urlparse(uri).path))
        profiles.append(profile)
        config=(profile/'user'/'registrymodifications.xcu').read_text(encoding='utf-8')
        assert 'MacroSecurityLevel' in config and '<value>3</value>' in config
        assert 'UpdateLinksMode' in config
        assert Path(command[-1]).read_bytes()==original
        assert Path(command[-1])!=source
        raise subprocess.TimeoutExpired(command,timeout)
    monkeypatch.setattr(libreoffice,'_call',call)
    report=libreoffice.export_pdf(source,tmp_path/'render')
    assert report['status']=='FAIL'
    assert source.read_bytes()==original
    assert all(not profile.exists() for profile in profiles)


def test_timeout_terminates_own_descendant_process(tmp_path):
    import sys
    import time
    ready=tmp_path/'child-ready'
    completed=tmp_path/'child-completed'
    child=tmp_path/'child.py'
    child.write_text('from pathlib import Path\nimport time\n'
                     +f'Path({str(ready)!r}).write_text("ready")\n'
                     +'time.sleep(3)\n'
                     +f'Path({str(completed)!r}).write_text("unexpected survivor")\n',encoding='utf-8')
    parent=tmp_path/'parent.py'
    parent.write_text('import subprocess,sys,time\n'
                      +f'subprocess.Popen([sys.executable,{str(child)!r}])\n'
                      +'time.sleep(30)\n',encoding='utf-8')
    with pytest.raises(subprocess.TimeoutExpired):
        libreoffice._call([sys.executable,str(parent)],timeout=2)
    assert ready.exists(), 'Child must actually have started before testing tree cleanup'
    time.sleep(1.3)
    assert not completed.exists()


def test_real_libreoffice_export_preserves_original(tmp_path):
    if libreoffice.available()['status']!='PASS':
        pytest.skip('Optional installed local LibreOffice is unavailable; no setup is attempted')
    source=make_source(tmp_path/'source.docx')
    from docx.oxml import OxmlElement
    doc=Document(source)
    doc.add_paragraph('中文字体与数学公式测试。')
    paragraph=doc.add_paragraph()
    math=OxmlElement('m:oMath')
    run=OxmlElement('m:r')
    text=OxmlElement('m:t')
    text.text='E=mc²'
    run.append(text)
    math.append(run)
    paragraph._p.append(math)
    doc.save(source)
    original=source.read_bytes()
    report=libreoffice.export_pdf(source,tmp_path/'render')
    assert report['status']=='PASS',report
    assert Path(report['pdf']).read_bytes().startswith(b'%PDF-')
    assert source.read_bytes()==original
    import pypdfium2 as pdfium
    with pdfium.PdfDocument(report['pdf']) as pdf:
        page=pdf[0]
        text=page.get_textpage()
        rendered=text.get_text_range()
        assert '12.5' in rendered
        assert '中文字体与数学公式测试' in rendered
        assert 'E=mc²' in ''.join(rendered.split())
        text.close()
        page.close()
