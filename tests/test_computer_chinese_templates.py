"""Measured manuscript template typography, bilingual roles and native edits."""
from __future__ import annotations

import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches
from lxml import etree
from PIL import Image

from manuweave.models import StyleProfile
from manuweave.presets import load_preset
from manuweave.templates import validate_template
from manuweave.word import apply_docx, audit_docx, inspect_docx

DATA = Path(__file__).resolve().parents[1] / "src/manuweave/data/templates/computer-chinese"
PROFILES = sorted((DATA / "profiles").glob("*.json"))
FACTS = json.loads((DATA / "source-facts.json").read_text(encoding="utf-8"))["documents"]


def _fixture(directory, marked_body=False):
    image = directory / "synthetic.png"
    Image.new("RGB", (120, 40), "white").save(image)
    doc = Document()
    doc.add_paragraph("合成模板格式验证", "Title")
    doc.add_paragraph("Synthetic manuscript typography", "Title")
    doc.add_paragraph("合成作者甲，合成作者乙")
    doc.add_paragraph("Synthetic Author A, Synthetic Author B")
    doc.add_paragraph("合成单位，合成城市 000000")
    doc.add_paragraph("Synthetic Institute, Synthetic City 000000")
    doc.add_paragraph("通信作者：合成作者甲，test@example.invalid")
    doc.add_paragraph("Correspondence: Synthetic Author A, test@example.invalid")
    doc.add_paragraph("摘要：本段仅为软件测试材料，无科研结论。")
    doc.add_paragraph("Abstract: This is synthetic test material without research claims.")
    doc.add_paragraph("关键词：原生文档；合成数据；格式验证")
    doc.add_paragraph("Keywords: native document; synthetic data; format verification")
    if marked_body:
        doc.add_paragraph("未编号引言的合成正文，受保护数值为 6。")
    doc.add_paragraph("0 引言", "Heading 1")
    doc.add_paragraph("受保护的合成数值为 12。测试参考标记为 [1]。")
    doc.add_paragraph("1.1 处理流程", "Heading 2")
    doc.add_paragraph("第二个受保护的合成数值为 24。")
    doc.add_paragraph("1.1.1 实现步骤", "Heading 3")
    paragraph = doc.add_paragraph()
    math = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.text = "E=mc²"
    run.append(text)
    math.append(run)
    paragraph._p.append(math)
    doc.add_paragraph("表1 合成数值")
    doc.add_paragraph("Table 1 Synthetic values")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "合成分组"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "12"
    table.cell(2, 0).text = "B"
    table.cell(2, 1).text = "24"
    doc.add_picture(str(image), width=Inches(2))
    doc.add_paragraph("图1 合成示意图")
    doc.add_paragraph("Figure 1 Synthetic diagram")
    doc.add_paragraph("参考文献")
    doc.add_paragraph("[1] Synthetic Author. 2026. Invented software test material.")
    source = directory / "synthetic.docx"
    doc.save(source)
    role_by_text = {
        "Synthetic manuscript typography": "title_en",
        "合成作者甲，合成作者乙": "author_zh",
        "Synthetic Author A, Synthetic Author B": "author_en",
        "合成单位，合成城市 000000": "affiliation_zh",
        "Synthetic Institute, Synthetic City 000000": "affiliation_en",
        "通信作者：合成作者甲，test@example.invalid": "correspondence_zh",
        "Correspondence: Synthetic Author A, test@example.invalid": "correspondence_en",
    }
    if marked_body:
        role_by_text["未编号引言的合成正文，受保护数值为 6。"] = "body_start"
    roles = {block["id"]: role_by_text[block["text"]] for block in inspect_docx(source)["blocks"]
             if block["text"] in role_by_text}
    return source, roles


@pytest.mark.parametrize("profile_file", PROFILES, ids=lambda p: p.stem)
def test_each_computer_template_has_located_word_source_evidence(profile_file):
    profile = json.loads(profile_file.read_text(encoding="utf-8"))
    StyleProfile.model_validate(profile, strict=True)
    validate_template(profile)
    assert load_preset(profile["id"])["id"] == profile["id"]
    fact = FACTS[profile["id"]]
    source = profile["sources"][0]
    assert source["sha256"] == fact["source"]["sha256"]
    assert len(source["sha256"]) == 64 and source["byte_length"] > 10000
    assert source["source_verified"] and not source["original_file_redistributed"]
    assert profile["template"]["assets"] == []
    assert profile["template"]["engine_support"] == {"docx": "partial", "latex": "unsupported"}
    assert profile["template"]["license"] == "rules_only_source_not_redistributed"
    requirements = profile["template"]["requirements"]
    for target, properties in [*profile["rules"].items(), ("page", profile["page"]),
                               ("sections", profile.get("sections", {}))]:
        for prop, expected in properties.items():
            assert any(r["target"] == target and r["property"] == prop and r["expected"] == expected
                       and r["certainty"] == "verified" and r["checking"] == "xml"
                       and r["applies_to"] == ["docx"] and r["source"]["locator"] for r in requirements)


@pytest.mark.parametrize("profile_file", PROFILES, ids=lambda p: p.stem)
def test_each_computer_template_formats_bilingual_native_content(profile_file, tmp_path):
    profile = json.loads(profile_file.read_text(encoding="utf-8"))
    marked_body = profile.get("sections", {}).get("body_starts_at") == "marked_body_start"
    source, roles = _fixture(tmp_path, marked_body=marked_body)
    output = tmp_path / "formatted.docx"
    result = apply_docx(source, output, profile, roles=roles)
    assert result["audit"]["content"]["status"] == "PASS"
    assert result["audit"]["style"]["status"] == "PASS"
    before, after = inspect_docx(source), inspect_docx(output)
    mutable = {"word/document.xml"}
    if "mirror_margins" in profile["page"]:
        mutable.add("word/settings.xml")
    assert {k: v for k, v in before["inventory"].items() if k not in mutable} == {
        k: v for k, v in after["inventory"].items() if k not in mutable}
    assert [b["text"] for b in before["blocks"]] == [b["text"] for b in after["blocks"]]
    assert after["objects"]["equations"] == after["objects"]["images"] == after["objects"]["tables"] == 1
    indexed = {block["text"]: block for block in after["blocks"]}
    for language, title in (("zh", "合成模板格式验证"), ("en", "Synthetic manuscript typography")):
        rule = profile["rules"].get("title_" + language, profile["rules"].get("title", {}))
        for prop, expected in rule.items():
            if prop == "alignment" and expected == "justify":
                expected = "both"
            assert indexed[title]["format"].get(prop) == expected
    for language, text in (("zh", "摘要：本段仅为软件测试材料，无科研结论。"),
                           ("en", "Abstract: This is synthetic test material without research claims.")):
        rule = profile["rules"].get("abstract_" + language, profile["rules"].get("abstract", {}))
        for prop, expected in rule.items():
            if prop == "alignment" and expected == "justify":
                expected = "both"
            assert indexed[text]["format"].get(prop) == expected
    if profile.get("sections"):
        assert [page["columns"] for page in after["page"]] == [1, 2]
    repeated = tmp_path / "repeated.docx"
    apply_docx(output, repeated, profile, roles=roles)
    assert audit_docx(output, repeated, profile, roles=roles)["style"]["status"] == "PASS"
    assert inspect_docx(repeated)["inventory"] == after["inventory"]
    with ZipFile(source) as original, ZipFile(output) as formatted:
        assert all(original.read(n) == formatted.read(n) for n in original.namelist() if n.startswith("word/media/"))
        assert b"gridSpan" in formatted.read("word/document.xml")
        math_ns = {"m": "http://schemas.openxmlformats.org/officeDocument/2006/math"}
        original_math = etree.fromstring(original.read("word/document.xml")).xpath("//m:oMath", namespaces=math_ns)
        formatted_math = etree.fromstring(formatted.read("word/document.xml")).xpath("//m:oMath", namespaces=math_ns)
        assert [etree.tostring(m, method="c14n") for m in original_math] == [
            etree.tostring(m, method="c14n") for m in formatted_math]


def test_real_template_sizes_and_instruction_runs_are_not_generic_defaults():
    cjc = load_preset("cn-cjc-word-template")
    jos = load_preset("cn-jos-word-2025")
    science = load_preset("cn-jsjkx-word-2024")
    jcad = load_preset("cn-jcad-template-2026")
    assert cjc["rules"]["title_zh"]["size_pt"] == 22
    assert cjc["rules"]["abstract_zh"]["size_pt"] == 9
    assert cjc["rules"]["abstract_en"]["size_pt"] == 10.5
    assert science["rules"]["title_zh"]["size_pt"] == 16
    assert jos["rules"]["body"]["size_pt"] == 9
    assert jos["rules"]["abstract_en"]["size_pt"] == 7.5
    assert abs(jos["page"]["width_mm"] - 184) < .05
    assert abs(jos["page"]["height_mm"] - 260) < .05
    assert jcad["page"]["margin_top_mm"] < 5
    assert jcad["page"]["margin_bottom_mm"] < 5
    assert len(PROFILES) >= 10
    assert not any(p.suffix in {".doc", ".docx", ".pdf", ".tex", ".sty", ".cls", ".png"} for p in DATA.rglob("*"))
