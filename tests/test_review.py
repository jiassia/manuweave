
import pytest
from docx import Document
from PIL import Image

from manuweave import pipeline
from manuweave.fonts import check_fonts
from manuweave.fonts import pdf_font_family
from manuweave.io import digest, read_json
from manuweave.review import accept


def test_font_inventory_reports_missing():
    report = check_fonts({"rules": {"body": {"font": "ManuWeaveDeliberatelyMissingFontZ92"}}})
    assert report["status"] in {"FAIL", "UNVERIFIED"}
    if report["status"] == "FAIL":
        assert report["missing"] == ["ManuWeaveDeliberatelyMissingFontZ92"]


def test_rendered_font_family_names_are_not_substitutes():
    assert pdf_font_family("ABCDEF+TimesNewRomanPSMT") == pdf_font_family("Times New Roman")
    assert pdf_font_family("TimesNewRomanPS-BoldMT") == pdf_font_family("Times New Roman")
    assert pdf_font_family("LinLibertineO") == pdf_font_family("Linux Libertine O")
    assert pdf_font_family("LinBiolinumOB") == pdf_font_family("Linux Biolinum O")
    assert pdf_font_family("Calibri") != pdf_font_family("Linux Libertine O")
    assert pdf_font_family("SimSun") == pdf_font_family("宋体")


def test_page_review_hash_coverage_and_gate(tmp_path, monkeypatch):
    # The page-review fixture controls rendering and its independent font check.
    monkeypatch.setattr("manuweave.fonts.check_fonts", lambda _: {"status": "PASS", "scope": "Synthetic font inventory"})
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("Sample manuscript")
    doc.save(source)

    def rendered(native, output, profile=None):
        output.mkdir()
        page = output / "page-1.png"
        Image.new("RGB", (20, 20), "white").save(page)
        return {"status": "PASS", "source_sha256": digest(native), "visual": {
            "status": "UNVERIFIED", "geometry": "PASS", "pages": [{"page": 1, "png": str(page), "sha256": digest(page)}]}}

    monkeypatch.setattr("manuweave.render.render", rendered)
    result = pipeline.run(source, tmp_path / "run", pipeline.resolve_profile("english-single"))
    assert result["status"] == "UNVERIFIED"
    render = read_json(tmp_path / "run" / "render.json")
    checks = dict.fromkeys(["glyphs", "overlaps", "equations", "tables", "figures", "heading_breaks", "headers_footers"], "PASS")
    review = {"reviewer": "test fixture", "method": "human_image_inspection", "source_sha256": render["source_sha256"],
              "pages": [{"page": 1, "sha256": "0" * 64, "checks": checks, "observations": "Synthetic gate test fixture"}]}
    with pytest.raises(ValueError, match="hash mismatch"):
        accept(result["state"], review)
    review["pages"][0]["sha256"] = render["visual"]["pages"][0]["sha256"]
    review["pages"][0]["checks"]["figures"] = "FAIL"
    assert accept(result["state"], review)["status"] == "FAIL"
    review["pages"][0]["checks"]["figures"] = "NOT_APPLICABLE"
    assert accept(result["state"], review)["status"] == "PASS"
    assert pipeline.run(state_path=result["state"])["status"] == "PASS"


def test_external_word_field_is_not_rendered(tmp_path):
    import zipfile
    from lxml import etree
    from manuweave.render import render
    from manuweave.word import W
    source = tmp_path / "field.docx"
    doc = Document()
    doc.add_paragraph("Field test")
    doc.save(source)
    with zipfile.ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    tree = etree.fromstring(parts["word/document.xml"])
    p = tree.find(".//" + W + "p")
    for text in ["INCLU", "DETEXT \"secret.txt\""]:
        run = etree.SubElement(p, W + "r")
        etree.SubElement(run, W + "instrText").text = text
    parts["word/document.xml"] = etree.tostring(tree)
    with zipfile.ZipFile(source, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    result = render(source, tmp_path / "blocked")
    assert result["status"] == "UNVERIFIED"
    assert "Active external" in result["reason"]
    assert not (tmp_path / "blocked").exists()
